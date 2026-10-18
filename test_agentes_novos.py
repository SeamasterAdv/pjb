import json
import os
import threading

import pytest
from openpyxl import load_workbook

from agentes import Orquestrador, planejar
from agentes.especialistas import AgenteEmail, AgenteResumoDia, _questao_valida
from google_gmail import GmailNaoConectado
from llm_nuvem import LLMComReserva
from test_agentes import LLMFalso


class LLMNovos(LLMFalso):
    """Responde também pelos agentes de planilha, estudos, resumo do dia e e-mail."""

    def __init__(self, cartoes=12, questoes=6, **kwargs):
        super().__init__(**kwargs)
        self.cartoes, self.questoes = cartoes, questoes
        self.prompts = []

    def chat(self, prompt, system_prompt=None, model=None, **kwargs):
        self.prompts.append((system_prompt, prompt))
        if "Agente de Planilhas" in system_prompt:
            return json.dumps({"titulo": "Controle de Gastos", "observacoes": "Valores de exemplo.", "abas": [{
                "nome": "Gastos", "colunas": ["Categoria", "Valor"], "tipos": ["texto", "moeda"],
                "linhas": [["Moradia", 1500], ["Mercado", "R$ 642,35"], ["Transporte", 230.5]], "total": True,
                "grafico": {"tipo": "barras", "titulo": "Gastos", "categorias": "Categoria", "valores": ["Valor"]}}]})
        if "Agente de Estudos" in system_prompt:
            return json.dumps({
                "titulo": "Revolução Francesa", "resumo": "Panorama da revolução de 1789.",
                "topicos": [{"titulo": "Causas", "explicacao": "Crise fiscal e desigualdade.", "pontos_chave": ["Crise fiscal da monarquia"]},
                            {"titulo": "Fases", "explicacao": "Monarquia constitucional, Convenção e Diretório.", "pontos_chave": ["Terror jacobino"]}],
                "flashcards": [{"frente": f"Pergunta {i}", "verso": f"Resposta {i}"} for i in range(self.cartoes)],
                "questoes": [{"enunciado": f"Questão {i}?", "alternativas": ["A) um", "B) dois", "C) três", "D) quatro"],
                              "correta": "B", "explicacao": "Porque sim."} for i in range(self.questoes)],
                "dicas": ["Revise em voz alta"]})
        if "resume manchetes" in system_prompt:
            return json.dumps({"temas": [{"tema": "direito", "resumo": "O STF julgou um caso importante."}]})
        if "Resumo do Dia" in system_prompt:
            return ("Bom dia! Hoje é terça-feira, 13 de outubro. Você tem uma reunião às 10h com a equipe. "
                    "Não há e-mails novos. Nas notícias de direito, o STF julgou um caso importante que pode mudar "
                    "a forma como contratos são revisados. Tenha um ótimo dia.")
        if "Agente de E-mail" in system_prompt:
            return json.dumps({"visao_geral": "Você tem dois e-mails; um pede resposta hoje.", "emails": [
                {"numero": 1, "prioridade": "alta", "resumo": "Cliente pede orçamento.", "acao": "responder",
                 "rascunho": "Olá, envio o orçamento até amanhã.\n\nAtenciosamente,"},
                {"numero": 2, "prioridade": "baixa", "resumo": "Newsletter.", "acao": "nenhuma", "rascunho": ""}]})
        return super().chat(prompt, system_prompt=system_prompt, model=model, **kwargs)


def executar(orquestrador, comando):
    pronto = threading.Event()
    tarefa = orquestrador.enviar(comando, ao_concluir=lambda t: pronto.set())
    assert tarefa is not None
    assert pronto.wait(30)
    return tarefa


def orquestrador(tmp_path, llm=None, **kwargs):
    return Orquestrador(str(tmp_path), llm=llm or LLMNovos(), abrir_ao_concluir=False, **kwargs)


# ------------------------------------------------------------ reconhecimento
@pytest.mark.parametrize("frase, tipo, plano", [
    ("crie uma planilha de controle de gastos mensais com gráfico", "planilha", ["planilha"]),
    ("faça uma planilha em Excel comparando os preços atuais dos carros elétricos", "planilha", ["pesquisa", "planilha"]),
    ("me ajude a estudar direito constitucional para a prova", "estudos", ["estudos"]),
    ("crie 20 flashcards e 10 questões sobre a Revolução Francesa", "estudos", ["estudos"]),
    ("faça um simulado sobre os meus documentos de direito penal", "estudos", ["biblioteca", "estudos"]),
    ("qual o resumo do dia?", "resumo_dia", ["resumo_dia"]),
    ("resuma meus e-mails", "email", ["email"]),
    ("quais são os e-mails importantes", "email", ["email"]),
])
def test_reconhece_os_novos_agentes(frase, tipo, plano):
    tarefa = planejar(frase)
    assert (tarefa.tipo, tarefa.plano) == (tipo, plano)


@pytest.mark.parametrize("frase", ["envie um email para o João", "conectar gmail", "tenho umas questões sobre o contrato",
                                   "abra a planilha de gastos"])
def test_frases_parecidas_seguem_o_fluxo_normal(frase):
    assert planejar(frase) is None


def test_quantidades_e_tema_do_estudo():
    tarefa = planejar("crie 20 flashcards e 10 questões sobre a Revolução Francesa")
    assert tarefa.quantidades == {"flashcards": 20, "questoes": 10}
    assert planejar("faça um simulado sobre os meus documentos de direito penal").tema == "direito penal"


def test_dados_pessoais_ficam_no_computador():
    assert planejar("qual o resumo do dia?").local_apenas
    assert planejar("resuma meus e-mails").local_apenas
    assert not planejar("crie uma planilha de controle de gastos").local_apenas


# ------------------------------------------------------------ planilha
def test_planilha_com_total_formulas_e_grafico(tmp_path):
    tarefa = executar(orquestrador(tmp_path), "crie uma planilha de controle de gastos com gráfico")

    assert tarefa.status == "concluída" and tarefa.avaliacao.cumprido
    livro = load_workbook(tarefa.arquivos[0])
    aba = livro["Gastos"]
    assert aba["B3"].value == 642.35                      # "R$ 642,35" virou número
    assert aba["B5"].value == "=SUM(B2:B4)"                # linha de total com fórmula
    assert len(aba._charts) == 1 and "Sobre" in livro.sheetnames
    assert "planilha 'Controle de Gastos' está pronta em Excel" in tarefa.mensagem_final


# ------------------------------------------------------------ estudos
def test_material_de_estudo_com_pdf_e_flashcards_para_o_anki(tmp_path):
    tarefa = executar(orquestrador(tmp_path, LLMNovos(cartoes=20, questoes=10)),
                      "crie 20 flashcards e 10 questões sobre a Revolução Francesa")

    assert tarefa.status == "concluída" and tarefa.avaliacao.cumprido
    pdf, anki = tarefa.arquivos
    assert pdf.endswith(".pdf") and anki.endswith("_flashcards_anki.txt")
    linhas = open(anki, encoding="utf-8").read().splitlines()
    assert linhas[0] == "#separator:tab" and len(linhas) == 22  # 2 de cabeçalho + 20 cartões
    assert "20 flashcards e 10 questões" in tarefa.mensagem_final


def test_poucos_flashcards_reprovam_no_critico(tmp_path):
    tarefa = executar(orquestrador(tmp_path, LLMNovos(cartoes=5, questoes=10)),
                      "crie 20 flashcards e 10 questões sobre a Revolução Francesa")
    assert any("5 flashcards (pedido: 20)" in d and not ok for d, ok in tarefa.avaliacoes[0].verificacoes)


def test_questao_valida_normaliza_letras_e_gabarito():
    questao = _questao_valida({"enunciado": "?", "alternativas": ["a) x", "B) y", "c. z", "w"], "correta": 2})
    assert questao["alternativas"] == ["x", "y", "z", "w"] and questao["correta"] == "C"
    assert _questao_valida({"enunciado": "?", "alternativas": ["x", "y"], "correta": "A"}) is None


# ------------------------------------------------------------ resumo do dia
def dados_do_dia():
    return {"data": "terça-feira, 13/10/2026, 08:00", "agenda": "10:00 Reunião com a equipe", "clima": None,
            "emails": [], "noticias": {"direito": [("STF julga revisão de contratos", "Tue, 13 Oct 2026")], "jogos": []}}


def test_resumo_do_dia_fala_e_salva_pdf_sem_mandar_a_agenda_para_a_nuvem(tmp_path):
    local, nuvem = LLMNovos(), LLMNovos()
    orq = orquestrador(tmp_path, local, llm_nuvem=LLMComReserva([], nuvem))
    orq.agentes["resumo_dia"] = AgenteResumoDia(coletar=dados_do_dia, llm_publico=nuvem, llm=local)

    tarefa = executar(orq, "qual o resumo do dia?")

    assert tarefa.status == "concluída"
    falado = tarefa.contexto["resposta"]
    assert "10:00 Reunião com a equipe" in falado and "Nenhum e-mail novo." in falado
    assert "Em direito: O STF julgou um caso importante." in falado   # resumo da nuvem
    assert tarefa.arquivos[0].endswith(".pdf")
    # A nuvem só recebeu as manchetes; a agenda nunca saiu do computador.
    assert nuvem.prompts and all("Reunião com a equipe" not in prompt for _, prompt in nuvem.prompts)


def test_resumo_falado_nao_menciona_o_que_nao_esta_configurado():
    from datetime import datetime

    dados = {"agenda": None, "clima": None, "emails": None,
             "noticias": {"jogos": [("Novo console é anunciado", "hoje")], "filmes": []}}
    falado = AgenteResumoDia.montar_falado(dados, {}, agora=datetime(2026, 10, 13, 8, 30))
    assert falado == "Bom dia! Hoje é terça-feira, 13 de outubro. Nas notícias. Em jogos, destaque para: Novo console é anunciado."
    assert "agenda" not in falado.lower() and "e-mail" not in falado.lower() and "clima" not in falado.lower()


# ------------------------------------------------------------ e-mail
EMAILS = [
    {"id": "1", "de": "cliente@empresa.com", "assunto": "Orçamento", "data": "hoje", "trecho": "Pode me mandar o orçamento?", "importante": True},
    {"id": "2", "de": "news@site.com", "assunto": "Newsletter", "data": "hoje", "trecho": "Novidades da semana", "importante": False},
]


def test_email_resume_e_rascunha_sem_enviar(tmp_path):
    orq = orquestrador(tmp_path)
    orq.agentes["email"] = AgenteEmail(listar=lambda consulta, limite: EMAILS, llm=orq.agentes["email"].llm)

    tarefa = executar(orq, "resuma meus e-mails")

    assert tarefa.status == "concluída" and tarefa.avaliacao.cumprido
    assert tarefa.arquivos[0].endswith(".docx")
    assert "1 rascunho(s) de resposta" in tarefa.mensagem_final and "Nada foi enviado" in tarefa.mensagem_final


def test_email_sem_gmail_vinculado_explica_o_que_fazer(tmp_path):
    def sem_gmail(consulta, limite):
        raise GmailNaoConectado("O Gmail não está vinculado. Diga 'conectar gmail'.")

    orq = orquestrador(tmp_path)
    orq.agentes["email"] = AgenteEmail(listar=sem_gmail, llm=orq.agentes["email"].llm)

    tarefa = executar(orq, "resuma meus e-mails")

    assert tarefa.status == "não realizada"
    assert "conectar gmail" in tarefa.mensagem_final
    assert not tarefa.relatorio


def test_conectar_gmail_sem_credenciais_explica_o_que_falta(monkeypatch, tmp_path):
    import google_gmail

    monkeypatch.setattr(google_gmail, "GOOGLE_CREDENTIALS_FILE", str(tmp_path / "nao_existe.json"))
    monkeypatch.setattr(google_gmail, "GOOGLE_CLIENT_ID", "")
    monkeypatch.setattr(google_gmail, "GOOGLE_CLIENT_SECRET", "")
    with pytest.raises(FileNotFoundError, match="Credenciais do Google não encontradas"):
        google_gmail._configuracao_cliente()
    assert google_gmail.conectar_gmail() == google_gmail.INSTRUCOES_VINCULO


def test_sem_emails_novos_avisa_sem_gerar_arquivo(tmp_path):
    orq = orquestrador(tmp_path)
    orq.agentes["email"] = AgenteEmail(listar=lambda consulta, limite: [], llm=orq.agentes["email"].llm)

    tarefa = executar(orq, "resuma meus e-mails")

    assert tarefa.status == "concluída" and not tarefa.arquivos
    assert tarefa.mensagem_final == "Você não tem e-mails novos para resumir agora."
