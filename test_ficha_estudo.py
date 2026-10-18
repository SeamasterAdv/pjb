"""Testes da ficha tática de estudo: formato, destaques, lei seca, teste de estresse, PDF/Word e o agente em etapas."""
import json
import unicodedata

import pymupdf
import pytest

from agente_leis import TextoDasLeis
from agentes import ficha_estudo
from agentes.base import Tarefa
from agentes.especialistas import TOPICOS_POR_CHAMADA, AgenteEstudos, _casar_topicos
from agentes.ficha_estudo import (aplicar_lacunas, conferir_lei_seca, destacar, gerar_docx_ficha, gerar_pdf_ficha,
                                  normalizar_ficha, quadro_de_prazos)

PAGINA_CF = """<html><body>
<p>Art. 37. A administração pública direta e indireta obedecerá aos princípios de legalidade.</p>
<p>§ 6º As pessoas jurídicas de direito público responderão pelos danos que seus agentes, nessa qualidade,
causarem a terceiros, assegurado o direito de regresso nos casos de dolo ou culpa.</p>
<p>Art. 38. Ao servidor público em exercício de mandato eletivo aplicam-se as disposições.</p>
</body></html>""".encode("cp1252")


def topico(titulo="Responsabilidade objetiva", **campos):
    base = {
        "titulo": titulo, "incidencia": "alta", "como_cai": "Troca 'objetiva' por 'subjetiva'.",
        "explicacao": ["O Estado responde sem culpa.", "Isso protege o cidadão."],
        "natureza": "Responsabilidade objetiva.", "requisitos": ["Conduta", "Dano", "Nexo"],
        "regra": "O Estado responde objetivamente.", "excecoes": ["Omissão genérica: subjetiva, salvo dever específico."],
        "prazos_e_efeitos": [{"item": "Prescrição", "valor": "5 anos", "fundamento": "Decreto 20.910/32"}],
        "lei_seca": [{"norma": "CF", "artigo": "37", "referencia": "art. 37, § 6º, CF", "palavra_chave": "nessa qualidade",
                      "pegadinha": "Troca por 'em qualquer situação'."}],
        "jurisprudencia": [{"referencia": "Tema 940 do STF", "tese": "Ação contra o Estado, não contra o agente.", "conferir": False}],
        "palavras_de_conflito": ["nessa qualidade"], "pegadinhas": ["Não confundir com subjetiva."],
        "exemplo": "Viatura bate num carro.", "pontos_chave": ["Ação: objetiva."], "fonte": "aula.pdf",
    }
    base.update(campos)
    return base


# ------------------------------------------------------------ conteúdo
def test_normaliza_o_formato_novo_e_o_antigo():
    ficha = normalizar_ficha({"titulo": "T", "area": "Direito", "resumo": "Visão.", "topicos": [
        topico(),
        {"titulo": "Antigo", "paragrafos": ["Um.", "Dois."], "atencao": "Cuidado.", "pontos_chave": ["P"]},
        {"sem": "título"},
    ]})
    assert ficha["area"] == "direito" and ficha["visao_geral"] == "Visão." == ficha["resumo"]
    novo, antigo = ficha["topicos"]
    assert novo["regra"] and novo["lei_seca"][0]["referencia"] == "art. 37, § 6º, CF"
    assert antigo["explicacao"] == ["Um.", "Dois."] and antigo["pegadinhas"] == ["Cuidado."] and antigo["incidencia"] == "media"
    assert normalizar_ficha({"area": "história"})["area"] == "geral"


def test_quadro_geral_de_prazos():
    ficha = normalizar_ficha({"topicos": [topico(), topico("Outro", prazos_e_efeitos=[{"item": "Recurso", "valor": "15 dias"}])]})
    assert [(p["topico"], p["valor"]) for p in quadro_de_prazos(ficha)] == [("Responsabilidade objetiva", "5 anos"), ("Outro", "15 dias")]


def test_destaca_as_palavras_que_mudam_o_sentido():
    html = destacar("Salvo se houver dolo, é vedado, independentemente de culpa; a salvaguarda não.")
    assert '<span class="mt">Salvo se</span>' in html
    assert '<span class="mt">é vedado</span>' in html and '<span class="mt">independentemente</span>' in html
    assert "salvaguarda" in html and '<span class="mt">salvaguarda' not in html
    # Palavras do próprio tópico e sobreposição viram um destaque só (HTML equilibrado).
    sobreposto = destacar("salvo se houver dolo", ["salvo se houver"])
    assert sobreposto == '<span class="mt">salvo se houver</span> dolo'
    assert destacar("<b>&</b>") == "&lt;b&gt;&amp;&lt;/b&gt;"


def test_marca_texto_do_topico_so_para_palavras_restritivas():
    from agentes.ficha_estudo import qualificadoras

    assert qualificadoras(["risco integral", "excludentes", "independentemente de culpa", "somente", "culpa anônima",
                           "não se aplica"]) == ["independentemente de culpa", "somente", "não se aplica"]
    # Números, quóruns e prazos também são trocados pela banca.
    assert qualificadoras(["três quintos", "em dois turnos", "tendente a abolir", "controle difuso"]) == \
        ["três quintos", "em dois turnos", "tendente a abolir"]


def test_lacuna_com_topico_de_outro_nome_vai_pelo_conteudo():
    ficha = normalizar_ficha({"topicos": [topico("Controle difuso (concreto) e seus efeitos", regra="Qualquer juiz afasta a lei."),
                                          topico("Efeitos da decisão no controle concentrado e modulação")]})
    aplicar_lacunas(ficha, [
        {"topico": "Controle repressivo difuso e reserva de plenário (não consta da ficha)", "tese": "Reserva de plenário."},
        {"topico": "", "tese": "A modulação dos efeitos exige dois terços do STF."},
    ])
    assert ficha["topicos"][0]["teste_de_estresse"] == ["Reserva de plenário."]
    assert ficha["topicos"][1]["teste_de_estresse"] == ["A modulação dos efeitos exige dois terços do STF."]
    assert not ficha["lacunas_gerais"]


def test_decreto_nao_e_decreto_lei():
    from agente_leis import candidatos_planalto, norma_por_texto

    decreto = norma_por_texto("Decreto 20.910/32")
    assert (decreto.tipo, decreto.nome) == ("decreto", "o Decreto 20.910")
    assert candidatos_planalto("20910", "", "decreto")[0].endswith("/decreto/antigos/d20910.htm")
    assert norma_por_texto("decreto-lei 3.365/41").tipo == "decreto-lei"


def test_lacunas_do_teste_de_estresse_entram_no_topico_certo():
    ficha = normalizar_ficha({"topicos": [topico(), topico("Prescrição contra a Fazenda")]})
    aplicadas = aplicar_lacunas(ficha, [
        {"topico": "Prescricao contra a fazenda", "tese": "Interrompe-se uma única vez."},
        {"topico": "Assunto que não existe", "tese": "Tese solta."},
        {"topico": "Prescrição contra a Fazenda", "tese": "Interrompe-se uma única vez."},  # repetida
        "texto sem tópico",
    ])
    assert aplicadas == 3
    assert ficha["topicos"][1]["teste_de_estresse"] == ["Interrompe-se uma única vez."]
    assert ficha["lacunas_gerais"] == ["Tese solta.", "texto sem tópico"]


# ------------------------------------------------------------- lei seca
def test_confere_a_lei_seca_no_planalto():
    ficha = normalizar_ficha({"topicos": [topico(lei_seca=[
        {"norma": "CF", "artigo": "37", "referencia": "art. 37, § 6º, CF"},
        {"norma": "CF/88", "referencia": "art. 999 da CF"},
        {"norma": "Lei 9.999/1990", "artigo": "1"},
        {"referencia": "Regimento interno, art. 5"},
    ])]})
    textos = TextoDasLeis(baixar=lambda url: PAGINA_CF if "constituicao" in url else (_ for _ in ()).throw(OSError("404")))
    contagem = conferir_lei_seca(ficha, textos=textos)
    certo, errado, fora, sem_norma = ficha["topicos"][0]["lei_seca"]
    assert certo["conferido"] is True and "nessa qualidade" in certo["texto_oficial"]
    assert errado["conferido"] is False
    assert fora["conferido"] is None and sem_norma["conferido"] is None
    assert contagem == {"conferidos": 1, "nao_encontrados": 1, "sem_acesso": 1}


# ------------------------------------------------------------ arquivos
def test_pdf_colorido_com_todas_as_partes(tmp_path):
    ficha = normalizar_ficha({
        "titulo": "Responsabilidade Civil do Estado", "area": "direito", "visao_geral": "Visão geral.",
        "como_a_banca_cobra": ["Prazos e exceções."], "topicos": [topico(teste_de_estresse=["Concessionárias também."])],
        "flashcards": [{"frente": "Prazo?", "verso": "5 anos."}],
        "questoes": [{"enunciado": "Assinale.", "alternativas": ["a", "b", "c", "d"], "correta": "C", "explicacao": "Art. 37.",
                      "topico": "Responsabilidade objetiva"}],
        "glossario": [{"termo": "Nexo", "definicao": "Ligação."}], "fontes": ["aula.pdf"],
    })
    ficha["topicos"][0]["lei_seca"][0].update(conferido=True, texto_oficial="responderão pelos danos, nessa qualidade")
    caminho = gerar_pdf_ficha(ficha, str(tmp_path / "ficha.pdf"))
    with pymupdf.open(caminho) as documento:  # NFKC desfaz as ligaduras da fonte ("ﬁcha" -> "ficha")
        texto = unicodedata.normalize("NFKC", " ".join(pagina.get_text() for pagina in documento))
    for esperado in ("Ficha tática de estudo", "COMO A BANCA COBRA", "Mapa de incidência", "NATUREZA JURÍDICA E REQUISITOS",
                     "REGRA", "EXCEÇÕES", "PRAZOS E EFEITOS", "LEI SECA", "TEXTO CONFERIDO NO PLANALTO", "STF E STJ",
                     "ACRESCENTADO NO TESTE DE ESTRESSE", "Teste de estresse", "Gabarito comentado", "Flashcards",
                     "Glossário", "Como estudar com esta ficha", "Página 1 de"):
        assert esperado in texto, esperado


def test_pdf_fora_do_direito_usa_rotulos_gerais(tmp_path):
    ficha = normalizar_ficha({"titulo": "Revolução Francesa", "area": "geral", "topicos": [topico("Causas")]})
    with pymupdf.open(gerar_pdf_ficha(ficha, str(tmp_path / "f.pdf"))) as documento:
        texto = documento[0].get_text() + (documento[1].get_text() if documento.page_count > 1 else "")
    assert "CONCEITO E ELEMENTOS" in texto and "NATUREZA JURÍDICA" not in texto


def test_word_com_caixas_e_marca_texto(tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    ficha = normalizar_ficha({"titulo": "Ficha", "area": "direito", "topicos": [topico()],
                              "flashcards": [{"frente": "P", "verso": "R"}]})
    documento = Document(gerar_docx_ficha(ficha, str(tmp_path / "ficha.docx")))
    destacados = [run.text for tabela in documento.tables for linha in tabela.rows for celula in linha.cells
                  for paragrafo in celula.paragraphs for run in paragrafo.runs if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "salvo" in destacados and "nessa qualidade" in destacados
    rotulos = [celula.paragraphs[0].text for tabela in documento.tables for celula in tabela.rows[0].cells]
    assert {"REGRA", "EXCEÇÕES", "LEI SECA", "STF E STJ"} <= set(rotulos)


# -------------------------------------------------------- o agente
def test_casa_as_fichas_com_o_plano():
    pedidos = [{"titulo": "Controle difuso"}, {"titulo": "ADI"}]
    recebidos = [{"titulo": "ADI genérica e interventiva"}, {"titulo": "Extra que ninguém pediu"}, {"titulo": "controle DIFUSO"}]
    casados = _casar_topicos(pedidos, recebidos)
    assert [t["titulo"] for t in casados] == ["Controle difuso", "ADI"]


class IAEtapas:
    """Responde cada etapa do agente de estudos com o formato novo, e guarda o que recebeu."""

    def __init__(self, topicos=7, lacunas=True):
        self.topicos, self.lacunas = topicos, lacunas
        self.etapas = []

    def chat(self, prompt, system_prompt=None, model=None, **kwargs):
        if "ANÁLISE CRÍTICA do assunto" in system_prompt:
            self.etapas.append(("analise", prompt))
            return json.dumps({"titulo": "Responsabilidade Civil do Estado", "area": "direito", "visao_geral": "Visão.",
                               "como_a_banca_cobra": ["Prazos"], "topicos": [
                                   {"titulo": f"Tópico {i}", "incidencia": "alta" if i == 1 else "baixa", "foco": "f"}
                                   for i in range(1, self.topicos + 1)]})
        if "FICHAS TÁTICAS" in system_prompt:
            self.etapas.append(("fichas", prompt))
            pedidos = [linha[2:].split(" (incidência")[0] for linha in prompt.splitlines() if linha.startswith("- Tópico")]
            return json.dumps({"topicos": [topico(t) for t in reversed(pedidos)] + [topico("Intruso")],
                               "fora_das_fontes": ["Histórico"]})
        if "revisão espaçada" in system_prompt:
            self.etapas.append(("fechamento", prompt))
            return json.dumps({"flashcards": [{"frente": f"P{i}", "verso": f"R{i}"} for i in range(25)],
                               "glossario": [{"termo": "Nexo", "definicao": "Ligação."}]})
        if "TESTE DE ESTRESSE" in system_prompt:
            self.etapas.append(("estresse", prompt))
            return json.dumps({"questoes": [{"enunciado": f"Q{i}?", "alternativas": ["a", "b", "c", "d"], "correta": "B",
                                             "explicacao": "Porque sim.", "topico": "Tópico 1", "estava_na_ficha": i > 0}
                                            for i in range(10)],
                               "lacunas": [{"topico": "Tópico 2", "tese": "Exceção do parágrafo único."}] if self.lacunas else []})
        raise AssertionError(f"etapa inesperada: {system_prompt[:80]}")


def test_agente_faz_as_quatro_etapas_e_gera_a_ficha(tmp_path):
    ia, conferidas = IAEtapas(topicos=7), []

    def conferir(ficha):
        conferidas.append(len(ficha["topicos"]))
        ficha["topicos"][0]["lei_seca"][0].update(conferido=True, texto_oficial="nessa qualidade")
        return {"conferidos": 7, "nao_encontrados": 1}

    agente = AgenteEstudos(llm=ia, conferir_lei=conferir)
    andamento = []
    agente.ao_progredir = lambda tarefa: andamento.append(tarefa.etapa_atual)
    tarefa = Tarefa(pedido="faz um resumo da matéria de responsabilidade civil", plano=["biblioteca", "estudos"],
                    formatos=["pdf", "docx"], tema="responsabilidade civil", materia=True, fontes_no_claude=True,
                    pasta=str(tmp_path / "saida"))
    tarefa.contexto["biblioteca"] = "[Fonte: aula.pdf | arquivo inteiro]\nO Estado responde objetivamente."

    resumo = agente.executar(tarefa)

    etapas = [nome for nome, _ in ia.etapas]
    lotes = -(-7 // TOPICOS_POR_CHAMADA)
    assert etapas[0] == "analise" and etapas.count("fichas") == lotes and etapas[-2:] == ["fechamento", "estresse"]
    analise = ia.etapas[0][1]
    assert analise.index("O Estado responde objetivamente") < analise.index("Pedido do usuário")  # fontes antes do pedido
    ficha = tarefa.contexto["conteudo"]
    assert [t["titulo"] for t in ficha["topicos"]] == [f"Tópico {i}" for i in range(1, 8)]  # ordem do plano, sem o intruso
    assert ficha["topicos"][1]["teste_de_estresse"] == ["Exceção do parágrafo único."]
    assert len(ficha["flashcards"]) == 20 and len(ficha["questoes"]) == 10  # o padrão da matéria: 20 e 10
    assert ficha["fora_das_fontes"] == ["Histórico"] and conferidas == [7]
    pdf, docx, anki = tarefa.arquivos
    assert pdf.endswith(".pdf") and docx.endswith(".docx") and anki.endswith("_flashcards_anki.txt")
    guardada = json.load(open(tmp_path / "saida" / "ficha.json", encoding="utf-8"))  # para refazer o visual depois
    assert len(guardada["topicos"]) == 7 and guardada["questoes"]
    assert "7 tópicos, 20 flashcards e 10 questões" in resumo
    assert "Conferiu 7 artigo(s) da lei seca" in resumo and "1 artigo(s) citado(s) não existem" in resumo
    assert "achou 1 lacuna(s)" in resumo
    assert andamento[0] == "analisando como o assunto cai na prova" and "conferindo a lei seca no Planalto" in andamento
    assert "escrevendo as fichas (7 de 7 prontas)" in andamento


def test_ficha_de_assunto_sem_materia_tem_menos_topicos(tmp_path):
    ia = IAEtapas(topicos=12, lacunas=False)
    agente = AgenteEstudos(llm=ia, conferir_lei=None)
    tarefa = Tarefa(pedido="quero estudar responsabilidade civil", plano=["estudos"], tema="responsabilidade civil",
                    pasta=str(tmp_path / "saida"))
    resumo = agente.executar(tarefa)
    assert len(tarefa.contexto["conteudo"]["topicos"]) == 8  # no máximo 8 fora do resumo da matéria
    assert "12 flashcards e 6 questões" in resumo and "lacuna" not in resumo
    assert "Sem fontes externas" in ia.etapas[0][1]


def test_sem_topicos_na_analise_falha(tmp_path):
    class Vazia(IAEtapas):
        def chat(self, prompt, system_prompt=None, model=None, **kwargs):
            return json.dumps({"titulo": "X", "topicos": []})

    tarefa = Tarefa(pedido="estudar X", plano=["estudos"], tema="X", pasta=str(tmp_path))
    with pytest.raises(ValueError, match="sem tópicos"):
        AgenteEstudos(llm=Vazia(), conferir_lei=None).executar(tarefa)
