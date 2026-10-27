"""Peças processuais: o reconhecimento do pedido, a conferência dos artigos, o Word e o PDF e o caminho pelos agentes
(IA falsa, sem nuvem nem Planalto de verdade)."""
import json
import os
import re
import threading

import pytest

from agentes import Orquestrador, planejar
from agentes.base import Tarefa
from agentes.pecas import (TIPOS_DE_PECA, AgentePeca, citacoes_do_texto, gerar_docx_peca, gerar_pdf_peca,
                           marcar_inexistentes, pedido_de_peca)
from llm_nuvem import LLMComReserva
from test_agentes import LLMFalso

ERRADOS = {"9999", "8888"}  # artigos que "não existem" nos testes


@pytest.mark.parametrize("frase, tipo", [
    ("faz uma petição inicial de indenização por danos morais contra a Viagem Fácil", "peticao_inicial"),
    ("redige um agravo de instrumento contra a decisão que negou a gratuidade", "agravo_de_instrumento"),
    ("USA o Claude pra produzir uma peça correspondente", ""),
    ("faz a peça cabível desse caso", ""),
    ("faz a apelação do réu condenado por furto", "apelacao_criminal"),
    ("faz uma apelação contra a sentença que julgou improcedente", "apelacao"),
    ("faz embargos de declaração da sentença omissa", "embargos_de_declaracao"),
    ("impetra um habeas corpus para o meu cliente preso", "habeas_corpus"),
    ("quero uma reclamação trabalhista de horas extras", "reclamacao_trabalhista"),
    ("Beatriz faz uma contestação trabalhista", "contestacao_trabalhista"),
    ("elabora um mandado de segurança contra o secretário de saúde", "mandado_de_seguranca"),
])
def test_reconhece_pedidos_de_peca(frase, tipo):
    assert pedido_de_peca(frase) == tipo


@pytest.mark.parametrize("frase", [
    "Produz um ensino de peça de agravo de instrumento pra eu estudar",
    "o que é agravo de instrumento?",
    "faz uma ficha de apelação",
    "faz um resumo sobre contestação",
    "me explica a petição inicial",
    "monte uma apresentação com base nos meus documentos sobre petição inicial",
    "faz um PDF a respeito do recurso especial",
    "faz um PDF sobre a reforma tributária",
])
def test_estudar_a_peca_nao_e_redigir(frase):
    assert pedido_de_peca(frase) is None


def test_planejar_monta_a_tarefa_da_peca():
    tarefa = planejar("Beatriz, redige um agravo de instrumento contra a decisão que negou a liminar")
    assert tarefa.tipo == "peca" and tarefa.plano == ["peca"]
    assert tarefa.contexto["tipo_peca"] == "agravo_de_instrumento" and tarefa.tema == "Agravo de instrumento"
    pelo_claude = planejar("usa o Claude pra produzir uma peça correspondente")
    assert pelo_claude.ia_preferida == "claude" and pelo_claude.contexto["tipo_peca"] == ""
    # Estudo da peça continua com os agentes de estudo.
    assert planejar("faz um resumo para a prova sobre agravo de instrumento").tipo == "estudos"


def test_cada_peca_tem_estrutura_fundamentos_e_cuidados():
    for chave, tipo in TIPOS_DE_PECA.items():
        assert tipo["nome"] and tipo["area"] and len(tipo["estrutura"]) >= 5, chave
        assert tipo["fundamentos"] and all(f.startswith("art. ") for f in tipo["fundamentos"]), chave
        assert any("pedido" in parte.lower() for parte in tipo["estrutura"]), chave


def test_citacoes_e_marcas_dos_artigos_que_nao_existem():
    texto = ("com fundamento no art. 1.015, I, do CPC e nos arts. 186 e 927 do Código Civil, além do artigo 5º, X, da "
             "Constituição Federal e do art. 9999 do CPC.")
    citacoes = {(c["norma"], c["artigo"]) for c in citacoes_do_texto(texto)}
    assert citacoes == {("cpc", "1.015"), ("cc", "186"), ("cc", "927"), ("cf", "5"), ("cpc", "9999")}
    peca = {"documentos": [{"preambulo": texto, "secoes": []}]}
    marcar_inexistentes(peca, [{"norma": "cpc", "artigo": "9999"}])
    assert peca["documentos"][0]["preambulo"].endswith("do art. 9999 [CONFERIR O NÚMERO] do CPC.")
    assert "1.015 [CONFERIR" not in peca["documentos"][0]["preambulo"]


# ------------------------------------------------------------------ o agente
ANALISE = {
    "tipo": "peticao_inicial", "por_que": "O voo foi cancelado e ainda não há processo: cabe a ação de indenização.",
    "juizo": "Juízo da Vara Cível da Comarca de [COMARCA]",
    "partes": [{"papel": "autora", "nome": "Kamila", "qualificacao": ""}, {"papel": "ré", "nome": "Viagem Fácil Turismo"}],
    "fatos": ["Voo cancelado sem aviso."],
    "teses": [{"titulo": "Responsabilidade objetiva", "argumento": "Fato do serviço.",
               "fundamentos": [{"norma": "CDC", "artigo": "14", "referencia": "art. 14, CDC"},
                               {"norma": "CPC", "artigo": "9999", "referencia": "art. 9999, CPC"}]}],
    "jurisprudencia": [], "pedidos": ["indenização por danos morais"], "provas": ["documentos"],
    "valor_da_causa": "R$ 10.000,00", "prazo": "", "faltando": ["CPF da autora", "endereço da ré"],
}
REDACAO = {"documentos": [{
    "parte": "peca", "enderecamento": "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO DA ___ VARA CÍVEL DA COMARCA DE [COMARCA]",
    "processo": "", "titulo": "AÇÃO DE INDENIZAÇÃO POR DANOS MORAIS",
    "preambulo": ("KAMILA, [NACIONALIDADE], portadora do CPF [CPF], por seu advogado, vem propor a presente AÇÃO DE INDENIZAÇÃO "
                  "POR DANOS MORAIS, com fundamento no art. 14 do CDC e no art. 319 do CPC, em face de VIAGEM FÁCIL TURISMO."),
    "secoes": [
        {"titulo": "I – DOS FATOS", "paragrafos": ["A autora comprou passagens e o voo foi cancelado sem aviso prévio."]},
        {"titulo": "II – DO DIREITO", "paragrafos": ["A responsabilidade é objetiva (art. 14 do CDC), e o art. 8888 do CPC também."]},
        {"titulo": "III – DOS PEDIDOS", "itens": ["a) a citação da ré", "b) a condenação em danos morais", "a procedência"]},
    ],
    "fechamento": ["Nestes termos,", "pede deferimento.", "Local…, Data…", "Advogado…, OAB/UF nº…"],
}], "valor_da_causa": "R$ 10.000,00"}


class IAPeca:
    """Imita o Claude (ou o Gemini) nas peças e no crítico, e guarda o que recebeu."""

    def __init__(self, nome="Claude"):
        self.nome, self.modelo, self.recebido = nome, f"{nome.lower()}-teste", []

    def disponivel(self):
        return True

    def chat(self, prompt, system_prompt=None, model=None, **kwargs):
        self.recebido.append((system_prompt, prompt))
        if "Agente de Peças" in system_prompt and "ANALISA" in system_prompt:
            return json.dumps(ANALISE, ensure_ascii=False)
        if "Agente de Peças" in system_prompt:
            return json.dumps(REDACAO, ensure_ascii=False)
        if "Agente Crítico" in system_prompt:
            return json.dumps({"nota": 9, "pontos_fortes": ["Peça certa"], "problemas": [], "correcao": ""})
        return "{}"


def conferir_falso(ficha):
    contagem = {"conferidos": 0, "nao_encontrados": 0}
    for topico in ficha["topicos"]:
        for item in topico["lei_seca"]:
            numero = re.sub(r"\D", "", str(item.get("artigo") or ""))
            if numero in ERRADOS:
                item["conferido"] = False
                contagem["nao_encontrados"] += 1
            else:
                item.update(conferido=True, texto_oficial=f"Texto oficial do artigo {item.get('artigo')}.")
                contagem["conferidos"] += 1
    return contagem


def _tarefa(tmp_path, **contexto):
    tarefa = Tarefa(pedido="faz uma petição inicial de indenização contra a Viagem Fácil", plano=["peca"], formatos=["docx", "pdf"],
                    tema="Petição inicial")
    tarefa.pasta = str(tmp_path / "peca")
    tarefa.contexto.update(tipo_peca="peticao_inicial", fatos=tarefa.pedido, **contexto)
    return tarefa


def test_agente_confere_os_artigos_antes_e_depois_de_redigir(tmp_path):
    ia = IAPeca()
    agente = AgentePeca(llm=ia, conferir_lei=conferir_falso, conferir_sumulas=None)
    tarefa = _tarefa(tmp_path)
    resumo = agente.executar(tarefa)

    sistema_analise, _ = ia.recebido[0]
    assert "MODO PROVA DA OAB" in sistema_analise and "peticao_inicial:" in sistema_analise
    _, prompt_redacao = ia.recebido[1]
    # O texto oficial dos conferidos vai junto; o artigo que não existe é proibido na redação.
    assert "Texto oficial do artigo 14." in prompt_redacao and "Texto oficial do artigo 319." in prompt_redacao
    assert "NÃO CITE" in prompt_redacao and "art. 9999, CPC" in prompt_redacao
    assert "Petição inicial (procedimento comum)" in prompt_redacao and "Dos pedidos" in prompt_redacao

    peca = tarefa.contexto["conteudo"]
    direito = peca["documentos"][0]["secoes"][1]["paragrafos"][0]
    assert "art. 8888 [CONFERIR O NÚMERO] do CPC" in direito  # citado só na redação, e não existe
    assert peca["conferencia"]["nao_encontrados"] == 2 and peca["conferencia"]["conferidos"] >= 3
    assert [os.path.splitext(a)[1] for a in tarefa.arquivos] == [".docx", ".pdf"]
    assert "Conferiu" in resumo and "não existem" in resumo
    assert json.load(open(os.path.join(tarefa.pasta, "peca.json"), encoding="utf-8"))["tipo"] == "peticao_inicial"


def test_modo_caso_real_leva_o_advogado_e_a_cidade(tmp_path):
    ia = IAPeca()
    prefs = {"pecas": {"modo": "real", "advogado": "Dra. Ana Souza", "oab": "DF 12.345", "cidade": "Brasília", "instrucoes": "pedidos curtos"}}
    AgentePeca(llm=ia, conferir_lei=conferir_falso, conferir_sumulas=None).executar(_tarefa(tmp_path, preferencias=prefs))
    sistema, _ = ia.recebido[1]
    assert "MODO CASO REAL" in sistema and "Dra. Ana Souza" in sistema and "Brasília" in sistema
    assert "pedidos curtos" in sistema and "MODO PROVA DA OAB" not in sistema


def test_word_e_pdf_no_formato_forense(tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    peca = {**REDACAO, "nome": "Petição inicial", "tipo": "peticao_inicial"}
    caminho_docx = gerar_docx_peca(peca, str(tmp_path / "peca.docx"))
    documento = Document(caminho_docx)
    assert documento.styles["Normal"].font.name == "Times New Roman"
    textos = [p.text for p in documento.paragraphs]
    assert "III – DOS PEDIDOS" in textos
    assert "a) a citação da ré" in textos and "c) a procedência" in textos  # a numeração é refeita
    assert "Dá-se à causa o valor de R$ 10.000,00." in textos
    destacados = [r.text for p in documento.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert "[CPF]" in destacados and "[COMARCA]" in destacados

    import pymupdf

    caminho_pdf = gerar_pdf_peca(peca, str(tmp_path / "peca.pdf"))
    with pymupdf.open(caminho_pdf) as pdf:
        texto = "".join(pagina.get_text() for pagina in pdf)
    assert "DOS PEDIDOS" in texto and "AÇÃO DE INDENIZAÇÃO" in texto


# ------------------------------------------------------------------ pelos agentes
def _orquestrador(tmp_path, ia, **extras):
    local = LLMFalso()
    return Orquestrador(str(tmp_path), llm=local, abrir_ao_concluir=False, llm_nuvem=LLMComReserva([ia], local),
                        preferencias=lambda: {"pecas": {"modo": "oab", "formato": "ambos"}}, **extras)


def _esperar(orquestrador, enviar):
    pronto = threading.Event()
    tarefa = enviar(lambda t: pronto.set())
    assert tarefa is not None
    assert pronto.wait(30)
    return tarefa


def test_peca_pelo_pedido_vai_para_a_nuvem_e_sai_em_word_e_pdf(tmp_path):
    ia = IAPeca("Claude")
    orq = _orquestrador(tmp_path, ia)
    tarefa = _esperar(orq, lambda fim: orq.enviar("faz uma petição inicial de indenização contra a Viagem Fácil", ao_concluir=fim))
    assert tarefa.status == "concluída", tarefa.erro
    assert tarefa.formatos == ["docx", "pdf"] and len(tarefa.arquivos) == 2
    assert any("Agente de Peças" in s for s, _ in ia.recebido) and any("Agente Crítico" in s for s, _ in ia.recebido)
    assert tarefa.contexto["modelos"] == ["Claude (claude-teste)"]
    assert "está pronta em Word e PDF" in tarefa.mensagem_final
    assert "CPF da autora" in tarefa.mensagem_final  # o que falta completar


def test_mensagem_de_inicio_da_peca(tmp_path):
    orq = _orquestrador(tmp_path, IAPeca())
    tarefa = planejar("redige um agravo de instrumento contra a decisão")
    orq._preferencias_da_peca(tarefa)
    mensagem = orq._inicio_peca(tarefa)
    assert "agravo de instrumento" in mensagem and "OAB" in mensagem
    assert "Os fatos vão para o Claude (o Gemini de reserva)" in mensagem


def test_peca_de_arquivo_le_o_caso_inteiro(tmp_path):
    caso = tmp_path / "caso_kamila.txt"
    caso.write_text("Enunciado: Kamila comprou um pacote da Viagem Fácil Turismo. " * 200, encoding="utf-8")
    ia = IAPeca()
    lidos = []
    from agentes.especialistas import AgenteBiblioteca

    biblioteca = AgenteBiblioteca(buscar=lambda consulta: "", ler_arquivo=lambda c, limite: lidos.append(limite) or open(c, encoding="utf-8").read())
    orq = _orquestrador(tmp_path, ia)
    orq.agentes["biblioteca"] = biblioteca
    tarefa = _esperar(orq, lambda fim: orq.enviar_peca("", "faz a peça desse caso", [str(caso)], ao_concluir=fim,
                                                         pedido="faz a peça desse caso"))
    assert tarefa.status == "concluída", tarefa.erro
    assert tarefa.plano == ["biblioteca", "peca"] and lidos[0] > 4000  # o caso inteiro, não só 4 mil letras
    _, prompt = ia.recebido[0]
    assert "ARQUIVOS DO CASO" in prompt and prompt.count("Kamila comprou") > 50
    assert "O pedido não diz a peça: escolha a CABÍVEL." in prompt
    assert "Por que essa peça:" in tarefa.mensagem_final


def test_so_no_pc_nao_manda_os_fatos_para_a_nuvem(tmp_path):
    ia = IAPeca()
    orq = _orquestrador(tmp_path, ia)
    orq.agentes["peca"] = AgentePeca(llm=IAPeca("Ollama"), conferir_lei=None, conferir_sumulas=None)
    tarefa = _esperar(orq, lambda fim: orq.enviar_peca("contestacao", "o autor diz que eu devo R$ 5 mil", ia="local", ao_concluir=fim))
    assert tarefa.status == "concluída", tarefa.erro
    assert not any("Agente de Peças" in s for s, _ in ia.recebido)


def test_faz_a_peca_desse_caso_logo_depois_do_arquivo_chegar(tmp_path):
    from fontes import AgenteFontes, Registro

    pedidos = []
    agente = AgenteFontes(Registro(str(tmp_path / "registro.json")), str(tmp_path / "recebidos"), str(tmp_path / "arquivados"),
                          peca=lambda caminhos, pedido, canal, ia: pedidos.append((caminhos, pedido, ia)) or "Certo, vou redigir.")
    foto = tmp_path / "caso.jpg"
    foto.write_bytes(b"\xff\xd8\xff" + b"0" * 100)
    agente.registro.registrar(str(foto), "caso.jpg", "whatsapp")
    assert agente.responder("USA o Claude pra produzir uma peça correspondente") == "Certo, vou redigir."
    caminhos, pedido, ia = pedidos[-1]
    assert caminhos == [str(foto)] and ia == "claude"
    # Sem arquivo citado nem recente, a peça é dos agentes de sempre.
    assert AgenteFontes(Registro(str(tmp_path / "vazio.json")), str(tmp_path), str(tmp_path),
                        peca=lambda *a: "não").responder("faz uma petição inicial de cobrança") is None


def test_recurso_sem_valor_da_causa_e_processo_sem_repetir(tmp_path):
    # No teste com o Claude de verdade (out/2026), o agravo saiu com "Dá-se à causa o valor de Não se aplica…" e
    # "Processo nº Processo de origem nº…".
    from docx import Document

    documento = {**REDACAO["documentos"][0], "processo": "Processo de origem nº [Nº DO PROCESSO DE ORIGEM]"}
    peca = {"tipo": "agravo_de_instrumento", "documentos": [documento],
            "valor_da_causa": "Não se aplica ao agravo de instrumento. Valor da causa na origem: [VALOR DA CAUSA]."}
    textos = [p.text for p in Document(gerar_docx_peca(peca, str(tmp_path / "agravo.docx"))).paragraphs]
    assert "Processo de origem nº [Nº DO PROCESSO DE ORIGEM]" in textos
    assert not any(texto.startswith("Dá-se à causa") for texto in textos)
    inicial = {**peca, "tipo": "peticao_inicial", "valor_da_causa": "[VALOR DA CAUSA]",
               "documentos": [{**documento, "processo": "0701234-55.2026.8.07.0001"}]}
    textos = [p.text for p in Document(gerar_docx_peca(inicial, str(tmp_path / "inicial.docx"))).paragraphs]
    assert "Processo nº 0701234-55.2026.8.07.0001" in textos and "Dá-se à causa o valor de [VALOR DA CAUSA]." in textos

