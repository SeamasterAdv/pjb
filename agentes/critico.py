"""Agente crítico: confere se o que foi entregue cumpre o que foi pedido."""
import json
import logging
import os
import re

from agentes.base import Agente, Avaliacao, Tarefa
# O crítico confere o texto contra o mesmo trecho das fontes que o redator recebeu; com menos,
# fatos verdadeiros do resto das fontes seriam apontados como inventados.
from agentes.especialistas import LIMITE_CONTEXTO_WEB, problemas_da_aula
from agentes.geradores import _lista

logger = logging.getLogger(__name__)

# Aprovação: nenhuma verificação objetiva falhou e o LLM deu pelo menos esta nota ao conteúdo.
NOTA_MINIMA = 7


def _contar_slides(caminho):
    from pptx import Presentation
    return len(Presentation(caminho).slides)


def _contar_paginas_pdf(caminho):
    import pymupdf
    with pymupdf.open(caminho) as documento:
        return documento.page_count


def _contar_palavras_docx(caminho):
    from docx import Document
    return sum(len(p.text.split()) for p in Document(caminho).paragraphs)


def _verificar_planilha(caminho: str, pedido: str) -> list:
    from openpyxl import load_workbook

    livro = load_workbook(caminho)
    abas = [livro[nome] for nome in livro.sheetnames if nome != "Sobre"]
    resultado = [(f"Planilha com {len(abas)} aba(s) de dados", bool(abas))]
    for aba in abas:
        resultado.append((f"Aba '{aba.title}' com {aba.max_row - 1} linha(s) de dados", aba.max_row >= 3))
    if re.search(r"gr[áa]fico", pedido, re.IGNORECASE):
        resultado.append(("Gráfico pedido foi incluído", any(aba._charts for aba in abas)))
    return resultado


def _verificar_estudos(tarefa: Tarefa, conteudo: dict) -> list:
    pedidos = {chave: tarefa.quantidade_estudo(chave) for chave in ("flashcards", "questoes")}
    resultado = []
    for chave, rotulo in (("flashcards", "flashcards"), ("questoes", "questões")):
        feitos = len(conteudo.get(chave) or [])
        # Aceita até 20% a menos: itens malformados são descartados pelo agente.
        resultado.append((f"{feitos} {rotulo} (pedido: {pedidos[chave]})", feitos >= max(1, int(pedidos[chave] * 0.8))))
    topicos = len(conteudo.get("topicos") or [])
    if tarefa.materia and tarefa.fontes_no_claude:
        # Resumo da matéria inteira, escrito pelo Claude: precisa cobrir mais que o básico.
        resultado.append((f"Matéria em {topicos} tópicos (mínimo 5)", topicos >= 5))
    else:
        resultado.append(("Material tem tópicos explicados", topicos >= 2))
    return resultado


def _verificar_aulas(tarefa: Tarefa, conteudo: dict) -> list:
    aulas = conteudo.get("aulas") or []
    pedidas = tarefa.quantidades.get("aulas")
    if pedidas:
        resultado = [(f"{len(aulas)} aulas no curso (pedido: {pedidas})", len(aulas) == pedidas)]
    else:
        resultado = [(f"{len(aulas)} aulas no curso (mínimo 3)", len(aulas) >= 3)]
    for aula in aulas:
        problemas = problemas_da_aula(aula)
        resultado.append((f"Aula {aula.get('numero')}: {'; '.join(problemas) or 'completa'}", not problemas))
    return resultado


def verificacoes_objetivas(tarefa: Tarefa) -> list:
    """Checagens que não dependem do LLM: arquivos, formatos e quantidades pedidas."""
    resultado = []
    conteudo = tarefa.contexto.get("conteudo") or {}
    if tarefa.tipo == "pesquisa" and not tarefa.formatos:
        resposta = tarefa.contexto.get("resposta", "")
        resultado.append(("A pesquisa produziu uma resposta", len(resposta.split()) >= 40))
        resultado.append(("A resposta se baseou em fontes da internet", bool(tarefa.contexto.get("web"))))
        return resultado
    if tarefa.tipo == "email" and not tarefa.contexto.get("emails_lidos"):
        return [("Não havia e-mails novos para resumir", True)]

    extensoes = {os.path.splitext(a)[1].lstrip(".").lower() for a in tarefa.arquivos}
    for formato in tarefa.formatos:
        resultado.append((f"Arquivo {formato.upper()} foi gerado", formato in extensoes))
    for arquivo in tarefa.arquivos:
        nome = os.path.basename(arquivo)
        existe = os.path.exists(arquivo) and os.path.getsize(arquivo) > 0
        resultado.append((f"{nome} existe e não está vazio", existe))
        if not existe:
            continue
        try:
            if arquivo.endswith(".pptx"):
                slides = _contar_slides(arquivo) - 1  # sem a capa
                if tarefa.contexto.get("fontes"):
                    slides -= 1  # nem o slide de fontes
                if tarefa.quantidade_slides:
                    resultado.append((
                        f"{slides} slides de conteúdo (pedido: {tarefa.quantidade_slides})",
                        slides == tarefa.quantidade_slides,
                    ))
                else:
                    resultado.append((f"{slides} slides de conteúdo (mínimo 3)", slides >= 3))
            elif arquivo.endswith(".docx") and tarefa.tipo == "documento":
                palavras = _contar_palavras_docx(arquivo)
                resultado.append((f"Word com {palavras} palavras (mínimo 150)", palavras >= 150))
            elif arquivo.endswith(".pdf"):
                paginas = _contar_paginas_pdf(arquivo)
                resultado.append((f"PDF com {paginas} página(s)", paginas >= 1))
            elif arquivo.endswith(".xlsx"):
                resultado.extend(_verificar_planilha(arquivo, tarefa.pedido))
        except Exception as exc:
            resultado.append((f"{nome} abre sem erro ({exc})", False))

    if tarefa.tipo in ("documento", "apresentacao"):
        resultado.extend(verificar_profundidade(conteudo))
    elif tarefa.tipo == "estudos":
        resultado.extend(_verificar_estudos(tarefa, conteudo))
    elif tarefa.tipo == "aulas":
        resultado.extend(_verificar_aulas(tarefa, conteudo))
    elif tarefa.tipo == "resumo_dia":
        palavras = len(tarefa.contexto.get("resposta", "").split())
        resultado.append((f"Resumo falado com {palavras} palavras (mínimo 40)", palavras >= 40))
    elif tarefa.tipo == "email":
        lidos, resumidos = tarefa.contexto.get("emails_lidos", 0), conteudo.get("emails_resumidos", 0)
        resultado.append((f"{resumidos} de {lidos} e-mail(s) resumidos", resumidos >= lidos))
    return resultado


def verificar_profundidade(conteudo: dict) -> list:
    """Reprova tópicos de uma ou duas palavras e parágrafos rasos, que o LLM costuma deixar passar."""
    resultado = []
    blocos = [b for b in (conteudo.get("slides") or conteudo.get("secoes") or []) if isinstance(b, dict)]
    # O modelo às vezes devolve um texto onde deveria vir uma lista; _lista trata os dois casos.
    topicos = [t for bloco in blocos for t in _lista(bloco.get("topicos"))]
    if topicos:
        media = sum(len(t.split()) for t in topicos) / len(topicos)
        resultado.append((f"Tópicos com {media:.1f} palavras em média (mínimo 5)", media >= 5))
    paragrafos = [p for bloco in blocos for p in _lista(bloco.get("paragrafos"))]
    if paragrafos:
        media = sum(len(p.split()) for p in paragrafos) / len(paragrafos)
        resultado.append((f"Parágrafos com {media:.0f} palavras em média (mínimo 12)", media >= 12))
    return resultado


def _enxugar(valor, profundidade: int = 0):
    """Versão curta de qualquer conteúdo (listas com até 6 itens, textos com até 300 letras)."""
    if isinstance(valor, dict):
        return {chave: _enxugar(v, profundidade + 1) for chave, v in valor.items() if v not in (None, "", [], {})}
    if isinstance(valor, list):
        return [_enxugar(v, profundidade + 1) for v in valor[:6]] + ([f"... e mais {len(valor) - 6}"] if len(valor) > 6 else [])
    if isinstance(valor, str):
        return valor[:300]
    return valor


def _resumo_conteudo(tarefa: Tarefa) -> str:
    conteudo = tarefa.contexto.get("conteudo")
    if not conteudo:
        return tarefa.contexto.get("resposta", "")[:3000]
    if tarefa.tipo == "aulas":
        # Curso: o plano e o começo de cada aula (o texto inteiro passaria do que o crítico precisa).
        enxuto = {
            "curso": conteudo.get("titulo"), "topicos_importantes": conteudo.get("topicos_importantes"),
            "aulas": [{"titulo": a.get("titulo"), "introducao": str(a.get("introducao") or "")[:200],
                       "secoes": [s.get("titulo") for s in a.get("secoes") or [] if isinstance(s, dict)],
                       "exercicios": len(a.get("exercicios") or []), "erro": a.get("erro")}
                      for a in conteudo.get("aulas") or []],
        }
        return json.dumps(enxuto, ensure_ascii=False, default=str)[:6000]
    if tarefa.tipo not in ("documento", "apresentacao"):
        # Planilha, estudos, resumo do dia e e-mail: estrutura própria, enviada de forma enxuta.
        resumo = {"conteudo": _enxugar(conteudo)}
        if tarefa.contexto.get("resposta"):
            resumo["texto_falado"] = tarefa.contexto["resposta"][:1500]
        return json.dumps(resumo, ensure_ascii=False, default=str)[:6000]
    enxuto = {
        "titulo": conteudo.get("titulo"),
        "subtitulo": conteudo.get("subtitulo"),
        "resumo": (conteudo.get("resumo") or "")[:400],
        "secoes": [
            {"titulo": s.get("titulo"), "inicio": " ".join(s.get("paragrafos") or [])[:300], "topicos": (s.get("topicos") or [])[:5]}
            for s in conteudo.get("secoes") or []
        ],
        "slides": [{"titulo": s.get("titulo"), "topicos": s.get("topicos")} for s in conteudo.get("slides") or []],
        "quantidade_de_fontes": len(conteudo.get("fontes") or []),
    }
    return json.dumps({k: v for k, v in enxuto.items() if v}, ensure_ascii=False)[:4000]


# Sem isto, o crítico reprovava a ficha ("o pedido era um resumo simples, mas está enviesado para concurso")
# e o material inteiro era reescrito à toa: o formato foi pedido pelo usuário para todo resumo de estudo.
FORMATO_COMBINADO_ESTUDOS = (
    "Formato combinado com o usuário: todo resumo de estudo é uma ficha tática focada na prova (como a banca cobra, "
    "incidência, natureza e requisitos, regra e exceções, prazos e efeitos, lei seca conferida no Planalto, STF/STJ, "
    "pegadinhas e teste de estresse), mesmo quando o pedido diz só \"resumo\". Esse foco em prova e esse formato não "
    "são problema, e os itens marcados para conferir são de propósito. Avalie a correção do conteúdo, se a regra e as "
    "exceções estão certas e completas e se cada tópico ensina de verdade.\n\n"
)


class AgenteCritico(Agente):
    nome = "critico"
    papel = "Analisa o resultado e decide se o pedido foi cumprido."
    instrucoes = (
        "Você é o Agente Crítico da Beatriz: um revisor exigente e justo do CONTEÚDO entregue. "
        "Avalie: o tema é o que foi pedido, o conteúdo é específico e útil (não genérico), está em "
        "português do Brasil e não há afirmações inventadas. Quando houver fontes, qualquer número, "
        "nome ou fato que não esteja nelas é um problema grave (nota máxima 5). "
        "Seja rigoroso com conteúdo raso: tópicos de uma ou duas palavras, frases vagas ou que só "
        "repetem o título valem nota baixa. NÃO avalie formato de arquivo, quantidade de slides "
        "ou páginas (isso já foi conferido automaticamente) nem imagens, cores ou layout. "
        "Nota: 9-10 excelente, 7-8 bom, 5-6 fraco, 0-4 não atende. "
        "Responda SOMENTE com JSON: "
        '{"nota": inteiro de 0 a 10, "pontos_fortes": [str], '
        '"problemas": [str], "correcao": "instrução objetiva do que o agente deve refazer"}'
    )

    def avaliar(self, tarefa: Tarefa) -> Avaliacao:
        verificacoes = verificacoes_objetivas(tarefa)
        falhas = [descricao for descricao, passou in verificacoes if not passou]
        fontes = tarefa.contexto.get("web", "")
        conferencia = (
            f"\n\nTRECHO DAS FONTES USADAS (confira se números, nomes e fatos do conteúdo aparecem aqui; "
            f"aponte como problema cada afirmação que não aparece):\n{fontes[:LIMITE_CONTEXTO_WEB]}"
            if fontes else ""
        )
        prompt = (
            f"Pedido do usuário: {tarefa.pedido}\n\n"
            f"{FORMATO_COMBINADO_ESTUDOS if tarefa.tipo == 'estudos' else ''}"
            f"Conteúdo entregue (resumido):\n{_resumo_conteudo(tarefa)}{conferencia}"
        )
        try:
            dados = self.pedir_json(prompt, max_tokens=600)
        except Exception as exc:
            logger.warning("Crítico não conseguiu avaliar com o LLM: %s", exc)
            dados = {"nota": NOTA_MINIMA if not falhas else 3,
                     "problemas": ["Avaliação do LLM indisponível; valeram só as verificações automáticas."]}
        try:
            nota = max(0, min(10, int(dados.get("nota", 0))))
        except (TypeError, ValueError):
            nota = 0
        problemas = [f"Verificação automática falhou: {f}" for f in falhas]
        problemas += [str(p) for p in _lista(dados.get("problemas"))]
        # O veredito vem das verificações objetivas + nota; o "cumprido" do modelo pequeno é pouco confiável.
        cumprido = not falhas and nota >= NOTA_MINIMA
        return Avaliacao(
            cumprido=cumprido,
            nota=nota,
            pontos_fortes=_lista(dados.get("pontos_fortes")),
            problemas=problemas,
            correcao=str(dados.get("correcao") or "; ".join(problemas)),
            verificacoes=verificacoes,
        )

    def executar(self, tarefa: Tarefa) -> str:
        avaliacao = self.avaliar(tarefa)
        tarefa.avaliacoes.append(avaliacao)
        veredito = "aprovado" if avaliacao.cumprido else "reprovado"
        return f"Resultado {veredito} com nota {avaliacao.nota}/10."
