"""Ficha tática de estudo: o formato dos resumos, a conferência da lei seca e o PDF/Word coloridos.

O método (análise crítica focada na prova), em quatro filtros:
1. Incidência: antes de escrever, como o assunto vira questão (o que a banca pergunta e o peso de cada tema).
2. Desconstrução: cada instituto no mesmo esqueleto (natureza e requisitos; a regra numa linha e o resto
   do espaço para as exceções; prazos e efeitos; lei seca com a palavra-chave; STF e STJ).
3. Leitura direcionada: primeiro a lei (o texto oficial vem do Planalto) e o destaque das palavras que
   mudam o sentido ("salvo", "apenas", "presume-se"...), que são as pegadinhas.
4. Teste de estresse: questões difíceis respondidas só com a ficha; o que faltou entra na ficha.

Fora do direito, os mesmos campos ganham rótulos equivalentes (conceito, regra geral, exceções, números e
fórmulas, texto-base, posições consolidadas). Nada aqui chama o LLM.
"""
import contextvars
import difflib
import html
import logging
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

logger = logging.getLogger(__name__)

ROTULOS = {
    "direito": {"natureza": "Natureza jurídica e requisitos", "regra": "Regra", "excecoes": "Exceções",
                "prazos": "Prazos e efeitos", "lei_seca": "Lei seca", "jurisprudencia": "STF e STJ"},
    "geral": {"natureza": "Conceito e elementos", "regra": "Regra geral", "excecoes": "Exceções e casos especiais",
              "prazos": "Números, datas e fórmulas", "lei_seca": "Texto-base", "jurisprudencia": "Posições consolidadas"},
}

# Cada parte da ficha tem sempre a mesma cor (e o rótulo escrito, para funcionar impresso em preto e branco).
CORES = {
    "marca": "#2E1A47", "ouro": "#C99A1E", "texto": "#1F2937", "suave": "#64748B",
    "banca": ("#BE185D", "#FDF0F6"), "natureza": ("#0F766E", "#ECFDF9"), "regra": ("#1D4ED8", "#EFF4FF"),
    "excecoes": ("#C2410C", "#FFF4ED"), "prazos": ("#7C3AED", "#F5F0FF"), "lei_seca": ("#334155", "#F1F5F9"),
    "jurisprudencia": ("#15803D", "#EFFBF3"), "pegadinha": ("#A16207", "#FEFCE8"), "exemplo": ("#475569", "#F8FAFC"),
    "estresse": ("#B91C1C", "#FEF2F2"), "destaque": "#FDE047",
}
# As escolhas da janela de estudos para o arquivo que está sendo gerado (preferencias_estudo.visual): cores do tema,
# letra, partes e marca-texto. Sem elas, o padrão.
_VISUAL = contextvars.ContextVar("visual_da_ficha", default=None)


def _visual() -> dict:
    return _VISUAL.get() or {}


def _cores() -> dict:
    return _visual().get("cores") or CORES


def _parte(chave: str) -> bool:
    return (_visual().get("partes") or {}).get(chave, True)


INCIDENCIAS = {"alta": ("ALTA", "#B91C1C"), "media": ("MÉDIA", "#B45309"), "baixa": ("BAIXA", "#64748B")}
# Legenda de cores (rótulos curtos, para caber numa linha).
LEGENDA = {
    "direito": [("banca", "Como cai"), ("natureza", "Requisitos"), ("regra", "Regra"), ("excecoes", "Exceções"),
                ("prazos", "Prazos"), ("lei_seca", "Lei seca"), ("jurisprudencia", "STF/STJ"), ("estresse", "Estresse")],
    "geral": [("banca", "Como cai"), ("natureza", "Conceito"), ("regra", "Regra"), ("excecoes", "Exceções"),
              ("prazos", "Números"), ("lei_seca", "Texto-base"), ("jurisprudencia", "Consensos"), ("estresse", "Estresse")],
}

# Palavras que mudam o sentido jurídico: aparecem com marca-texto em toda a ficha.
PALAVRAS_DE_CONFLITO = [
    r"salvo(?: se)?", r"exceto", r"ressalvad\w*", r"independentemente", r"presume-se", r"presumem-se", r"presun[çc][ãa]o",
    r"apenas", r"somente", r"exclusivamente", r"vedad[oa]s?", r"(?:fica|é|são) vedad[oa]s?", r"proibid[oa]s?",
    r"não se aplica\w*", r"desde que", r"sob pena", r"privativ[oa]s?", r"facultativ[oa]s?", r"obrigatoriamente",
    r"inclusive", r"em regra", r"necessariamente", r"imprescrit[íi]ve(?:l|is)", r"intransfer[íi]ve(?:l|is)",
    r"irrenunci[áa]ve(?:l|is)", r"absolutamente", r"relativamente", r"ex tunc", r"ex nunc", r"de of[íi]cio",
]
RE_CONFLITO = re.compile(r"(?<![\w-])(?:" + "|".join(PALAVRAS_DE_CONFLITO) + r")(?![\w-])", re.IGNORECASE)
# As palavras de conflito que a IA lista para o tópico só ganham marca-texto se forem restritivas ou
# qualificadoras: ela às vezes lista conceitos ("risco integral"), e marca-texto em tudo não destaca nada.
RE_QUALIFICADORA = re.compile(
    r"\b(?:salvo|exceto|apenas|somente|exclusiv\w*|vedad\w*|proibid\w*|independ\w*|presum\w*|desde que|sob pena|privativ\w*|"
    r"facultativ\w*|obrigatori\w*|inclusive|em regra|necessariamente|imprescrit\w*|absolut\w*|relativ\w*|nunca|sempre|"
    r"qualquer|tod[oa]s?|nenhum\w*|ex tunc|ex nunc|de oficio|nao|mediante|ate|no maximo|no minimo|previ[oa]s?|expressa\w*|"
    r"tacit\w*|solidari\w*|subsidiari\w*|concorrente|dispensad\w*|ressalvad\w*|unicamente|tao somente|"
    # números, quóruns e prazos, que a banca troca ("três quintos", "em dois turnos", "anteriores à Constituição")
    r"\d+|um|uma|dois|duas|tres|quatro|cinco|seis|sete|oito|nove|dez|quintos?|tercos?|metade|maioria|unanim\w*|"
    r"turnos?|dias?|meses|anos?|horas?|anterior\w*|posterior\w*|tendente|precipuamente|originari\w*)\b"
)


def qualificadoras(extras) -> list[str]:
    return [e for e in extras if e and RE_QUALIFICADORA.search(_normal(e))]


# ----------------------------------------------------------------- conteúdo
def _lista(valor) -> list:
    if not valor:
        return []
    if isinstance(valor, (str, dict)):
        return [valor]
    return [item for item in valor if item not in (None, "", [], {})]


def _textos(valor) -> list[str]:
    return [str(item).strip() for item in _lista(valor) if str(item).strip()]


def _incidencia(valor) -> str:
    texto = str(valor or "").lower()
    if texto.startswith("a"):
        return "alta"
    if texto.startswith("b"):
        return "baixa"
    return "media"


def _item_lei(item) -> dict | None:
    if isinstance(item, str):
        item = {"referencia": item}
    if not isinstance(item, dict):
        return None
    referencia = str(item.get("referencia") or "").strip()
    norma, artigo = str(item.get("norma") or "").strip(), str(item.get("artigo") or "").strip()
    if not referencia and (norma or artigo):
        referencia = f"art. {artigo} {norma}".strip()
    if not referencia:
        return None
    return {**item, "referencia": referencia, "norma": norma, "artigo": artigo,
            "palavra_chave": str(item.get("palavra_chave") or "").strip(), "pegadinha": str(item.get("pegadinha") or "").strip()}


def _item_juris(item) -> dict | None:
    if isinstance(item, str):
        item = {"tese": item}
    if not isinstance(item, dict) or not (item.get("tese") or item.get("referencia")):
        return None
    return {**item, "referencia": str(item.get("referencia") or "").strip(), "tese": str(item.get("tese") or "").strip(),
            "conferir": bool(item.get("conferir"))}


def _prazos(valor) -> list[dict]:
    saida = []
    for item in _lista(valor):
        if isinstance(item, str):
            item = {"item": item, "valor": ""}
        if isinstance(item, dict) and (item.get("item") or item.get("valor")):
            saida.append({"item": str(item.get("item") or "").strip(), "valor": str(item.get("valor") or "").strip(),
                          "fundamento": str(item.get("fundamento") or "").strip()})
    return saida


def normalizar_topico(topico: dict) -> dict:
    """Aceita o formato novo e o antigo ("paragrafos"/"explicacao" texto, "atencao")."""
    explicacao = _textos(topico.get("explicacao")) or _textos(topico.get("paragrafos"))
    requisitos = _textos(topico.get("requisitos"))
    natureza = topico.get("natureza")
    if isinstance(topico.get("natureza_e_requisitos"), dict):
        natureza = natureza or topico["natureza_e_requisitos"].get("natureza")
        requisitos = requisitos or _textos(topico["natureza_e_requisitos"].get("requisitos"))
    return {
        "titulo": str(topico.get("titulo") or "").strip(),
        "incidencia": _incidencia(topico.get("incidencia")),
        "como_cai": str(topico.get("como_cai") or "").strip(),
        "explicacao": explicacao,
        "natureza": str(natureza or "").strip(),
        "requisitos": requisitos,
        "regra": str(topico.get("regra") or "").strip(),
        "excecoes": _textos(topico.get("excecoes")),
        "prazos_e_efeitos": _prazos(topico.get("prazos_e_efeitos")),
        "lei_seca": [i for i in map(_item_lei, _lista(topico.get("lei_seca"))) if i],
        "jurisprudencia": [i for i in map(_item_juris, _lista(topico.get("jurisprudencia"))) if i],
        "palavras_de_conflito": _textos(topico.get("palavras_de_conflito")),
        "pegadinhas": _textos(topico.get("pegadinhas")) or _textos(topico.get("atencao")),
        "exemplo": str(topico.get("exemplo") or "").strip(),
        "pontos_chave": _textos(topico.get("pontos_chave")),
        "teste_de_estresse": _textos(topico.get("teste_de_estresse")),
        "fonte": str(topico.get("fonte") or "").strip(),
    }


def normalizar_ficha(conteudo: dict) -> dict:
    ficha = dict(conteudo)
    ficha["area"] = "direito" if str(conteudo.get("area") or "").lower().startswith("dir") else "geral"
    ficha["visao_geral"] = str(conteudo.get("visao_geral") or conteudo.get("resumo") or "").strip()
    ficha["resumo"] = ficha["visao_geral"]  # os outros pontos do sistema (relatório, WhatsApp) leem "resumo"
    ficha["como_a_banca_cobra"] = _textos(conteudo.get("como_a_banca_cobra"))
    ficha["topicos"] = [normalizar_topico(t) for t in _lista(conteudo.get("topicos")) if isinstance(t, dict) and t.get("titulo")]
    ficha["glossario"] = [g for g in _lista(conteudo.get("glossario")) if isinstance(g, dict) and g.get("termo") and g.get("definicao")]
    ficha["fora_das_fontes"] = _textos(conteudo.get("fora_das_fontes"))
    ficha["lacunas_gerais"] = _textos(conteudo.get("lacunas_gerais"))
    ficha["dicas"] = _textos(conteudo.get("dicas"))
    return ficha


def quadro_de_prazos(ficha: dict) -> list[dict]:
    """Todos os prazos e efeitos da matéria numa tabela só (para revisar na véspera)."""
    linhas = []
    for topico in ficha.get("topicos") or []:
        for prazo in topico.get("prazos_e_efeitos") or []:
            linhas.append({**prazo, "topico": topico["titulo"]})
    return linhas


def _normal(texto: str) -> str:
    import unicodedata

    texto = unicodedata.normalize("NFKD", str(texto or "").lower())
    return re.sub(r"\s+", " ", "".join(c for c in texto if not unicodedata.combining(c))).strip()


def topico_mais_parecido(ficha: dict, titulo: str) -> dict | None:
    topicos = ficha.get("topicos") or []
    nomes = [_normal(t["titulo"]) for t in topicos]
    alvo = _normal(titulo)
    if not alvo:
        return None
    for nome, topico in zip(nomes, topicos):
        if alvo == nome or alvo in nome or nome in alvo:
            return topico
    parecido = difflib.get_close_matches(alvo, nomes, n=1, cutoff=0.55)
    return topicos[nomes.index(parecido[0])] if parecido else None


def _palavras(texto: str) -> set:
    return {p for p in re.findall(r"[a-z]{5,}", _normal(texto))}


def topico_pelo_conteudo(ficha: dict, texto: str) -> dict | None:
    """O tópico que mais divide palavras com o texto (quando a IA cita um tópico com outro nome)."""
    alvo = _palavras(texto)
    melhor, pontos = None, 0
    for topico in ficha.get("topicos") or []:
        comuns = alvo & (_palavras(topico["titulo"]) | _palavras(topico.get("regra") or ""))
        peso = len(comuns & _palavras(topico["titulo"])) * 2 + len(comuns)  # o título conta mais
        if peso > pontos:
            melhor, pontos = topico, peso
    return melhor if pontos >= 3 else None


def aplicar_lacunas(ficha: dict, lacunas) -> int:
    """O que o teste de estresse achou que faltava entra no tópico certo (ou num quadro geral)."""
    aplicadas = 0
    for lacuna in _lista(lacunas):
        if isinstance(lacuna, str):
            lacuna = {"tese": lacuna}
        if not isinstance(lacuna, dict) or not str(lacuna.get("tese") or "").strip():
            continue
        tese = str(lacuna["tese"]).strip()
        topico = (topico_mais_parecido(ficha, lacuna.get("topico") or "")
                  or topico_pelo_conteudo(ficha, f"{lacuna.get('topico') or ''} {tese}"))
        destino = topico["teste_de_estresse"] if topico else ficha.setdefault("lacunas_gerais", [])
        if tese not in destino:
            destino.append(tese)
            aplicadas += 1
    return aplicadas


# ------------------------------------------------------------- lei seca
RE_NUMERO_ARTIGO = re.compile(r"(\d{1,2}\.\d{3}|\d{1,4})(?:\s*-\s*([A-Za-z])\b)?")


def _numero_do_item(item: dict) -> str:
    from agente_leis import _numero_do_artigo

    for campo in (item.get("artigo"), re.sub(r"(?i)^.*?\bart(?:igo)?s?\.?\s*", "", item.get("referencia") or "")):
        achado = RE_NUMERO_ARTIGO.search(str(campo or ""))
        if achado:
            return _numero_do_artigo(achado.group(1), achado.group(2) or "")
    return ""


def conferir_lei_seca(ficha: dict, textos=None) -> dict:
    """Busca no Planalto o texto oficial de cada artigo citado na lei seca.

    Marca cada item: "conferido" True (achou o artigo; o caput vai para a ficha), False (a lei abriu, mas
    o artigo não existe nela: provável número errado) ou None (lei fora do alcance ou sem internet).
    """
    from agente_leis import (AgenteLeis, TextoDasLeis, norma_por_texto, sem_anotacoes,
                             trecho_do_artigo)

    textos = textos or TextoDasLeis()
    buscador = AgenteLeis(textos=textos, explicador=lambda *_: "")
    itens = []
    for topico in ficha.get("topicos") or []:
        for item in topico.get("lei_seca") or []:
            norma = norma_por_texto(f"{item.get('norma', '')} {item.get('referencia', '')}")
            numero = _numero_do_item(item)
            item["conferido"] = None
            if norma and numero:
                itens.append((item, norma, numero))
    normas = {n.chave: n for _, n, _ in itens}

    def abrir(norma):
        try:
            return norma.chave, buscador._linhas(norma)
        except Exception as exc:
            logger.info("Lei seca: não abri %s (%s)", norma.nome, exc)
            return norma.chave, None

    with ThreadPoolExecutor(max_workers=4, thread_name_prefix="lei-seca") as executor:
        paginas = dict(executor.map(abrir, normas.values()))
    contagem = {"conferidos": 0, "nao_encontrados": 0, "sem_acesso": 0}
    for item, norma, numero in itens:
        linhas = paginas.get(norma.chave)
        if linhas is None:
            contagem["sem_acesso"] += 1
            continue
        trecho = trecho_do_artigo(linhas, numero)
        if not trecho:
            item["conferido"] = False
            contagem["nao_encontrados"] += 1
            continue
        texto = " ".join(sem_anotacoes(linha) for linha in trecho)
        item["texto_oficial"] = texto if len(texto) <= 700 else texto[:700].rsplit(" ", 1)[0] + " [...]"
        item["conferido"] = True
        item["norma_nome"] = norma.nome
        contagem["conferidos"] += 1
    return contagem


def conferir_jurisprudencia(ficha: dict, vade=None) -> dict:
    """Confere no vade-mécum do PC cada súmula citada na jurisprudência ("Súmula 7 do STJ", "SV 11", "OJ 191 da SDI-1").

    Achou: o enunciado oficial entra no item ("texto_oficial"), a marca "conferir" sai e a situação fica (uma súmula
    cancelada aparece como cancelada). O tribunal tem súmulas no banco, mas não essa: "conferido" False e "conferir".
    Sem as súmulas daquele tribunal no banco (ou sem número na citação), o item fica como veio.
    """
    import vade_mecum
    from agenda_linguagem import preparar
    from agente_leis import pedido_de_sumula

    vade = vade or vade_mecum.vade_padrao()
    total = vade.total()["sumulas"]
    contagem = {"conferidas": 0, "nao_encontradas": 0, "canceladas": 0}
    for topico in ficha.get("topicos") or []:
        for item in topico.get("jurisprudencia") or []:
            pedido = pedido_de_sumula(preparar(f"{item.get('referencia', '')} {item.get('tese', '')}")[1])
            if not pedido or not pedido["tribunal"]:
                continue
            if not any(chave.startswith(pedido["tribunal"] + "_") for chave in total):
                continue  # o banco ainda não tem as súmulas desse tribunal
            sumula = vade.sumula(pedido["tribunal"], pedido["tipo"], pedido["numero"], pedido["orgao"])
            if not sumula:
                item.update(conferido=False, conferir=True)
                contagem["nao_encontradas"] += 1
                continue
            item.update(conferido=True, conferir=False, texto_oficial=sumula["texto"], situacao=sumula["situacao"])
            contagem["conferidas"] += 1
            if sumula["situacao"] in ("cancelada", "revogada", "superada"):
                contagem["canceladas"] += 1
    return contagem


# ------------------------------------------------------------------ PDF
def destacar(texto: str, extras=()) -> str:
    """HTML do texto com marca-texto nas palavras de conflito (as da lista e as do próprio tópico)."""
    seguro = html.escape(str(texto or ""))
    if not _visual().get("marca_texto", True):
        return seguro
    intervalos = [m.span() for m in RE_CONFLITO.finditer(seguro)]
    padroes = [re.escape(html.escape(p.strip())) for p in extras if p and len(p.strip()) > 2]
    if padroes:
        expressao = r"(?<![\w-])(?:" + "|".join(padroes) + r")(?![\w-])"
        intervalos += [m.span() for m in re.finditer(expressao, seguro, flags=re.IGNORECASE)]
    if not intervalos:
        return seguro
    # Trechos que se sobrepõem ("salvo se houver" e "salvo se") viram um destaque só.
    unidos = []
    for inicio, fim in sorted(intervalos):
        if unidos and inicio <= unidos[-1][1]:
            unidos[-1][1] = max(unidos[-1][1], fim)
        else:
            unidos.append([inicio, fim])
    partes, posicao = [], 0
    for inicio, fim in unidos:
        partes += [seguro[posicao:inicio], '<span class="mt">', seguro[inicio:fim], "</span>"]
        posicao = fim
    partes.append(seguro[posicao:])
    return "".join(partes)


def _css() -> str:
    c = _cores()
    letra = 12 if _visual().get("letra") == "grande" else 10.5
    caixas = "\n".join(
        f".{nome} {{ background-color: {fundo}; border-left: 4px solid {borda}; }}\n.{nome} .rot {{ color: {borda}; }}"
        for nome, (borda, fundo) in ((k, v) for k, v in c.items() if isinstance(v, tuple))
    )
    return f"""
body {{ font-family: sans-serif; font-size: {letra}pt; color: {c['texto']}; line-height: 1.4; }}
.topo {{ background-color: {c['marca']}; padding: 12pt 14pt; margin-bottom: 8pt; }}
.titulo {{ color: #FFFFFF; font-size: 21pt; font-weight: bold; margin: 0; }}
.sub {{ color: {c['ouro']}; font-size: 10pt; margin: 2pt 0 0 0; }}
.data {{ color: #CBD5E1; font-size: 8pt; margin: 2pt 0 0 0; }}
h1 {{ color: {c['marca']}; font-size: 15pt; border-bottom: 2px solid {c['ouro']}; margin: 14pt 0 6pt 0; }}
h2 {{ color: {c['marca']}; font-size: 13.5pt; border-bottom: 1.5px solid {c['ouro']}; margin: 16pt 0 4pt 0; }}
.num {{ background-color: {c['marca']}; color: #FFFFFF; padding: 1pt 5pt; }}
.selo {{ color: #FFFFFF; font-size: 7.5pt; font-weight: bold; padding: 1pt 4pt; }}
.caixa {{ padding: 5pt 8pt; margin: 5pt 0; }}
.rot {{ font-size: 7.5pt; font-weight: bold; letter-spacing: 1px; margin: 0 0 2pt 0; }}
.caixa p {{ margin: 2pt 0; }}
.caixa ul {{ margin: 2pt 0 2pt 0; }}
li {{ margin: 1.5pt 0; }}
.regra-texto {{ font-weight: bold; font-size: 11pt; }}
.citacao {{ font-family: serif; font-style: italic; color: #111827; }}
.ok {{ color: #15803D; font-size: 7.5pt; font-weight: bold; }}
.alerta {{ color: #B91C1C; font-size: 7.5pt; font-weight: bold; }}
.mt {{ background-color: {c['destaque']}; font-weight: bold; }}
.cai-linha {{ color: {c["banca"][0]}; margin: 2pt 0 4pt 0; }}
.fonte {{ color: {c['suave']}; font-size: 8pt; }}
.entenda p {{ text-align: justify; margin: 3pt 0; }}
table {{ border-collapse: collapse; width: 100%; margin: 4pt 0; }}
th {{ padding: 3pt 5pt; text-align: left; font-size: 9pt; font-weight: bold; }}
td {{ border: 1px solid #E2E8F0; padding: 3pt 5pt; font-size: 9pt; vertical-align: top; }}
.t-prazos th {{ color: {c['prazos'][0]}; border-bottom: 2px solid {c['prazos'][0]}; }}
.t-padrao th {{ color: {c['marca']}; border-bottom: 2px solid {c['marca']}; }}
.legenda td {{ border: none; font-size: 8pt; padding: 2pt 4pt; }}
.questao {{ margin: 6pt 0 2pt 0; font-weight: bold; color: {c['marca']}; }}
.pequeno {{ font-size: 8.5pt; color: {c['suave']}; }}
{caixas}
"""


def _caixa(classe: str, rotulo: str, corpo: str) -> str:
    return f'<div class="caixa {classe}"><p class="rot">{html.escape(rotulo.upper())}</p>{corpo}</div>'


def _ul(itens, extras=()) -> str:
    return "<ul>" + "".join(f"<li>{destacar(i, extras)}</li>" for i in itens) + "</ul>"


def _selo(incidencia: str) -> str:
    nome, cor = INCIDENCIAS.get(incidencia, INCIDENCIAS["media"])
    return f'<span class="selo" style="background-color: {cor}">INCIDÊNCIA {nome}</span>'


def _html_topico(numero: int, topico: dict, rotulos: dict) -> str:
    extras = qualificadoras(topico["palavras_de_conflito"])
    partes = [f'<h2><span class="num">{numero}</span> {html.escape(topico["titulo"])} &nbsp;{_selo(topico["incidencia"])}</h2>']
    if topico["como_cai"] and _parte("banca"):
        partes.append(f'<p class="cai-linha"><b>Como cai:</b> {destacar(topico["como_cai"], extras)}</p>')
    if topico["explicacao"]:
        partes.append('<div class="entenda">' + "".join(f"<p>{destacar(p, extras)}</p>" for p in topico["explicacao"]) + "</div>")
    if topico["natureza"] or topico["requisitos"]:
        corpo = f'<p>{destacar(topico["natureza"], extras)}</p>' if topico["natureza"] else ""
        if topico["requisitos"]:
            corpo += "<p><b>Requisitos:</b></p>" + _ul(topico["requisitos"], extras)
        partes.append(_caixa("natureza", rotulos["natureza"], corpo))
    if topico["regra"]:
        partes.append(_caixa("regra", rotulos["regra"], f'<p class="regra-texto">{destacar(topico["regra"], extras)}</p>'))
    if topico["excecoes"]:
        partes.append(_caixa("excecoes", rotulos["excecoes"], _ul(topico["excecoes"], extras)))
    if topico["prazos_e_efeitos"] and _parte("prazos"):
        linhas = "".join(f"<tr><td>{destacar(p['item'], extras)}</td><td><b>{destacar(p['valor'], extras)}</b></td>"
                         f"<td>{html.escape(p.get('fundamento') or '')}</td></tr>" for p in topico["prazos_e_efeitos"])
        partes.append(f'<p class="rot" style="color: {_cores()["prazos"][0]}">{html.escape(rotulos["prazos"].upper())}</p>'
                      f'<table class="t-prazos"><tr><th>O quê</th><th>Prazo / efeito</th><th>Fundamento</th></tr>{linhas}</table>')
    if topico["lei_seca"] and _parte("lei_seca"):
        corpo = ""
        for item in topico["lei_seca"]:
            selo = ""
            if item.get("conferido") is True:
                selo = ' <span class="ok">TEXTO CONFERIDO NO PLANALTO</span>'
            elif item.get("conferido") is False:
                selo = ' <span class="alerta">NÃO ACHEI ESTE ARTIGO NO PLANALTO: CONFIRA O NÚMERO</span>'
            corpo += f"<p><b>{html.escape(item['referencia'])}</b>{selo}</p>"
            if item.get("palavra_chave"):
                corpo += f'<p>Palavra-chave: <span class="mt">{html.escape(item["palavra_chave"])}</span></p>'
            if item.get("texto_oficial"):
                corpo += f'<p class="citacao">“{destacar(item["texto_oficial"], extras + [item.get("palavra_chave", "")])}”</p>'
            if item.get("pegadinha"):
                corpo += f"<p><b>Pegadinha:</b> {destacar(item['pegadinha'], extras)}</p>"
        partes.append(_caixa("lei_seca", rotulos["lei_seca"], corpo))
    if topico["jurisprudencia"] and _parte("jurisprudencia"):
        itens = []
        for item in topico["jurisprudencia"]:
            aviso = ' <span class="alerta">CONFERIR O NÚMERO</span>' if item["conferir"] else ""
            if item.get("situacao") in ("cancelada", "revogada", "superada"):
                aviso += f' <span class="alerta">SÚMULA {item["situacao"].upper()}</span>'
            ref = f"<b>{html.escape(item['referencia'])}</b>: " if item["referencia"] else ""
            oficial = ""
            if item.get("texto_oficial"):
                oficial = f'<br/><span class="citacao">Texto oficial: “{destacar(item["texto_oficial"], extras)}”</span>'
            itens.append(f"<li>{ref}{destacar(item['tese'], extras)}{aviso}{oficial}</li>")
        partes.append(_caixa("jurisprudencia", rotulos["jurisprudencia"], "<ul>" + "".join(itens) + "</ul>"))
    if topico["pegadinhas"] and _parte("pegadinhas"):
        partes.append(_caixa("pegadinha", "Pegadinhas", _ul(topico["pegadinhas"], extras)))
    if topico["exemplo"] and _parte("exemplo"):
        partes.append(_caixa("exemplo", "Exemplo", f"<p>{destacar(topico['exemplo'], extras)}</p>"))
    if topico["teste_de_estresse"]:
        partes.append(_caixa("estresse", "Acrescentado no teste de estresse", _ul(topico["teste_de_estresse"], extras)))
    if topico["pontos_chave"]:
        partes.append('<p class="rot" style="color: #475569">PARA REVISAR</p>' + _ul(topico["pontos_chave"], extras))
    if topico["fonte"]:
        partes.append(f'<p class="fonte">Fonte: {html.escape(topico["fonte"])}</p>')
    return "".join(partes)


LETRAS = "ABCD"


COMO_ESTUDAR = [
    "Comece pelo mapa de incidência: estude primeiro os tópicos de peso alto.",
    "Em cada tópico, leia o 'Como cai' antes da explicação: você passa a ler procurando a resposta da prova.",
    "Decore a regra numa linha e gaste o tempo nas exceções: é delas que saem as questões.",
    "Leia a lei seca em voz alta e repare no marca-texto: a banca copia o artigo trocando uma dessas palavras.",
    "Faça o teste de estresse sem olhar. Errou? Volte ao tópico indicado e reescreva a tese com as suas palavras.",
    "Revise pelos flashcards (Anki) e pelo quadro de prazos na véspera.",
]


def html_da_ficha(ficha: dict) -> str:
    rotulos = ROTULOS[ficha.get("area", "geral")]
    partes = [
        '<div class="topo">'
        f'<p class="titulo">{html.escape(ficha.get("titulo") or "Ficha de estudo")}</p>'
        f'<p class="sub">Ficha tática de estudo · análise crítica focada na prova</p>'
        f'<p class="data">Gerado pela Beatriz em {datetime.now():%d/%m/%Y às %H:%M}</p></div>'
    ]
    if ficha.get("visao_geral"):
        partes.append(f"<p>{destacar(ficha['visao_geral'])}</p>")
    if ficha.get("como_a_banca_cobra") and _parte("banca"):
        partes.append(_caixa("banca", "Como a banca cobra este assunto", _ul(ficha["como_a_banca_cobra"])))
    topicos = ficha.get("topicos") or []
    if topicos and _parte("mapa"):
        linhas = "".join(f"<tr><td>{n}. {html.escape(t['titulo'])}</td><td>{_selo(t['incidencia'])}</td>"
                         f"<td>{destacar(t['como_cai'])}</td></tr>" for n, t in enumerate(topicos, start=1))
        partes.append(f'<h1>Mapa de incidência</h1><table class="t-padrao"><tr><th>Tópico</th><th>Peso</th>'
                      f"<th>Como vira questão</th></tr>{linhas}</table>")
    chips = "".join(f'<td><span class="selo" style="background-color: {_cores()[k][0]}">{html.escape(v.upper())}</span></td>'
                    for k, v in LEGENDA[ficha.get("area", "geral")] if _parte(k) and (k != "estresse" or ficha.get("questoes")))
    partes.append(f'<p class="pequeno">Como ler: cada cor é sempre a mesma parte da ficha. '
                  f'<span class="mt">Marca-texto</span> = palavra que muda o sentido (pegadinha).</p>'
                  f'<table class="legenda"><tr>{chips}</tr></table>')
    for numero, topico in enumerate(topicos, start=1):
        partes.append(_html_topico(numero, topico, rotulos))
    prazos = quadro_de_prazos(ficha) if _parte("prazos") else []
    if prazos:
        linhas = "".join(f"<tr><td>{html.escape(p['topico'])}</td><td>{destacar(p['item'])}</td><td><b>{destacar(p['valor'])}</b></td>"
                         f"<td>{html.escape(p.get('fundamento') or '')}</td></tr>" for p in prazos)
        partes.append(f'<h1>Quadro geral: {html.escape(rotulos["prazos"].lower())}</h1><table class="t-prazos"><tr><th>Tópico</th>'
                      f"<th>O quê</th><th>Prazo / efeito</th><th>Fundamento</th></tr>{linhas}</table>")
    if ficha.get("lacunas_gerais"):
        partes.append(_caixa("estresse", "Acrescentado no teste de estresse (geral)", _ul(ficha["lacunas_gerais"])))
    questoes = ficha.get("questoes") or []
    if questoes:
        partes.append('<h1>Teste de estresse</h1><p class="pequeno">Responda sem olhar a ficha. Errou? A questão diz '
                      "qual tópico tinha a resposta: volte nele.</p>")
        for n, q in enumerate(questoes, start=1):
            topico = f' <span class="pequeno">(tópico: {html.escape(q["topico"])})</span>' if q.get("topico") else ""
            partes.append(f'<p class="questao">Questão {n}{topico}</p><p>{destacar(q["enunciado"])}</p>'
                          + "<ul>" + "".join(f"<li>{l}) {destacar(a)}</li>" for l, a in zip(LETRAS, q["alternativas"])) + "</ul>")
        partes.append("<h1>Gabarito comentado</h1>" + "".join(
            f"<p><b>{n}. {q['correta']}</b>: {destacar(q.get('explicacao') or '')}</p>" for n, q in enumerate(questoes, start=1)))
    flashcards = ficha.get("flashcards") or []
    if flashcards:
        linhas = "".join(f"<tr><td>{destacar(c['frente'])}</td><td>{destacar(c['verso'])}</td></tr>" for c in flashcards)
        partes.append('<h1>Flashcards</h1><p class="pequeno">Os mesmos cartões estão no arquivo para o Anki, na mesma pasta '
                      f'(Arquivo &gt; Importar).</p><table class="t-padrao"><tr><th>Pergunta</th><th>Resposta</th></tr>{linhas}</table>')
    if ficha.get("glossario") and _parte("glossario"):
        linhas = "".join(f"<tr><td><b>{html.escape(g['termo'])}</b></td><td>{destacar(g['definicao'])}</td></tr>"
                         for g in ficha["glossario"])
        partes.append(f'<h1>Glossário</h1><table class="t-padrao"><tr><th>Termo</th><th>Definição</th></tr>{linhas}</table>')
    if ficha.get("fora_das_fontes"):
        partes.append("<h1>Complementado além das suas fontes</h1><p class=\"pequeno\">Veio do conhecimento geral da IA, não "
                      "dos seus arquivos: confira no material da aula.</p>" + _ul(ficha["fora_das_fontes"]))
    if _parte("como_estudar"):
        partes.append("<h1>Como estudar com esta ficha</h1>" + _ul(COMO_ESTUDAR))
    fontes = _textos(ficha.get("fontes"))
    if fontes:
        partes.append('<h1>Fontes</h1><ol class="pequeno">' + "".join(f"<li>{html.escape(f)}</li>" for f in fontes) + "</ol>")
    return "".join(partes)


def gerar_pdf_ficha(ficha: dict, caminho: str, visual: dict = None) -> str:
    """O PDF da ficha. `visual` (preferencias_estudo.visual) traz as cores, a letra, as partes e o marca-texto; sem ele,
    vale o que ficou guardado na ficha (`ficha["visual"]`) ou o padrão."""
    marca = _VISUAL.set(visual if visual is not None else ficha.get("visual"))
    try:
        return _gerar_pdf_ficha(ficha, caminho)
    finally:
        _VISUAL.reset(marca)


def _gerar_pdf_ficha(ficha: dict, caminho: str) -> str:
    import pymupdf

    historia = pymupdf.Story(html=html_da_ficha(ficha), user_css=_css())
    pagina = pymupdf.paper_rect("a4")
    area = pagina + (42, 42, -42, -50)
    escritor = pymupdf.DocumentWriter(caminho)
    mais = True
    while mais:
        dispositivo = escritor.begin_page(pagina)
        mais, _ = historia.place(area)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()
    documento = pymupdf.open(caminho)
    total = documento.page_count
    titulo = (ficha.get("titulo") or "")[:70]
    for indice, pagina_pdf in enumerate(documento, start=1):
        largura, altura = pagina_pdf.rect.width, pagina_pdf.rect.height
        pagina_pdf.insert_text((42, altura - 24), titulo, fontsize=7.5, color=(0.39, 0.45, 0.55))
        pagina_pdf.insert_text((largura - 100, altura - 24), f"Página {indice} de {total}", fontsize=7.5,
                               color=(0.39, 0.45, 0.55))
    documento.saveIncr()
    documento.close()
    return caminho


# ----------------------------------------------------------------- Word
def _rgb(cor: str):
    from docx.shared import RGBColor

    return RGBColor.from_string(cor.lstrip("#").upper())


def _texto_com_destaques(paragrafo, texto: str, extras=(), negrito=False, italico=False, cor=None):
    """Escreve o texto no parágrafo do Word com marca-texto amarelo nas palavras de conflito."""
    from docx.enum.text import WD_COLOR_INDEX

    marcado = destacar(texto, extras)
    for pedaco in re.split(r'(<span class="mt">.*?</span>)', marcado):
        if not pedaco:
            continue
        destaque = pedaco.startswith('<span class="mt">')
        conteudo = html.unescape(re.sub(r"<[^>]+>", "", pedaco))
        run = paragrafo.add_run(conteudo)
        run.bold = negrito or destaque
        run.italic = italico
        if destaque:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if cor:
            run.font.color.rgb = _rgb(cor)
    return paragrafo


def _sombrear(celula, fundo: str, borda: str):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    propriedades = celula._tc.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), fundo.lstrip("#"))
    propriedades.append(sombra)
    bordas = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        elemento = OxmlElement(f"w:{lado}")
        elemento.set(qn("w:val"), "single")
        elemento.set(qn("w:sz"), "24" if lado == "left" else "2")
        elemento.set(qn("w:color"), (borda if lado == "left" else fundo).lstrip("#"))
        bordas.append(elemento)
    propriedades.append(bordas)


def _caixa_docx(doc, classe: str, rotulo: str, escrever):
    from docx.shared import Pt

    borda, fundo = _cores()[classe]
    tabela = doc.add_table(rows=1, cols=1)
    celula = tabela.cell(0, 0)
    _sombrear(celula, fundo, borda)
    titulo = celula.paragraphs[0]
    run = titulo.add_run(rotulo.upper())
    run.bold, run.font.size, run.font.color.rgb = True, Pt(8), _rgb(borda)
    escrever(celula)
    doc.add_paragraph()


def _lista_docx(celula, itens, extras=()):
    for item in itens:
        _texto_com_destaques(celula.add_paragraph(style="List Bullet"), item, extras)


def _item_juris_docx(item: dict) -> str:
    texto = f"{item['referencia']}: {item['tese']}".strip(": ")
    if item["conferir"]:
        texto += " (conferir o número)"
    if item.get("situacao") in ("cancelada", "revogada", "superada"):
        texto += f" (súmula {item['situacao']})"
    if item.get("texto_oficial"):
        texto += f" Texto oficial: “{item['texto_oficial']}”"
    return texto


def gerar_docx_ficha(ficha: dict, caminho: str, visual: dict = None) -> str:
    """O Word da ficha, com as mesmas escolhas do PDF (veja `gerar_pdf_ficha`)."""
    marca = _VISUAL.set(visual if visual is not None else ficha.get("visual"))
    try:
        return _gerar_docx_ficha(ficha, caminho)
    finally:
        _VISUAL.reset(marca)


def _gerar_docx_ficha(ficha: dict, caminho: str) -> str:
    from docx import Document
    from docx.shared import Pt

    rotulos = ROTULOS[ficha.get("area", "geral")]
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(12 if _visual().get("letra") == "grande" else 10.5)
    for nivel, tamanho in (("Title", 24), ("Heading 1", 15), ("Heading 2", 13)):
        doc.styles[nivel].font.color.rgb = _rgb(_cores()["marca"])
        doc.styles[nivel].font.size = Pt(tamanho)
    doc.add_paragraph(ficha.get("titulo") or "Ficha de estudo", style="Title")
    _texto_com_destaques(doc.add_paragraph(), "Ficha tática de estudo · análise crítica focada na prova", cor=_cores()["ouro"])
    if ficha.get("visao_geral"):
        _texto_com_destaques(doc.add_paragraph(), ficha["visao_geral"])
    if ficha.get("como_a_banca_cobra") and _parte("banca"):
        _caixa_docx(doc, "banca", "Como a banca cobra este assunto", lambda c: _lista_docx(c, ficha["como_a_banca_cobra"]))
    topicos = ficha.get("topicos") or []
    if topicos and _parte("mapa"):
        doc.add_heading("Mapa de incidência", level=1)
        tabela = doc.add_table(rows=1, cols=3)
        tabela.style = "Table Grid"
        for celula, texto in zip(tabela.rows[0].cells, ("Tópico", "Peso", "Como vira questão")):
            celula.text = texto
        for n, topico in enumerate(topicos, start=1):
            celulas = tabela.add_row().cells
            celulas[0].text = f"{n}. {topico['titulo']}"
            celulas[1].text = INCIDENCIAS[topico["incidencia"]][0]
            celulas[2].text = topico["como_cai"]
    for n, topico in enumerate(topicos, start=1):
        extras = qualificadoras(topico["palavras_de_conflito"])
        doc.add_heading(f"{n}. {topico['titulo']} (incidência {INCIDENCIAS[topico['incidencia']][0].lower()})", level=2)
        if topico["como_cai"] and _parte("banca"):
            paragrafo = doc.add_paragraph()
            paragrafo.add_run("Como cai: ").bold = True
            _texto_com_destaques(paragrafo, topico["como_cai"], extras, cor=_cores()["banca"][0])
        for texto in topico["explicacao"]:
            _texto_com_destaques(doc.add_paragraph(), texto, extras)
        if topico["natureza"] or topico["requisitos"]:
            def natureza(c, t=topico, e=extras):
                if t["natureza"]:
                    _texto_com_destaques(c.add_paragraph(), t["natureza"], e)
                _lista_docx(c, [f"Requisito: {r}" for r in t["requisitos"]], e)
            _caixa_docx(doc, "natureza", rotulos["natureza"], natureza)
        if topico["regra"]:
            _caixa_docx(doc, "regra", rotulos["regra"],
                        lambda c, t=topico, e=extras: _texto_com_destaques(c.add_paragraph(), t["regra"], e, negrito=True))
        if topico["excecoes"]:
            _caixa_docx(doc, "excecoes", rotulos["excecoes"], lambda c, t=topico, e=extras: _lista_docx(c, t["excecoes"], e))
        if topico["prazos_e_efeitos"] and _parte("prazos"):
            _caixa_docx(doc, "prazos", rotulos["prazos"], lambda c, t=topico, e=extras: _lista_docx(
                c, [f"{p['item']}: {p['valor']}" + (f" ({p['fundamento']})" if p.get("fundamento") else "") for p in t["prazos_e_efeitos"]], e))
        if topico["lei_seca"] and _parte("lei_seca"):
            def lei(c, t=topico, e=extras):
                for item in t["lei_seca"]:
                    selo = {True: " (texto conferido no Planalto)", False: " (não achei no Planalto: confira o número)"}.get(item.get("conferido"), "")
                    _texto_com_destaques(c.add_paragraph(), item["referencia"] + selo, negrito=True)
                    if item.get("palavra_chave"):
                        _texto_com_destaques(c.add_paragraph(), f"Palavra-chave: {item['palavra_chave']}", e + [item["palavra_chave"]])
                    if item.get("texto_oficial"):
                        _texto_com_destaques(c.add_paragraph(), f"“{item['texto_oficial']}”", e + [item.get("palavra_chave", "")], italico=True)
                    if item.get("pegadinha"):
                        _texto_com_destaques(c.add_paragraph(), f"Pegadinha: {item['pegadinha']}", e)
            _caixa_docx(doc, "lei_seca", rotulos["lei_seca"], lei)
        if topico["jurisprudencia"] and _parte("jurisprudencia"):
            _caixa_docx(doc, "jurisprudencia", rotulos["jurisprudencia"], lambda c, t=topico, e=extras: _lista_docx(
                c, [_item_juris_docx(j) for j in t["jurisprudencia"]], e))
        if topico["pegadinhas"] and _parte("pegadinhas"):
            _caixa_docx(doc, "pegadinha", "Pegadinhas", lambda c, t=topico, e=extras: _lista_docx(c, t["pegadinhas"], e))
        if topico["exemplo"] and _parte("exemplo"):
            _caixa_docx(doc, "exemplo", "Exemplo",
                        lambda c, t=topico, e=extras: _texto_com_destaques(c.add_paragraph(), t["exemplo"], e))
        if topico["teste_de_estresse"]:
            _caixa_docx(doc, "estresse", "Acrescentado no teste de estresse",
                        lambda c, t=topico, e=extras: _lista_docx(c, t["teste_de_estresse"], e))
        for ponto in topico["pontos_chave"]:
            _texto_com_destaques(doc.add_paragraph(style="List Bullet"), ponto, extras)
        if topico["fonte"]:
            _texto_com_destaques(doc.add_paragraph(), f"Fonte: {topico['fonte']}", cor=_cores()["suave"])
    prazos = quadro_de_prazos(ficha) if _parte("prazos") else []
    if prazos:
        doc.add_heading(f"Quadro geral: {rotulos['prazos'].lower()}", level=1)
        tabela = doc.add_table(rows=1, cols=4)
        tabela.style = "Table Grid"
        for celula, texto in zip(tabela.rows[0].cells, ("Tópico", "O quê", "Prazo / efeito", "Fundamento")):
            celula.text = texto
        for p in prazos:
            for celula, texto in zip(tabela.add_row().cells, (p["topico"], p["item"], p["valor"], p.get("fundamento") or "")):
                celula.text = texto
    if ficha.get("lacunas_gerais"):
        _caixa_docx(doc, "estresse", "Acrescentado no teste de estresse (geral)", lambda c: _lista_docx(c, ficha["lacunas_gerais"]))
    questoes = ficha.get("questoes") or []
    if questoes:
        doc.add_heading("Teste de estresse", level=1)
        doc.add_paragraph("Responda sem olhar a ficha. Errou? A questão diz qual tópico tinha a resposta: volte nele.")
        for n, q in enumerate(questoes, start=1):
            titulo = f"Questão {n}" + (f" (tópico: {q['topico']})" if q.get("topico") else "")
            doc.add_paragraph().add_run(titulo).bold = True
            _texto_com_destaques(doc.add_paragraph(), q["enunciado"])
            for letra, alternativa in zip(LETRAS, q["alternativas"]):
                _texto_com_destaques(doc.add_paragraph(), f"{letra}) {alternativa}")
        doc.add_heading("Gabarito comentado", level=1)
        for n, q in enumerate(questoes, start=1):
            _texto_com_destaques(doc.add_paragraph(), f"{n}. {q['correta']}: {q.get('explicacao') or ''}")
    if ficha.get("flashcards"):
        doc.add_heading("Flashcards", level=1)
        tabela = doc.add_table(rows=1, cols=2)
        tabela.style = "Table Grid"
        tabela.rows[0].cells[0].text, tabela.rows[0].cells[1].text = "Pergunta", "Resposta"
        for cartao in ficha["flashcards"]:
            celulas = tabela.add_row().cells
            celulas[0].text, celulas[1].text = cartao["frente"], cartao["verso"]
    if ficha.get("glossario") and _parte("glossario"):
        doc.add_heading("Glossário", level=1)
        for item in ficha["glossario"]:
            paragrafo = doc.add_paragraph()
            paragrafo.add_run(f"{item['termo']}: ").bold = True
            _texto_com_destaques(paragrafo, item["definicao"])
    if ficha.get("fora_das_fontes"):
        doc.add_heading("Complementado além das suas fontes", level=1)
        for item in ficha["fora_das_fontes"]:
            doc.add_paragraph(item, style="List Bullet")
    fontes = _textos(ficha.get("fontes"))
    if fontes:
        doc.add_heading("Fontes", level=1)
        for fonte in fontes:
            doc.add_paragraph(fonte, style="List Number")
    doc.save(caminho)
    return caminho
