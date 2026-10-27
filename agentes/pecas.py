"""Peças processuais: petição inicial, contestação, recursos, embargos, HC, MS, trabalhistas e penais.

Pedido do usuário (out/2026): "a capacidade de criar peças processuais". Fluxo do `AgentePeca` (escolha do usuário: os
fatos podem ir para o Claude, com o Gemini de reserva, e por último o Ollama):
1. **Análise** (JSON): a peça cabível (se o pedido não disser qual, a IA escolhe e diz por quê, como na 2ª fase da OAB),
   o juízo, as partes, os fatos, as teses com os fundamentos (norma e artigo), os pedidos e o que falta saber.
2. **Conferência**: cada artigo dos fundamentos é procurado no vade-mécum do PC (ou no Planalto) e cada súmula no banco
   de súmulas (`ficha_estudo.conferir_lei_seca` e `conferir_jurisprudencia`); o que não existe sai da redação, e o texto
   oficial do que existe vai junto, para a citação sair certa.
3. **Redação** (JSON): a peça na estrutura do tipo (`TIPOS_DE_PECA`), com interposição e razões separadas nos recursos.
4. **Depois**: os artigos citados no texto final que não foram conferidos são procurados também; o que não existe fica
   marcado "[CONFERIR O NÚMERO]". O Word (Times 12, justificado, recuo, margens de 3 e 2 cm) e o PDF saem iguais; os dados
   que faltam ficam entre colchetes, em amarelo no Word.

Modo (janela de estudos, `pecas.modo`): "oab" (prova: sem identificar ninguém, "Nome…", "Advogado…, OAB…", "Local…,
Data…") ou "real" (os dados do pedido; o resto entre colchetes). Nada aqui inventa nomes, documentos ou números de processo.
"""
import html
import json
import logging
import os
import re
from datetime import datetime

from agentes import ficha_estudo, geradores
from agentes.base import Agente, Tarefa, normalizar

logger = logging.getLogger(__name__)

# Cada peça: o nome, a área, como reconhecer no pedido (sem acentos), a estrutura em ordem, os fundamentos de cabimento e
# os cuidados que mais derrubam a peça. "recurso": True separa a interposição (ao juízo de origem) das razões.
TIPOS_DE_PECA = {
    "peticao_inicial": {
        "nome": "Petição inicial (procedimento comum)", "area": "cível",
        "reconhecer": r"peticao inicial|\binicial\b|acao (de )?(indenizac|cobranc|obrigac|declarat|revisional|repetic|rescis)\w*",
        "estrutura": ["Endereçamento ao juízo competente", "Qualificação do autor (art. 319, II, CPC)",
                      "Nome da ação e fundamento legal", "Qualificação do réu", "Dos fatos", "Do direito",
                      "Da tutela de urgência (art. 300, CPC), se cabível", "Dos pedidos (certos e determinados, arts. 322 e 324, CPC)",
                      "Das provas (art. 319, VI, CPC)", "Da audiência de conciliação ou mediação (art. 319, VII, CPC)",
                      "Do valor da causa (arts. 291 e 292, CPC)", "Fechamento"],
        "fundamentos": ["art. 319, CPC", "art. 320, CPC", "art. 292, CPC"],
        "cuidados": "Pedidos certos e determinados, cada um com a causa de pedir; valor da causa pela soma dos pedidos; "
                    "gratuidade (arts. 98 e 99, CPC) só se o caso indicar.",
    },
    "contestacao": {
        "nome": "Contestação", "area": "cível", "reconhecer": r"contestac\w*",
        "estrutura": ["Endereçamento ao juízo da causa, com o número do processo", "Qualificação do réu",
                      "Síntese da petição inicial", "Das preliminares (art. 337, CPC)", "Da prejudicial de mérito (prescrição/decadência), se houver",
                      "Do mérito, com impugnação específica de cada fato (art. 341, CPC)", "Da reconvenção (art. 343, CPC), se cabível",
                      "Dos pedidos", "Das provas", "Fechamento"],
        "fundamentos": ["art. 335, CPC", "art. 336, CPC", "art. 337, CPC", "art. 341, CPC"],
        "cuidados": "Prazo de 15 dias úteis (art. 335); toda a defesa de uma vez (art. 336, eventualidade); pedidos: acolher as "
                    "preliminares, julgar improcedente e condenar em honorários (art. 85, CPC).",
    },
    "replica": {
        "nome": "Réplica (impugnação à contestação)", "area": "cível", "reconhecer": r"\breplica\b|impugnac\w* a contestac\w*",
        "estrutura": ["Endereçamento ao juízo da causa, com o número do processo", "Qualificação resumida do autor",
                      "Síntese da contestação", "Das preliminares arguidas", "Do mérito", "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 350, CPC", "art. 351, CPC"],
        "cuidados": "Prazo de 15 dias úteis; rebater cada preliminar e cada fato novo trazido pelo réu.",
    },
    "apelacao": {
        "nome": "Apelação (cível)", "area": "cível", "recurso": True, "reconhecer": r"apelac\w*",
        "estrutura": ["Interposição ao juízo que deu a sentença: partes, processo, cabimento (art. 1.009, CPC), tempestividade, "
                      "preparo (art. 1.007, CPC) e pedido de remessa ao tribunal",
                      "Razões ao tribunal: endereçamento (Egrégio Tribunal, Colenda Câmara)", "Da tempestividade e do preparo",
                      "Síntese do processo e da sentença", "Das preliminares (nulidades), se houver",
                      "Das razões para a reforma da sentença", "Do pedido de efeito suspensivo, se for o caso (art. 1.012, CPC)",
                      "Do pedido: conhecimento e provimento para reformar (ou anular) a sentença", "Fechamento"],
        "fundamentos": ["art. 1.009, CPC", "art. 1.010, CPC", "art. 1.003, CPC", "art. 1.007, CPC"],
        "cuidados": "Prazo de 15 dias úteis (art. 1.003, § 5º); a interposição vai ao juízo de primeiro grau e as razões ao "
                    "tribunal; requisitos do art. 1.010 (fato, direito, razões e pedido de nova decisão).",
    },
    "apelacao_criminal": {
        "nome": "Apelação criminal", "area": "penal", "recurso": True,
        "reconhecer": r"apelac\w* (criminal|penal|no processo penal|da sentenca condenatoria)",
        "estrutura": ["Interposição ao juiz da sentença (art. 593, CPP), com a tempestividade", "Razões ao tribunal (art. 600, CPP)",
                      "Síntese da acusação e da sentença", "Das preliminares (nulidades)", "Do mérito: absolvição (art. 386, CPP)",
                      "Subsidiariamente: dosimetria (arts. 59 e 68, CP), regime (art. 33, CP), substituição (art. 44, CP)",
                      "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 593, CPP", "art. 600, CPP", "art. 386, CPP"],
        "cuidados": "Prazo de 5 dias para interpor (art. 593) e 8 dias para as razões (art. 600).",
    },
    "agravo_de_instrumento": {
        "nome": "Agravo de instrumento", "area": "cível", "reconhecer": r"agravo de instrumento",
        "estrutura": ["Endereçamento direto ao tribunal (Desembargador Presidente do Tribunal de Justiça)",
                      "Qualificação do agravante e do agravado, com o processo de origem",
                      "Interposição contra a decisão interlocutória, com pedido de efeito suspensivo ou tutela recursal",
                      "Do cabimento (art. 1.015, CPC, com o inciso) e da tempestividade",
                      "Das peças obrigatórias (art. 1.017, CPC)", "Dos nomes e endereços dos advogados (art. 1.016, IV, CPC)",
                      "Síntese da decisão agravada", "Das razões do pedido de reforma",
                      "Do efeito suspensivo ou da antecipação da tutela recursal (art. 1.019, I, CPC)", "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 1.015, CPC", "art. 1.016, CPC", "art. 1.017, CPC", "art. 1.019, CPC", "art. 1.003, CPC"],
        "cuidados": "Vai direto ao tribunal; prazo de 15 dias úteis; dizer o inciso do art. 1.015; pedir a intimação do "
                    "agravado para responder (art. 1.019, II).",
    },
    "agravo_interno": {
        "nome": "Agravo interno", "area": "cível", "reconhecer": r"agravo interno",
        "estrutura": ["Endereçamento ao relator da decisão", "Qualificação e processo", "Do cabimento (art. 1.021, CPC) e da tempestividade",
                      "Síntese da decisão monocrática", "Da impugnação específica dos fundamentos (art. 1.021, § 1º, CPC)",
                      "Dos pedidos: retratação ou julgamento pelo colegiado", "Fechamento"],
        "fundamentos": ["art. 1.021, CPC"],
        "cuidados": "Prazo de 15 dias úteis; impugnar especificamente cada fundamento da decisão.",
    },
    "embargos_de_declaracao": {
        "nome": "Embargos de declaração", "area": "cível", "reconhecer": r"embargos (de declarac\w*|declaratorios)",
        "estrutura": ["Endereçamento ao juízo (ou relator) que deu a decisão", "Qualificação e processo",
                      "Da tempestividade (5 dias, art. 1.023, CPC)", "Da decisão embargada",
                      "Do vício: omissão, contradição, obscuridade ou erro material (art. 1.022, CPC)",
                      "Do efeito modificativo, se houver (art. 1.023, § 2º, CPC)", "Do prequestionamento, se for o caso (art. 1.025, CPC)",
                      "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 1.022, CPC", "art. 1.023, CPC", "art. 1.026, CPC"],
        "cuidados": "Sem preparo; interrompem o prazo dos outros recursos (art. 1.026); apontar o vício com precisão.",
    },
    "recurso_especial": {
        "nome": "Recurso especial", "area": "cível", "recurso": True, "reconhecer": r"recurso especial|\bresp\b",
        "estrutura": ["Interposição ao Presidente ou Vice-Presidente do tribunal de origem (art. 1.029, CPC)",
                      "Razões ao Superior Tribunal de Justiça", "Da tempestividade e do preparo",
                      "Do cabimento (art. 105, III, CF, com a alínea)", "Do prequestionamento",
                      "Da relevância da questão de direito federal (art. 105, § 2º, CF)", "Síntese do caso e do acórdão",
                      "Das razões: violação da lei federal ou divergência (com cotejo analítico, art. 1.029, § 1º, CPC)",
                      "Do pedido", "Fechamento"],
        "fundamentos": ["art. 105, CF", "art. 1.029, CPC", "art. 1.003, CPC"],
        "cuidados": "Só questão de direito (Súmula 7 do STJ); prequestionamento; 15 dias úteis.",
    },
    "recurso_extraordinario": {
        "nome": "Recurso extraordinário", "area": "constitucional", "recurso": True, "reconhecer": r"recurso extraordinario",
        "estrutura": ["Interposição ao Presidente ou Vice-Presidente do tribunal de origem (art. 1.029, CPC)",
                      "Razões ao Supremo Tribunal Federal", "Da tempestividade e do preparo", "Do cabimento (art. 102, III, CF, com a alínea)",
                      "Do prequestionamento", "Da repercussão geral (art. 102, § 3º, CF; art. 1.035, CPC)", "Síntese do caso e do acórdão",
                      "Das razões: a ofensa à Constituição", "Do pedido", "Fechamento"],
        "fundamentos": ["art. 102, CF", "art. 1.029, CPC", "art. 1.035, CPC"],
        "cuidados": "Preliminar formal e fundamentada de repercussão geral; ofensa direta à Constituição.",
    },
    "recurso_inominado": {
        "nome": "Recurso inominado (juizado especial)", "area": "cível", "recurso": True, "reconhecer": r"recurso inominado",
        "estrutura": ["Interposição ao juiz do juizado (art. 41, Lei 9.099/1995)", "Razões à Turma Recursal",
                      "Da tempestividade (10 dias) e do preparo (art. 42, Lei 9.099/1995)", "Síntese da sentença",
                      "Das razões para a reforma", "Do pedido", "Fechamento"],
        "fundamentos": ["art. 41, Lei 9.099/1995", "art. 42, Lei 9.099/1995"],
        "cuidados": "Prazo de 10 dias; preparo em até 48 horas da interposição; assinado por advogado.",
    },
    "mandado_de_seguranca": {
        "nome": "Mandado de segurança", "area": "constitucional", "reconhecer": r"mandado de seguranca",
        "estrutura": ["Endereçamento ao juízo competente pela autoridade coatora", "Qualificação do impetrante",
                      "Impetração contra ato da autoridade coatora, com a pessoa jurídica a que ela se vincula (art. 6º, Lei 12.016/2009)",
                      "Dos fatos", "Do cabimento e da tempestividade (120 dias, art. 23, Lei 12.016/2009)",
                      "Do direito líquido e certo (prova pré-constituída)", "Da liminar (art. 7º, III, Lei 12.016/2009)",
                      "Dos pedidos: liminar, notificação da autoridade, ciência ao órgão de representação, oitiva do Ministério "
                      "Público e concessão da segurança", "Do valor da causa", "Fechamento"],
        "fundamentos": ["art. 5º, CF", "art. 1º, Lei 12.016/2009", "art. 7º, Lei 12.016/2009", "art. 23, Lei 12.016/2009"],
        "cuidados": "Sem dilação probatória; indicar a autoridade coatora certa; prazo decadencial de 120 dias.",
    },
    "habeas_corpus": {
        "nome": "Habeas corpus", "area": "penal", "reconhecer": r"habeas corpus|\bhc\b",
        "estrutura": ["Endereçamento à autoridade acima da coatora", "Impetrante, paciente e autoridade coatora (art. 654, CPP)",
                      "Dos fatos", "Do constrangimento ilegal (art. 648, CPP, com o inciso)", "Da liminar",
                      "Dos pedidos: liminar, informações da autoridade, parecer do Ministério Público e concessão da ordem "
                      "(alvará de soltura ou salvo-conduto)", "Fechamento"],
        "fundamentos": ["art. 5º, CF", "art. 647, CPP", "art. 648, CPP", "art. 654, CPP"],
        "cuidados": "Qualquer pessoa pode impetrar; apontar a hipótese do art. 648; não cabe para discutir prova a fundo.",
    },
    "liberdade_provisoria": {
        "nome": "Pedido de liberdade provisória ou revogação da prisão", "area": "penal",
        "reconhecer": r"liberdade provisoria|relaxamento (de|da) prisao|revogac\w* (de|da) (prisao )?preventiva",
        "estrutura": ["Endereçamento ao juízo criminal", "Qualificação do requerente e do processo", "Dos fatos",
                      "Da ausência dos requisitos da prisão preventiva (art. 312, CPP) ou da ilegalidade da prisão",
                      "Das medidas cautelares diversas da prisão (art. 319, CPP)", "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 5º, CF", "art. 310, CPP", "art. 312, CPP", "art. 316, CPP", "art. 319, CPP", "art. 321, CPP"],
        "cuidados": "Prisão ilegal: relaxamento; legal mas desnecessária: liberdade provisória ou revogação da preventiva.",
    },
    "resposta_a_acusacao": {
        "nome": "Resposta à acusação", "area": "penal", "reconhecer": r"resposta a acusac\w*",
        "estrutura": ["Endereçamento ao juízo criminal, com o processo", "Qualificação do acusado",
                      "Apresentação da resposta (arts. 396 e 396-A, CPP)", "Síntese da denúncia",
                      "Das preliminares (nulidades, rejeição da denúncia, art. 395, CPP)", "Da absolvição sumária (art. 397, CPP)",
                      "Dos pedidos", "Do rol de testemunhas (até 8, art. 401, CPP)", "Fechamento"],
        "fundamentos": ["art. 396, CPP", "art. 396-A, CPP", "art. 397, CPP"],
        "cuidados": "Prazo de 10 dias; é o momento do rol de testemunhas; pedir absolvição sumária pelo inciso certo.",
    },
    "memoriais": {
        "nome": "Memoriais (alegações finais por escrito)", "area": "penal", "reconhecer": r"memoriais|alegac\w* finais",
        "estrutura": ["Endereçamento ao juízo, com o processo", "Qualificação do acusado",
                      "Apresentação dos memoriais (art. 403, § 3º, CPP)", "Síntese do processo", "Das preliminares (nulidades)",
                      "Do mérito: absolvição (art. 386, CPP, com o inciso)",
                      "Subsidiariamente: desclassificação, pena mínima (art. 59, CP), atenuantes, regime (art. 33, CP), "
                      "substituição (art. 44, CP), sursis (art. 77, CP)", "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 403, CPP", "art. 386, CPP"],
        "cuidados": "Prazo de 5 dias; pedidos subsidiários em ordem, do mais amplo ao mais restrito.",
    },
    "queixa_crime": {
        "nome": "Queixa-crime", "area": "penal", "reconhecer": r"queixa.crime",
        "estrutura": ["Endereçamento ao juízo criminal competente", "Qualificação do querelante, por procurador com poderes "
                      "especiais (art. 44, CPP)", "Qualificação do querelado", "Dos fatos, com todas as circunstâncias (art. 41, CPP)",
                      "Do direito: a tipificação", "Dos pedidos: recebimento, citação, condenação e valor mínimo de reparação "
                      "(art. 387, IV, CPP)", "Do rol de testemunhas", "Fechamento"],
        "fundamentos": ["art. 30, CPP", "art. 41, CPP", "art. 44, CPP", "art. 38, CPP"],
        "cuidados": "Decadência de 6 meses (art. 38, CPP); procuração com poderes especiais e menção ao fato.",
    },
    "reclamacao_trabalhista": {
        "nome": "Reclamação trabalhista", "area": "trabalhista", "reconhecer": r"reclamac\w* trabalhista|reclamatoria",
        "estrutura": ["Endereçamento à Vara do Trabalho", "Qualificação do reclamante (com CTPS e PIS)",
                      "Qualificação da reclamada (CNPJ)", "Do contrato de trabalho", "Da justiça gratuita (art. 790, §§ 3º e 4º, CLT)",
                      "Dos fatos e fundamentos de cada pedido", "Dos pedidos (líquidos, com o valor de cada um, art. 840, § 1º, CLT)",
                      "Dos honorários de sucumbência (art. 791-A, CLT)", "Do valor da causa", "Das provas", "Fechamento"],
        "fundamentos": ["art. 840, CLT", "art. 7º, CF", "art. 791-A, CLT"],
        "cuidados": "Pedido sem valor é extinto sem resolução do mérito (art. 840, § 3º); prescrição (art. 7º, XXIX, CF; art. 11, CLT).",
    },
    "contestacao_trabalhista": {
        "nome": "Contestação trabalhista", "area": "trabalhista", "reconhecer": r"contestac\w* trabalhista|defesa trabalhista",
        "estrutura": ["Endereçamento à Vara do Trabalho, com o processo", "Qualificação da reclamada",
                      "Das preliminares", "Da prescrição (art. 7º, XXIX, CF; art. 11, CLT)", "Do mérito, pedido a pedido",
                      "Da compensação (art. 767, CLT)", "Dos pedidos", "Das provas", "Fechamento"],
        "fundamentos": ["art. 847, CLT", "art. 818, CLT", "art. 767, CLT", "art. 11, CLT"],
        "cuidados": "Impugnar cada pedido e os valores; ônus da prova (art. 818).",
    },
    "recurso_ordinario_trabalhista": {
        "nome": "Recurso ordinário trabalhista", "area": "trabalhista", "recurso": True, "reconhecer": r"recurso ordinario",
        "estrutura": ["Interposição ao juiz da Vara do Trabalho (art. 895, I, CLT)", "Razões ao Tribunal Regional do Trabalho",
                      "Da tempestividade (8 dias), do depósito recursal (art. 899, CLT) e das custas (art. 789, CLT)",
                      "Síntese da sentença", "Das razões para a reforma", "Do pedido", "Fechamento"],
        "fundamentos": ["art. 895, CLT", "art. 899, CLT", "art. 789, CLT"],
        "cuidados": "Prazo de 8 dias; depósito recursal e custas para a empresa recorrente.",
    },
    "embargos_a_execucao": {
        "nome": "Embargos à execução", "area": "cível", "reconhecer": r"embargos (a|na) execuc\w*|embargos do devedor",
        "estrutura": ["Endereçamento ao juízo da execução (distribuição por dependência, art. 914, § 1º, CPC)",
                      "Qualificação do embargante e do embargado", "Da tempestividade (15 dias, art. 915, CPC)",
                      "Síntese da execução", "Das matérias de defesa (art. 917, CPC)",
                      "Do excesso de execução, com o valor correto e o demonstrativo (art. 917, § 3º, CPC)",
                      "Do efeito suspensivo (art. 919, § 1º, CPC)", "Dos pedidos", "Do valor da causa", "Fechamento"],
        "fundamentos": ["art. 914, CPC", "art. 915, CPC", "art. 917, CPC", "art. 919, CPC"],
        "cuidados": "Não depende de penhora; excesso sem o valor correto é rejeitado.",
    },
    "impugnacao_ao_cumprimento": {
        "nome": "Impugnação ao cumprimento de sentença", "area": "cível", "reconhecer": r"impugnac\w* ao cumprimento",
        "estrutura": ["Endereçamento ao juízo do cumprimento, com o processo", "Qualificação do executado",
                      "Da tempestividade (15 dias depois do prazo do art. 523, CPC)", "Das matérias (art. 525, § 1º, CPC)",
                      "Do excesso, com o valor correto e o demonstrativo (art. 525, § 4º, CPC)",
                      "Do efeito suspensivo (art. 525, § 6º, CPC)", "Dos pedidos", "Fechamento"],
        "fundamentos": ["art. 525, CPC", "art. 523, CPC"],
        "cuidados": "Só as matérias do § 1º; o excesso pede o valor certo.",
    },
    "acao_de_alimentos": {
        "nome": "Ação de alimentos", "area": "família", "reconhecer": r"(acao|pedido) de alimentos",
        "estrutura": ["Endereçamento à Vara de Família", "Qualificação do alimentando, representado (se menor)",
                      "Qualificação do alimentante", "Dos fatos e do vínculo", "Do direito: necessidade e possibilidade "
                      "(art. 1.694, § 1º, CC)", "Dos alimentos provisórios (art. 4º, Lei 5.478/1968)", "Dos pedidos",
                      "Do valor da causa (12 prestações, art. 292, III, CPC)", "Fechamento"],
        "fundamentos": ["art. 1.694, CC", "art. 1.696, CC", "art. 2º, Lei 5.478/1968", "art. 4º, Lei 5.478/1968"],
        "cuidados": "Intervenção do Ministério Público quando houver incapaz; valor da causa pelas 12 prestações.",
    },
}
# Ordem de reconhecimento: as específicas antes das gerais ("contestação trabalhista" antes de "contestação").
ORDEM_DE_RECONHECIMENTO = [
    "contestacao_trabalhista", "apelacao_criminal", "recurso_ordinario_trabalhista", "impugnacao_ao_cumprimento",
    "embargos_a_execucao", "embargos_de_declaracao", "agravo_de_instrumento", "agravo_interno", "recurso_especial",
    "recurso_extraordinario", "recurso_inominado", "mandado_de_seguranca", "habeas_corpus", "liberdade_provisoria",
    "resposta_a_acusacao", "memoriais", "queixa_crime", "reclamacao_trabalhista", "acao_de_alimentos", "replica",
    "contestacao", "apelacao", "peticao_inicial",
]
# "Faz a peça desse caso", "redige a peça cabível", "uma petição": peça sem o tipo (a IA escolhe a cabível).
RE_PECA_GENERICA = re.compile(r"\b(pecas?( processua\w*| cabive\w*| correspondente| juridica| pratica)?|peticao|peticionar)\b")
# Verbo de criar ou de peticionar: "faz", "redige", "elabora", "interpõe", "impetra", "ajuíza".
RE_FAZER_PECA = re.compile(
    r"\b(crie|cria|criar|faca|faz|fazer|gere|gera|gerar|monte|monta|montar|prepare|prepara|preparar|elabore|elabora|elaborar|"
    r"escreva|escreve|escrever|produza|produz|produzir|redija|redige|redigir|minut\w*|desenvolv\w*|quero|preciso|"
    r"interp(or|onha|oe)|impetr\w*|ajuiz\w*|propo(r|nha)|opo(r|nha))\b")
# Estudar a peça não é redigir a peça: "ficha de agravo de instrumento", "ensino de peça pra eu estudar", "o que é apelação".
RE_ESTUDO_DA_PECA = re.compile(
    r"\b(estud\w*|ensin\w*|resum\w*|fichas?|flashcards?|explic\w*|aulas?|curso|material|revis\w*|questoes|simulado|"
    r"o que e|como funciona|quando cabe|para que serve|pra que serve|prazo d[aeo]|apresentac\w*|slides?|power ?point|planilhas?)\b"
    # "um documento sobre petição inicial", "PDF a respeito do agravo": a peça é o assunto, não a entrega.
    r"|\b(sobre|a respeito d\w+|acerca d\w+)\b.{0,30}\b(pecas?|peticao|recursos?|agravo|apelac\w*|contestac\w*|embargos|"
    r"habeas|mandado|reclamac\w*|queixa|memoriais)\b")


def tipo_no_texto(texto_normalizado: str) -> str | None:
    """O tipo de peça citado ("redige um agravo de instrumento" -> "agravo_de_instrumento"); None se não houver."""
    for chave in ORDEM_DE_RECONHECIMENTO:
        if re.search(TIPOS_DE_PECA[chave]["reconhecer"], texto_normalizado):
            return chave
    return None


def pedido_de_peca(comando: str):
    """(tipo, True) para "faz uma petição inicial de…" ("" = a peça cabível); None quando não é pedido de peça."""
    texto = normalizar(comando)
    if not RE_FAZER_PECA.search(texto) or RE_ESTUDO_DA_PECA.search(texto):
        return None
    tipo = tipo_no_texto(texto)
    if tipo is None and not RE_PECA_GENERICA.search(texto):
        return None
    # "apelação" sozinho, num pedido que fala de crime, é a criminal.
    if tipo == "apelacao" and re.search(r"\b(crime|criminal|penal|reu|denuncia|condenad\w*|pena)\b", texto):
        tipo = "apelacao_criminal"
    return tipo or ""


def nome_da_peca(tipo: str) -> str:
    return TIPOS_DE_PECA[tipo]["nome"] if tipo in TIPOS_DE_PECA else "peça processual"


def catalogo_para_o_prompt() -> str:
    return "\n".join(f"- {chave}: {tipo['nome']} ({tipo['area']})" for chave, tipo in TIPOS_DE_PECA.items())


# ------------------------------------------------------------------ citações
# "art. 319 do CPC", "arts. 186 e 927 do Código Civil", "artigo 5º, X, da CF", "art. 1.015, I, do CPC": os números e, no
# trecho até a próxima citação, a lei.
RE_ARTIGOS = re.compile(
    r"\b(?:arts?\.?|artigos?)\s*((?:\d{1,2}\.\d{3}|\d{1,4})[ºo°]?(?:-[A-Z])?"
    r"(?:\s*(?:,|e)\s*(?:\d{1,2}\.\d{3}|\d{1,4})[ºo°]?(?:-[A-Z])?)*)", re.IGNORECASE)
RE_NUMERO = re.compile(r"(\d{1,2}\.\d{3}|\d{1,4})[ºo°]?(?:-([A-Z]))?")


def _citacoes(texto: str):
    """(achado, norma, [artigos]) de cada citação com a lei; a lei vem do trecho até a próxima citação (70 letras no máximo)."""
    from agente_leis import norma_por_texto

    achados = list(RE_ARTIGOS.finditer(texto or ""))
    for indice, achado in enumerate(achados):
        fim = achados[indice + 1].start() if indice + 1 < len(achados) else len(texto)
        resto = re.split(r"\.\s|;|\n", texto[achado.end():min(fim, achado.end() + 70)])[0]
        norma = norma_por_texto(resto)
        if norma is None:
            continue
        artigos = [n.group(1) + (f"-{n.group(2).upper()}" if n.group(2) else "") for n in RE_NUMERO.finditer(achado.group(1))]
        yield achado, norma, resto, artigos


def _sem_ligacao_no_fim(resto: str) -> str:
    """', I, do CPC e nos' -> ', I, do CPC' (as palavras de ligação que levam à próxima citação saem)."""
    return re.sub(r"(?:[\s,.]+(?:e|do|da|dos|das|no|na|nos|nas|al[ée]m|bem|como|c/c|combinado com))+[\s,.]*$", "",
                  resto, flags=re.IGNORECASE).rstrip(" ,.")


def citacoes_do_texto(texto: str) -> list[dict]:
    """[{"norma", "artigo", "referencia"}] dos artigos citados com a lei ("art. 319 do CPC"); os sem lei ficam de fora."""
    unicos = {}
    for _, norma, resto, artigos in _citacoes(texto):
        for artigo in artigos:
            unicos.setdefault((norma.chave, artigo), {
                "norma": norma.chave, "artigo": artigo, "referencia": f"art. {artigo}{_sem_ligacao_no_fim(resto)}"[:90]})
    return list(unicos.values())


def conferir_fundamentos(fundamentos: list[dict], jurisprudencia: list[dict], conferir_lei=None, conferir_sumulas=None) -> dict:
    """Confere os artigos (vade-mécum do PC ou Planalto) e as súmulas (banco de súmulas). Marca cada item como a ficha de
    estudo marca ("conferido" True, False ou None) e devolve a contagem."""
    ficha = {"topicos": [{"lei_seca": fundamentos, "jurisprudencia": jurisprudencia}]}
    contagem = {}
    if fundamentos and conferir_lei:
        try:
            contagem.update(conferir_lei(ficha) or {})
        except Exception as exc:
            logger.warning("Não deu para conferir os artigos da peça: %s", exc)
    if jurisprudencia and conferir_sumulas:
        try:
            contagem.update(conferir_sumulas(ficha) or {})
        except Exception as exc:
            logger.warning("Não deu para conferir as súmulas da peça: %s", exc)
    return contagem


# ------------------------------------------------------------------ o agente
REGRAS_GERAIS = (
    "Regras: nunca invente nomes, CPF, CNPJ, endereços, números de processo, datas ou valores que não estejam no caso; o que "
    "faltar fica entre colchetes, em maiúsculas ([NOME DO AUTOR], [Nº DO PROCESSO], [VALOR]). Cite só artigos e súmulas "
    "que existem; prefira o fundamento legal à doutrina. Linguagem forense clara, sem rebuscamento desnecessário. Português do "
    "Brasil."
)
MODO_OAB = ("MODO PROVA DA OAB (2ª fase): não identifique ninguém além do que o enunciado dá; use 'Nome…', 'nacionalidade…', "
            "'estado civil…', 'profissão…', 'CPF…', 'endereço…' na qualificação que o caso não trouxer; feche com 'Local…, "
            "Data…' e 'Advogado…, OAB/UF nº…'. Nada de inventar dados para preencher.")
MODO_REAL = ("MODO CASO REAL: use os dados que o caso traz; o que faltar fica entre colchetes. Feche com a cidade e a data "
             "dadas abaixo (ou [CIDADE], [DATA]) e com o advogado dado abaixo (ou [ADVOGADO], OAB/[UF] nº [NÚMERO]).")


class AgentePeca(Agente):
    nome = "peca"
    papel = "Redige a peça processual na estrutura certa, com os artigos conferidos no vade-mécum, em Word e PDF."
    instrucoes = (
        "Você é o Agente de Peças da Beatriz, advogada experiente e examinadora da 2ª fase da OAB. Nesta etapa você ANALISA "
        "o caso antes de redigir. Responda SOMENTE com um objeto JSON neste formato:\n"
        '{"tipo": str, "por_que": str, "juizo": str, "processo": str, "decisao_atacada": str, '
        '"partes": [{"papel": str, "nome": str, "qualificacao": str}], "fatos": [str], '
        '"teses": [{"titulo": str, "argumento": str, "fundamentos": [{"norma": str, "artigo": str, "referencia": str}]}], '
        '"jurisprudencia": [{"referencia": str, "tese": str}], "pedidos": [str], "provas": [str], "valor_da_causa": str, '
        '"prazo": str, "faltando": [str]}\n'
        "'tipo' é a chave de uma das peças da lista abaixo: a pedida, ou, se o pedido não disser, a peça CABÍVEL no caso (é "
        "isso que a banca cobra: confira o momento processual, a decisão atacada e quem você representa) e 'por_que' explica "
        "a escolha em uma ou duas frases; 'juizo' é o endereçamento completo; 'norma' é a sigla (CF, CC, CPC, CP, CPP, CLT, "
        "CDC, ECA) ou a lei com número ('Lei 9.099/1995'); 'artigo' só o número ('319' ou '1.015'); 'referencia' como se "
        "cita ('art. 1.015, I, CPC'); 'faltando' lista os dados que o caso não traz e que a peça precisa.\n"
    )
    instrucoes_redacao = (
        "Você é o Agente de Peças da Beatriz, advogada experiente. Redija a peça completa a partir da análise e dos textos "
        "oficiais dos artigos dados. Responda SOMENTE com um objeto JSON neste formato:\n"
        '{"documentos": [{"parte": "interposicao", "razoes" ou "peca", "enderecamento": str, "processo": str, '
        '"titulo": str, "preambulo": str, "secoes": [{"titulo": str, "paragrafos": [str], "itens": [str]}], '
        '"fechamento": [str]}], "valor_da_causa": str}\n'
        "Regras da forma: 'enderecamento' em maiúsculas ('EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO DA … VARA CÍVEL DA "
        "COMARCA DE …'); 'titulo' é o nome da peça em maiúsculas ('AÇÃO DE INDENIZAÇÃO POR DANOS MORAIS', 'AGRAVO DE "
        "INSTRUMENTO'); 'preambulo' é o parágrafo de abertura (quem, qualificado, por seu advogado, vem propor/interpor/"
        "apresentar o quê, com qual fundamento, contra quem); as seções seguem a ESTRUTURA dada, com títulos numerados em "
        "romanos e maiúsculas ('I – DOS FATOS'); parágrafos completos, de 3 a 6 frases, que ligam o fato ao artigo; os "
        "pedidos vão em 'itens', um por item, sem a letra (a numeração é feita depois); 'fechamento' são as linhas finais "
        "('Nestes termos,', 'pede deferimento.', o local e a data, o advogado). Nos recursos com interposição e razões, são dois "
        "documentos: a interposição (curta, ao juízo de origem) e as razões (ao tribunal). Cite os artigos como aparecem nos "
        "textos oficiais. 'processo' é só o número (ou [Nº DO PROCESSO]); 'valor_da_causa' é só o valor ('R$ 10.000,00' ou "
        "[VALOR DA CAUSA]) e fica vazio quando a peça não tem valor da causa (recursos, contestação, réplica, HC).\n"
    )

    def __init__(self, conferir_lei=ficha_estudo.conferir_lei_seca, conferir_sumulas=ficha_estudo.conferir_jurisprudencia,
                 **kwargs):
        super().__init__(**kwargs)
        self.conferir_lei = conferir_lei
        self.conferir_sumulas = conferir_sumulas

    @staticmethod
    def _preferencias(tarefa: Tarefa) -> dict:
        return ((tarefa.contexto.get("preferencias") or {}).get("pecas") or {})

    def _modo(self, tarefa: Tarefa) -> str:
        prefs = self._preferencias(tarefa)
        if prefs.get("modo", "oab") == "oab":
            return MODO_OAB
        assinatura = (f"Advogado: {prefs.get('advogado') or '[ADVOGADO]'}, OAB {prefs.get('oab') or '[UF] nº [NÚMERO]'}. "
                      f"Cidade: {prefs.get('cidade') or '[CIDADE]'}. Data: {datetime.now():%d/%m/%Y}.")
        return f"{MODO_REAL} {assinatura}"

    def _sistema_da(self, instrucoes: str, tarefa: Tarefa) -> str:
        extra = self._preferencias(tarefa).get("instrucoes")
        return self._sistema(f"{instrucoes}{REGRAS_GERAIS}\n{self._modo(tarefa)}"
                             + (f"\nInstruções do usuário para as peças: {extra}" if extra else ""))

    def _caso(self, tarefa: Tarefa) -> str:
        """O caso: os arquivos lidos (enunciado, documentos) e os fatos escritos, antes do pedido."""
        partes = []
        if tarefa.contexto.get("biblioteca"):
            partes.append("ARQUIVOS DO CASO (enunciado e documentos que o usuário mandou):\n" + tarefa.contexto["biblioteca"])
        fatos = (tarefa.contexto.get("fatos") or "").strip()
        if fatos and fatos != tarefa.pedido:
            partes.append(f"FATOS ESCRITOS PELO USUÁRIO:\n{fatos}")
        partes.append(f"PEDIDO DO USUÁRIO: {tarefa.pedido}")
        return "\n\n".join(partes)

    def _analisar(self, tarefa: Tarefa) -> dict:
        pedido = tarefa.contexto.get("tipo_peca") or ""
        escolha = (f"A peça pedida é: {pedido} ({nome_da_peca(pedido)})." if pedido in TIPOS_DE_PECA else
                   "O pedido não diz a peça: escolha a CABÍVEL.")
        analise = self.pedir_json(f"{self._caso(tarefa)}\n\n{escolha}{_bloco_correcao_peca(tarefa)}", max_tokens=6000,
                                  sistema=self._sistema_da(f"{self.instrucoes}PEÇAS:\n{catalogo_para_o_prompt()}\n", tarefa))
        tipo = analise.get("tipo") if analise.get("tipo") in TIPOS_DE_PECA else pedido
        analise["tipo"] = tipo if tipo in TIPOS_DE_PECA else "peticao_inicial"
        return analise

    def _redigir(self, tarefa: Tarefa, analise: dict, conferidos: list, descartados: list) -> dict:
        tipo = TIPOS_DE_PECA[analise["tipo"]]
        estrutura = "\n".join(f"{n}. {parte}" for n, parte in enumerate(tipo["estrutura"], start=1))
        textos = "\n".join(f"- {i['referencia']}: “{i['texto_oficial']}”" for i in conferidos if i.get("texto_oficial"))
        fora = (f"\nNÃO CITE (não existem na lei; a análise errou o número): {', '.join(descartados)}." if descartados else "")
        prompt = (f"{self._caso(tarefa)}\n\nPEÇA: {tipo['nome']}\nESTRUTURA (nesta ordem):\n{estrutura}\n"
                  f"CUIDADOS: {tipo['cuidados']}\n\nANÁLISE:\n{json.dumps(analise, ensure_ascii=False)}\n\n"
                  f"TEXTOS OFICIAIS DOS ARTIGOS CONFERIDOS:\n{textos or '(nenhum conferido: cite com cuidado)'}{fora}"
                  f"{_bloco_correcao_peca(tarefa)}")
        return self.pedir_json(prompt, max_tokens=14000, sistema=self._sistema_da(self.instrucoes_redacao, tarefa))

    def executar(self, tarefa: Tarefa) -> str:
        analise = self._analisar(tarefa)
        tipo = TIPOS_DE_PECA[analise["tipo"]]
        # Os fundamentos da análise e os de cabimento do tipo, conferidos antes de redigir.
        fundamentos = [dict(f) for t in analise.get("teses") or [] if isinstance(t, dict)
                       for f in t.get("fundamentos") or [] if isinstance(f, dict)]
        fundamentos += [{"norma": ref.split(", ")[-1], "artigo": ref.split(",")[0].replace("art.", "").strip(), "referencia": ref}
                        for ref in tipo["fundamentos"]]
        juris = [dict(j) for j in analise.get("jurisprudencia") or [] if isinstance(j, dict)]
        juris = [{**j, "conferir": False} for j in juris]
        contagem = conferir_fundamentos(fundamentos, juris, self.conferir_lei, self.conferir_sumulas)
        conferidos = [f for f in fundamentos if f.get("conferido")]
        descartados = [f.get("referencia") or f.get("artigo") for f in fundamentos if f.get("conferido") is False]
        descartados = [d for d in dict.fromkeys(descartados + [j.get("referencia") for j in juris if j.get("conferido") is False]) if d]
        peca = self._redigir(tarefa, analise, conferidos, descartados)
        documentos = [d for d in peca.get("documentos") or [] if isinstance(d, dict) and (d.get("secoes") or d.get("preambulo"))]
        if not documentos:
            raise ValueError("a redação veio sem a peça")
        peca = {"tipo": analise["tipo"], "nome": tipo["nome"], "documentos": documentos,
                "valor_da_causa": peca.get("valor_da_causa") or analise.get("valor_da_causa") or "",
                "por_que": analise.get("por_que") or "", "faltando": [str(f) for f in analise.get("faltando") or [] if f],
                "prazo": analise.get("prazo") or ""}

        # Depois: os artigos do texto final que não foram conferidos antes; o que não existe fica marcado.
        ja = {(_chave_norma(f.get("norma")), _so_numero(f.get("artigo"))) for f in fundamentos}
        novos = [c for c in citacoes_do_texto(texto_da_peca(peca)) if (c["norma"], _so_numero(c["artigo"])) not in ja]
        depois = conferir_fundamentos(novos, [], self.conferir_lei, None)
        inexistentes = [c for c in novos if c.get("conferido") is False]
        if inexistentes:
            marcar_inexistentes(peca, inexistentes)
        conferencia = {
            "conferidos": contagem.get("conferidos", 0) + depois.get("conferidos", 0),
            "nao_encontrados": len(descartados) + len(inexistentes),
            "sumulas": contagem.get("conferidas", 0), "sumulas_canceladas": contagem.get("canceladas", 0),
            "artigos": [f.get("referencia") for f in conferidos + [c for c in novos if c.get("conferido")]],
            "descartados": descartados, "marcados": [c["referencia"] for c in inexistentes],
        }
        peca["conferencia"] = conferencia
        tarefa.contexto["conteudo"] = peca
        tarefa.tema = tipo["nome"]

        formatos = [f for f in tarefa.formatos if f in ("pdf", "docx")] or ["docx", "pdf"]
        os.makedirs(tarefa.pasta, exist_ok=True)
        with open(os.path.join(tarefa.pasta, "peca.json"), "w", encoding="utf-8") as arquivo:
            json.dump(peca, arquivo, ensure_ascii=False, indent=1)
        nome = geradores.nome_arquivo_seguro(tipo["nome"].split(" (")[0])
        gerar = {"docx": gerar_docx_peca, "pdf": gerar_pdf_peca}
        tarefa.arquivos = [gerar[f](peca, os.path.join(tarefa.pasta, f"{nome}.{f}")) for f in ("docx", "pdf") if f in formatos]

        resumo = f"Redigiu {tipo['nome'].lower()} ({len(documentos)} documento(s))."
        if conferencia["conferidos"]:
            resumo += f" Conferiu {conferencia['conferidos']} artigo(s) no vade-mécum."
        if conferencia["nao_encontrados"]:
            resumo += f" {conferencia['nao_encontrados']} citação(ões) não existem: tiradas ou marcadas para conferir."
        if conferencia["sumulas"]:
            resumo += f" Conferiu {conferencia['sumulas']} súmula(s)."
        return resumo


def _bloco_correcao_peca(tarefa: Tarefa) -> str:
    avaliacao = tarefa.avaliacao
    if not avaliacao or avaliacao.cumprido:
        return ""
    problemas = "\n".join(f"- {p}" for p in avaliacao.problemas)
    return f"\n\nA VERSÃO ANTERIOR FOI REPROVADA PELO REVISOR:\n{problemas}\nCorreção pedida: {avaliacao.correcao}"


def _so_numero(artigo) -> str:
    return re.sub(r"[^\dA-Z-]", "", str(artigo or "").upper().replace(".", ""))


def _chave_norma(norma) -> str:
    from agente_leis import norma_por_texto

    achada = norma_por_texto(str(norma or ""))
    return achada.chave if achada else str(norma or "")


def texto_da_peca(peca: dict) -> str:
    partes = []
    for documento in peca.get("documentos") or []:
        partes += [documento.get("enderecamento") or "", documento.get("preambulo") or ""]
        for secao in documento.get("secoes") or []:
            partes += [secao.get("titulo") or ""] + list(secao.get("paragrafos") or []) + list(secao.get("itens") or [])
    return "\n".join(str(p) for p in partes)


def marcar_inexistentes(peca: dict, inexistentes: list[dict]):
    """Põe "[CONFERIR O NÚMERO]" logo depois do número de cada artigo citado que não existe na lei."""
    errados = {(c["norma"], _so_numero(c["artigo"])) for c in inexistentes}

    def marcar(texto: str) -> str:
        texto = str(texto or "")
        saida, posicao = [], 0
        for achado, norma, _, artigos in _citacoes(texto):
            if any((norma.chave, _so_numero(a)) in errados for a in artigos):
                saida += [texto[posicao:achado.end()], " [CONFERIR O NÚMERO]"]
                posicao = achado.end()
        return "".join(saida) + texto[posicao:]

    for documento in peca.get("documentos") or []:
        documento["preambulo"] = marcar(documento.get("preambulo"))
        for secao in documento.get("secoes") or []:
            secao["paragrafos"] = [marcar(p) for p in secao.get("paragrafos") or []]
            secao["itens"] = [marcar(i) for i in secao.get("itens") or []]


# ------------------------------------------------------------------ Word e PDF
RE_COLCHETES = re.compile(r"(\[[^\]]{1,80}\])")
LETRAS = "abcdefghijklmnopqrstuvwxyz"


def _linha_do_processo(documento: dict) -> str:
    """"Processo nº 123…"; sem repetir quando a IA já escreveu "Processo de origem nº…"."""
    processo = str(documento.get("processo") or "").strip()
    if not processo:
        return ""
    return processo if re.match(r"(?i)(autos|processo|proc\.)", processo) else f"Processo nº {processo}"


def _linha_do_valor(peca: dict) -> str:
    """"Dá-se à causa o valor de R$ 10.000,00." só nas peças que têm valor da causa e com um valor de verdade (ou o
    colchete para completar); num recurso, a IA às vezes escreve "Não se aplica…", que não vai para a peça."""
    valor = str(peca.get("valor_da_causa") or "").strip().rstrip(".")
    estrutura = " ".join(TIPOS_DE_PECA.get(peca.get("tipo"), {}).get("estrutura", [])).lower()
    if not valor or "valor da causa" not in estrutura or re.search(r"(?i)n[ãa]o se aplica|n[ãa]o h[áa]", valor):
        return ""
    if not re.search(r"R\$|\[", valor):
        return ""
    return f"Dá-se à causa o valor de {valor}."


def _itens(secao: dict) -> list[str]:
    itens = [re.sub(r"^\s*(?:[a-z]\)|\d+[.)]|-)\s*", "", str(i)).strip() for i in secao.get("itens") or [] if str(i).strip()]
    return [f"{LETRAS[n % 26]}) {item}" for n, item in enumerate(itens)]


def gerar_docx_peca(peca: dict, caminho: str) -> str:
    """O Word da peça: Times New Roman 12, justificado, 1,5 de entrelinha, recuo de 2,5 cm, margens de 3 e 2 cm; os
    colchetes (dados a completar) em amarelo e as citações a conferir em vermelho."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
    from docx.shared import Cm, Pt

    doc = Document()
    for secao in doc.sections:
        secao.top_margin, secao.left_margin, secao.bottom_margin, secao.right_margin = Cm(3), Cm(3), Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)

    def escrever(texto: str, negrito=False, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, recuo=Cm(2.5)):
        paragrafo = doc.add_paragraph()
        paragrafo.alignment = alinhamento
        paragrafo.paragraph_format.first_line_indent = recuo
        for pedaco in RE_COLCHETES.split(str(texto or "")):
            if not pedaco:
                continue
            run = paragrafo.add_run(pedaco)
            run.bold = negrito
            if RE_COLCHETES.fullmatch(pedaco):
                run.font.highlight_color = WD_COLOR_INDEX.RED if "CONFERIR" in pedaco else WD_COLOR_INDEX.YELLOW
        return paragrafo

    for indice, documento in enumerate(peca.get("documentos") or []):
        if indice:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        escrever(documento.get("enderecamento"), negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, recuo=Cm(0))
        for _ in range(3):
            doc.add_paragraph()  # o espaço tradicional depois do endereçamento
        if _linha_do_processo(documento):
            escrever(_linha_do_processo(documento), alinhamento=WD_ALIGN_PARAGRAPH.LEFT, recuo=Cm(0))
        if documento.get("preambulo"):
            escrever(documento["preambulo"], recuo=Cm(4))
        if documento.get("titulo") and documento.get("titulo") not in (documento.get("preambulo") or ""):
            escrever(documento["titulo"], negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, recuo=Cm(0))
        for secao in documento.get("secoes") or []:
            if secao.get("titulo"):
                escrever(secao["titulo"], negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.LEFT, recuo=Cm(0))
            for paragrafo in secao.get("paragrafos") or []:
                escrever(paragrafo)
            for item in _itens(secao):
                escrever(item, recuo=Cm(1.25))
        if indice == len(peca["documentos"]) - 1 and _linha_do_valor(peca):
            escrever(_linha_do_valor(peca))
        for linha in documento.get("fechamento") or []:
            escrever(linha, alinhamento=WD_ALIGN_PARAGRAPH.CENTER, recuo=Cm(0))
    doc.save(caminho)
    return caminho


def _html_texto(texto: str) -> str:
    seguro = html.escape(str(texto or ""))
    return RE_COLCHETES.sub(lambda a: f'<span class="{"conferir" if "CONFERIR" in a.group(1) else "falta"}">{a.group(1)}</span>', seguro)


def html_da_peca(peca: dict) -> str:
    partes = []
    for indice, documento in enumerate(peca.get("documentos") or []):
        quebra = ' style="page-break-before: always"' if indice else ""
        partes.append(f'<p class="enderecamento"{quebra}>{_html_texto(documento.get("enderecamento"))}</p>'
                      '<p class="espaco">&nbsp;</p>')
        if _linha_do_processo(documento):
            partes.append(f'<p class="processo">{_html_texto(_linha_do_processo(documento))}</p>')
        if documento.get("preambulo"):
            partes.append(f'<p class="preambulo">{_html_texto(documento["preambulo"])}</p>')
        if documento.get("titulo") and documento.get("titulo") not in (documento.get("preambulo") or ""):
            partes.append(f'<p class="titulo">{_html_texto(documento["titulo"])}</p>')
        for secao in documento.get("secoes") or []:
            if secao.get("titulo"):
                partes.append(f'<p class="secao">{_html_texto(secao["titulo"])}</p>')
            partes += [f"<p>{_html_texto(p)}</p>" for p in secao.get("paragrafos") or []]
            partes += [f'<p class="item">{_html_texto(i)}</p>' for i in _itens(secao)]
        if indice == len(peca["documentos"]) - 1 and _linha_do_valor(peca):
            partes.append(f"<p>{_html_texto(_linha_do_valor(peca))}</p>")
        partes += [f'<p class="fecho">{_html_texto(l)}</p>' for l in documento.get("fechamento") or []]
    return "".join(partes)


CSS_DA_PECA = """
body { font-family: serif; font-size: 12pt; line-height: 1.5; color: #111; }
p { text-align: justify; text-indent: 2.5cm; margin: 0 0 6pt 0; }
.enderecamento { font-weight: bold; text-indent: 0; text-align: left; }
.espaco { margin-bottom: 40pt; }
.processo { text-indent: 0; text-align: left; }
.preambulo { text-indent: 4cm; }
.titulo { font-weight: bold; text-indent: 0; text-align: center; margin: 10pt 0; }
.secao { font-weight: bold; text-indent: 0; text-align: left; margin-top: 12pt; }
.item { text-indent: 0; margin-left: 1.25cm; }
.fecho { text-indent: 0; text-align: center; margin: 2pt 0; }
.falta { background-color: #FDE047; }
.conferir { background-color: #FCA5A5; font-weight: bold; }
"""


def gerar_pdf_peca(peca: dict, caminho: str) -> str:
    """O PDF da peça, igual ao Word (A4, margens de 3 e 2 cm), com o número da página no rodapé."""
    import pymupdf

    historia = pymupdf.Story(html=html_da_peca(peca), user_css=CSS_DA_PECA)
    pagina = pymupdf.paper_rect("a4")
    cm = 72 / 2.54
    area = pagina + (3 * cm, 3 * cm, -2 * cm, -2 * cm)
    escritor = pymupdf.DocumentWriter(caminho)
    mais = True
    while mais:
        dispositivo = escritor.begin_page(pagina)
        mais, _ = historia.place(area)
        historia.draw(dispositivo)
        escritor.end_page()
    escritor.close()
    documento = pymupdf.open(caminho)
    for indice, folha in enumerate(documento, start=1):
        folha.insert_text((folha.rect.width - 2 * cm - 20, folha.rect.height - 1.2 * cm), str(indice), fontsize=9,
                          color=(0.4, 0.4, 0.4))
    documento.saveIncr()
    documento.close()
    return caminho
