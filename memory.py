import os
import time
import glob
import json
import logging
import re
import hashlib
import math
import mimetypes
import shutil
import threading
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from collections import OrderedDict
from datetime import datetime
import numpy as np
import requests
from bs4 import BeautifulSoup
import chromadb
import pypdf

from config import (
    CAMINHO_DOCUMENTOS,
    MODOS_PASTAS,
    MODO_CONVERSA_PADRAO,
    PASTAS_OBSIDIAN,
    VAULT_OBSIDIAN,
    CAMINHO_MANIFESTO_MEMORIA,
    CAMINHO_AUDITORIA_MEMORIA,
    IDADE_MAXIMA_DOCUMENTO_DIAS,
    INTERVALO_OBSERVADOR_MEMORIA,
    NOTELOOKLM_PATH,
    ARQUIVO_MEMORIA_USUARIO,
    PASTA_MEMORIA_USUARIO,
    AUTO_CARREGAR_FONTES,
    OLLAMA_HOST,
    OLLAMA_KEEP_ALIVE,
    OLLAMA_EMBED_MODEL,
    TAMANHO_TRECHO,
    SOBREPOSICAO_TRECHO,
)

logger = logging.getLogger(__name__)

# Cache
_CACHE_INDEXACAO = {"timestamp": 0, "arquivos_processados": set(), "mtime_map": {}, "pendentes": []}
_OBSERVADOR_MEMORIA = None
_INDEXACAO_LOCK = threading.Lock()
_CACHE_MEMORIA = None
_TIMESTAMP_CACHE_MEMORIA = 0
_INDEXACAO_EM_ANDAMENTO = False
_INDEXACAO_PENDENTE = False
_STATUS_INDEXACAO = {
    "status": "inativa",
    "mensagem": "Biblioteca pronta para uso.",
    "progresso": 0,
    "total": 0,
}


def _atualizar_status_indexacao(status: str, mensagem: str | None = None, progresso: int | None = None, total: int | None = None):
    global _STATUS_INDEXACAO
    _STATUS_INDEXACAO["status"] = status
    if mensagem is not None:
        _STATUS_INDEXACAO["mensagem"] = mensagem
    if progresso is not None:
        _STATUS_INDEXACAO["progresso"] = progresso
    if total is not None:
        _STATUS_INDEXACAO["total"] = total


def status_indexacao():
    return dict(_STATUS_INDEXACAO)


def agendar_indexacao_em_background():
    """Inicia a indexação em thread separada e evita vários scans concorrentes."""
    global _INDEXACAO_PENDENTE
    if not AUTO_CARREGAR_FONTES:
        _atualizar_status_indexacao("inativa", "Carregamento automático de fontes desativado.")
        return False

    with _INDEXACAO_LOCK:
        if _INDEXACAO_EM_ANDAMENTO:
            _INDEXACAO_PENDENTE = True
            _atualizar_status_indexacao("carregando", "Biblioteca já está sendo carregada em segundo plano.")
            return False

    _atualizar_status_indexacao("carregando", "Biblioteca de conhecimento sendo carregada em segundo plano.")

    def aquecer_e_indexar():
        _aquecer_embedding()
        indexar_pasta_documentos()

    threading.Thread(target=aquecer_e_indexar, daemon=True).start()
    return True

PASTA_DOCUMENTOS = CAMINHO_DOCUMENTOS
os.makedirs(PASTA_DOCUMENTOS, exist_ok=True)
MODO_CONVERSA_ATUAL = MODO_CONVERSA_PADRAO
ARQUIVO_HISTORICO_CONVERSAS = os.path.join(os.path.dirname(__file__), "historico_conversas.jsonl")

# --- Embedding ---
# A busca por significado usa um modelo multilíngue servido pelo Ollama (bge-m3 por padrão),
# sem PyTorch. Cada modelo tem a própria coleção e o próprio manifesto: vetores de modelos
# diferentes não se comparam, então trocar o modelo começa uma indexação completa em vez de
# misturá-los. Sem o modelo no Ollama, a biblioteca usa o embedding padrão do ChromaDB (em
# inglês), também numa coleção só dele, e avisa no log.

from chromadb.utils.embedding_functions.ollama_embedding_function import OllamaEmbeddingFunction

LOTE_EMBEDDING = 32  # trechos por chamada ao Ollama
TEMPO_LIMITE_EMBEDDING = 120  # segundos por lote
TAMANHO_CACHE_PERGUNTAS = 64


class EmbeddingOllama(OllamaEmbeddingFunction):
    """Embedding do Ollama em lotes, que mantém o modelo carregado e guarda as últimas perguntas.

    Herda o nome "ollama" do ChromaDB, que grava modelo e endereço na configuração da coleção.
    Cada pergunta passa por até três buscas (com e sem o filtro de palavras); o cache evita
    pedir o mesmo vetor ao Ollama de novo.
    """

    def __init__(self, url=OLLAMA_HOST, model_name=OLLAMA_EMBED_MODEL, timeout=TEMPO_LIMITE_EMBEDDING):
        super().__init__(url=url, model_name=model_name, timeout=timeout)
        self._perguntas = OrderedDict()
        self._lock_perguntas = threading.Lock()

    def __call__(self, input):
        textos = list(input)
        vetores = []
        for inicio in range(0, len(textos), LOTE_EMBEDDING):
            resposta = self._client.embed(
                model=self.model_name,
                input=textos[inicio:inicio + LOTE_EMBEDDING],
                keep_alive=OLLAMA_KEEP_ALIVE,
            )
            vetores.extend(np.array(vetor, dtype=np.float32) for vetor in resposta["embeddings"])
        return vetores

    def embed_query(self, input):
        textos = list(input)
        with self._lock_perguntas:
            prontos = {texto: self._perguntas[texto] for texto in textos if texto in self._perguntas}
        faltando = [texto for texto in dict.fromkeys(textos) if texto not in prontos]
        if faltando:
            prontos.update(zip(faltando, self(faltando)))
        with self._lock_perguntas:
            for texto in textos:
                self._perguntas[texto] = prontos[texto]
                self._perguntas.move_to_end(texto)
            while len(self._perguntas) > TAMANHO_CACHE_PERGUNTAS:
                self._perguntas.popitem(last=False)
        return [prontos[texto] for texto in textos]


def _verificar_modelo_embedding(modelo):
    """'ok', 'ausente' (o Ollama respondeu e não tem o modelo) ou 'sem_ollama' (não respondeu)."""
    try:
        resposta = requests.post(f"{OLLAMA_HOST}/api/show", json={"model": modelo}, timeout=(1, 3))
    except requests.RequestException:
        return "sem_ollama"
    if resposta.status_code == 404:
        return "ausente"
    return "ok" if resposta.ok else "sem_ollama"


def _sufixo_modelo(modelo):
    """'bge-m3' e 'bge-m3:latest' viram 'bge_m3', que nomeia a coleção e o manifesto."""
    nome = (modelo or "").strip().lower()
    if nome.endswith(":latest"):
        nome = nome[: -len(":latest")]
    return re.sub(r"[^a-z0-9]+", "_", nome).strip("_") or "modelo"


def _escolher_embedding(modelo=OLLAMA_EMBED_MODEL):
    """(função de embedding, ou None para o padrão do ChromaDB; sufixo das coleções; aviso ou None)."""
    modelo = (modelo or "").strip()
    if not modelo:
        return None, "padrao", "OLLAMA_EMBED_MODEL está vazio: a biblioteca usa o embedding padrão do ChromaDB, em inglês."
    situacao = _verificar_modelo_embedding(modelo)
    if situacao == "ausente":
        return None, "padrao", (
            f"O modelo de embedding '{modelo}' não está no Ollama: a busca por significado usa o padrão "
            f"do ChromaDB (em inglês) até você rodar 'ollama pull {modelo}' e abrir a Beatriz de novo."
        )
    try:
        funcao = EmbeddingOllama(OLLAMA_HOST, modelo)
    except Exception as exc:  # pacote 'ollama' ausente
        return None, "padrao", f"Não consegui preparar o embedding do Ollama ({exc}): a biblioteca usa o padrão do ChromaDB."
    aviso = None
    if situacao == "sem_ollama":
        # Provavelmente o Ollama ainda está subindo: fica com a coleção do modelo, e o que
        # falhar agora (indexação, busca por significado) volta a funcionar quando ele responder.
        aviso = (
            f"O Ollama não respondeu ao abrir a biblioteca: a indexação e a busca por significado com "
            f"'{modelo}' voltam quando ele estiver no ar; até lá, só a busca por palavras."
        )
    return funcao, _sufixo_modelo(modelo), aviso


FUNCAO_EMBEDDING, SUFIXO_EMBEDDING, AVISO_EMBEDDING = _escolher_embedding()
NOME_COLECAO_DOCS = f"documentos_usuario_v3_{SUFIXO_EMBEDDING}"
_AVISO_EMBEDDING_REGISTRADO = False
if AVISO_EMBEDDING:
    print(f"Aviso: {AVISO_EMBEDDING}")


def _avisar_embedding():
    """Registra o aviso do embedding no log uma vez (no import, o log do main ainda não existe)."""
    global _AVISO_EMBEDDING_REGISTRADO
    if AVISO_EMBEDDING and not _AVISO_EMBEDDING_REGISTRADO:
        _AVISO_EMBEDDING_REGISTRADO = True
        logger.warning(AVISO_EMBEDDING)
        _registrar_auditoria("embedding_degradado", colecao=NOME_COLECAO_DOCS, aviso=AVISO_EMBEDDING)


_EMBEDDING_AQUECIDO = False


def _aquecer_embedding():
    """Carrega o modelo de embedding no Ollama (~3 s na primeira vez), para a primeira pergunta não esperar."""
    global _EMBEDDING_AQUECIDO
    if FUNCAO_EMBEDDING is None or _EMBEDDING_AQUECIDO:
        return
    try:
        FUNCAO_EMBEDDING(["Beatriz"])
        _EMBEDDING_AQUECIDO = True
    except Exception as exc:  # Ollama fora do ar: tenta de novo na próxima indexação
        logger.warning("Não consegui carregar o modelo de embedding '%s' no Ollama: %s", OLLAMA_EMBED_MODEL, exc)


def status_embedding():
    return {
        "modelo": OLLAMA_EMBED_MODEL if FUNCAO_EMBEDDING is not None else "padrão do ChromaDB",
        "colecao": NOME_COLECAO_DOCS,
        "aviso": AVISO_EMBEDDING,
    }


def _abrir_colecao_docs(cliente):
    if FUNCAO_EMBEDDING is None:
        return cliente.get_or_create_collection(name=NOME_COLECAO_DOCS)
    return cliente.get_or_create_collection(name=NOME_COLECAO_DOCS, embedding_function=FUNCAO_EMBEDDING)


# As coleções antigas (documentos_usuario_v2, com o embedding em inglês e cortes fixos)
# ficam no banco sem uso, para não apagar dados.
banco_memoria_path = os.path.join(NOTELOOKLM_PATH, "banco_memoria")
os.makedirs(banco_memoria_path, exist_ok=True)
CAMINHO_MANIFESTO_MEMORIA = os.path.join(banco_memoria_path, f"manifesto_fontes_v3_{SUFIXO_EMBEDDING}.json")
BANCO_EM_RAM = False
_MANIFESTO_EM_RAM = {}

try:
    chroma_client = chromadb.PersistentClient(path=banco_memoria_path)
    colecao_memoria = chroma_client.get_or_create_collection(name="memoria_usuario")
    colecao_docs = _abrir_colecao_docs(chroma_client)
except Exception as exc:
    print(f"Aviso: banco persistente indisponível, usando memória em RAM ({exc})")
    # O banco em RAM começa vazio a cada abertura, então o manifesto do disco (que diz o que já
    # foi indexado no banco persistente) não vale: este fica só na memória desta sessão.
    BANCO_EM_RAM = True
    CAMINHO_MANIFESTO_MEMORIA = os.path.join(banco_memoria_path, f"manifesto_fontes_v3_{SUFIXO_EMBEDDING}_ram.json")
    chroma_client = chromadb.Client()
    colecao_memoria = chroma_client.get_or_create_collection(name="memoria_usuario")
    colecao_docs = _abrir_colecao_docs(chroma_client)


# --- Palavras-chave ---
# A busca começa pelas palavras-chave da pergunta: elas filtram os candidatos por um
# índice (FTS do ChromaDB ou o índice invertido do histórico) em vez de varrer tudo.

PALAVRAS_VAZIAS = {
    "a", "o", "as", "os", "um", "uma", "uns", "umas", "de", "do", "da", "dos", "das", "em", "no", "na",
    "nos", "nas", "por", "pelo", "pela", "pelos", "pelas", "para", "pra", "pro", "com", "sem", "sob",
    "sobre", "entre", "ate", "e", "ou", "mas", "que", "se", "como", "quando", "onde", "qual", "quais",
    "quem", "quanto", "quanta", "quantos", "quantas", "porque", "por que", "eu", "voce", "ele", "ela",
    "eles", "elas", "nos", "me", "te", "lhe", "meu", "minha", "meus", "minhas", "seu", "sua", "seus",
    "suas", "isso", "isto", "esse", "essa", "este", "esta", "aquele", "aquela", "aqui", "ali", "la",
    "ja", "nao", "sim", "muito", "mais", "menos", "tambem", "so", "ser", "estar", "ter", "haver", "fazer",
    "foi", "era", "sao", "esta", "estao", "tem", "tinha", "vai", "vou", "pode", "poderia", "sobre",
    "beatriz", "favor", "diga", "fale", "explique", "sabe", "saber", "algo", "coisa", "coisas",
    "falei", "falou", "falamos", "falar", "disse", "dizer", "conversamos", "perguntei", "contei",
    "quero", "queria", "preciso", "gostaria", "lembra", "lembrar", "mostre", "mostra", "conte", "conta",
    "agora", "hoje", "ontem", "amanha", "depois", "antes", "sempre", "nunca", "tudo", "nada", "todo", "toda",
}

_VARIANTES_LETRA = {"a": "[aáàâã]", "e": "[eéê]", "i": "[ií]", "o": "[oóôõ]", "u": "[uúü]", "c": "[cç]"}


def _sem_acentos(texto: str) -> str:
    return unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode("ascii").lower()


def _raiz(palavra: str) -> str:
    """Reduz plural e acentos: 'Memórias' e 'memoria' viram 'memoria'."""
    palavra = _sem_acentos(palavra)
    # 'informação' e 'informações' ficam com a mesma raiz 'informac'.
    if len(palavra) > 5 and palavra.endswith(("oes", "aes", "aos")):
        return palavra[:-3]
    if len(palavra) > 4 and palavra.endswith("ao"):
        return palavra[:-2]
    if len(palavra) > 4 and palavra.endswith(("res", "zes")):
        return palavra[:-2]
    if len(palavra) > 3 and palavra.endswith("s"):
        return palavra[:-1]
    return palavra


def extrair_palavras_chave(texto: str, limite: int = 8) -> list[str]:
    """Raízes das palavras relevantes da pergunta, das mais longas (mais específicas) às mais curtas."""
    vistas = []
    for palavra in re.findall(r"[\wÀ-ÿ]+", texto or ""):
        normal = _sem_acentos(palavra)
        if len(normal) < 3 or normal in PALAVRAS_VAZIAS or (normal.isdigit() and len(normal) < 4):
            continue
        raiz = _raiz(normal)
        if raiz not in vistas:
            vistas.append(raiz)
    return sorted(vistas, key=len, reverse=True)[:limite]


def raizes_do_texto(texto: str) -> set[str]:
    return {_raiz(palavra) for palavra in re.findall(r"[\wÀ-ÿ]+", texto or "") if len(palavra) >= 3}


def _regex_palavra(raiz: str) -> str:
    """Regex sem diferenciar maiúsculas e acentos, que casa a raiz no início de uma palavra."""
    return "(?i)\\b" + "".join(_VARIANTES_LETRA.get(letra, re.escape(letra)) for letra in raiz)


def _cobertura(palavras_chave: list[str], texto: str) -> float:
    """Fração (ponderada pelo tamanho) das palavras-chave presentes no texto."""
    if not palavras_chave:
        return 0.0
    normal = _sem_acentos(texto)
    pesos = {raiz: len(raiz) for raiz in palavras_chave}
    encontrados = sum(peso for raiz, peso in pesos.items() if re.search(r"\b" + re.escape(raiz), normal))
    return encontrados / sum(pesos.values())


def mudar_modo_conversa(novo_modo):
    """Altera o modo de conversa e cria a pasta associada quando necessário."""
    global MODO_CONVERSA_ATUAL

    modo = (novo_modo or "").strip().lower()
    if not modo:
        return f"Modo atual: {MODO_CONVERSA_ATUAL}"

    if modo not in MODOS_PASTAS:
        modos_disponiveis = ", ".join(MODOS_PASTAS.keys())
        return f"Modo inválido: '{novo_modo}'. Modos disponíveis: {modos_disponiveis}."

    MODO_CONVERSA_ATUAL = modo
    pasta = MODOS_PASTAS[modo]
    os.makedirs(pasta, exist_ok=True)
    return f"Modo alterado para '{modo}'. Pasta ativa: {pasta}"


def listar_arquivos_notelooklm(base_path=None, extensoes=(".txt", ".md", ".json", ".jsonl", ".pdf", ".png", ".jpg", ".jpeg", ".ppt", ".pptx", ".docx", ".xlsx", ".html")):
    base = base_path or PASTA_DOCUMENTOS
    arquivos = []
    for ext in extensoes:
        arquivos.extend(glob.glob(os.path.join(base, "**", f"*{ext}"), recursive=True))
        arquivos.extend(glob.glob(os.path.join(base, f"*{ext}"), recursive=False))
    arquivos = sorted(set(arquivos))
    return arquivos


def _fontes_raiz():
    return [
        (PASTA_DOCUMENTOS, "documento_local"),
        (VAULT_OBSIDIAN, "obsidian"),
        (ARQUIVO_HISTORICO_CONVERSAS, "historico"),
    ]


def _id_fonte(caminho):
    return hashlib.sha256(os.path.realpath(caminho).encode("utf-8")).hexdigest()[:24]


def _hash_conteudo(conteudo):
    return hashlib.sha256(conteudo.encode("utf-8", errors="ignore")).hexdigest()


def _registrar_auditoria(evento, **dados):
    try:
        os.makedirs(os.path.dirname(CAMINHO_AUDITORIA_MEMORIA), exist_ok=True)
        with open(CAMINHO_AUDITORIA_MEMORIA, "a", encoding="utf-8") as arquivo:
            arquivo.write(json.dumps({"evento": evento, "timestamp": datetime.now().isoformat(), **dados}, ensure_ascii=False) + "\n")
    except OSError:
        pass


def _resolver_pasta_obsidian(pasta=None):
    """Retorna uma pasta da Vault sem permitir caminhos fora dela."""
    nome_pasta = (pasta or MODO_CONVERSA_ATUAL or "geral").strip().lower()
    caminho = PASTAS_OBSIDIAN.get(nome_pasta)
    if caminho is None:
        caminho = os.path.join(VAULT_OBSIDIAN, nome_pasta)

    vault = os.path.realpath(VAULT_OBSIDIAN)
    caminho_real = os.path.realpath(caminho)
    if os.path.commonpath([vault, caminho_real]) != vault:
        raise ValueError("A pasta informada precisa estar dentro da Vault do Obsidian.")
    return caminho_real


def criar_nota_obsidian(titulo: str, conteudo: str, pasta: str = None) -> str:
    """Cria ou atualiza uma nota Markdown na pasta de memória escolhida."""
    try:
        titulo_limpo = os.path.basename((titulo or "").strip())
        if not titulo_limpo:
            return "Erro ao salvar nota no Obsidian: informe um título."
        if not titulo_limpo.lower().endswith(".md"):
            titulo_limpo += ".md"

        caminho_pasta = _resolver_pasta_obsidian(pasta)
        os.makedirs(caminho_pasta, exist_ok=True)
        caminho_arquivo = os.path.join(caminho_pasta, titulo_limpo)
        data_atual = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        nota = (
            "---\n"
            f"data: {data_atual}\n"
            "criado_por: Beatriz\n"
            f"pasta: {os.path.basename(caminho_pasta)}\n"
            "---\n\n"
            f"{conteudo or ''}\n"
        )
        with open(caminho_arquivo, "w", encoding="utf-8") as arquivo:
            arquivo.write(nota)

        indexar_pasta_documentos()
        return f"Nota '{titulo_limpo}' salva na Vault do Obsidian em '{os.path.basename(caminho_pasta)}'."
    except Exception as exc:
        return f"Erro ao salvar nota no Obsidian: {exc}"


def buscar_notas_obsidian(termo: str, pasta: str = None, limite: int = 10) -> str:
    """Busca notas da Vault que contenham todas as palavras-chave do termo, em qualquer ordem."""
    consulta = (termo or "").strip().lower()
    if not consulta:
        return "Informe um termo para buscar nas notas do Obsidian."

    try:
        vault = os.path.realpath(VAULT_OBSIDIAN)
        base = _resolver_pasta_obsidian(pasta) if pasta else vault
        if os.path.commonpath([vault, base]) != vault:
            return "Erro ao buscar notas: a pasta precisa estar dentro da Vault do Obsidian."

        palavras_chave = extrair_palavras_chave(consulta) or [_raiz(consulta)]

        def contem_todas(texto):
            return _cobertura(palavras_chave, texto) == 1.0

        # Primeiro pelo índice da biblioteca: só os arquivos cujos trechos têm as palavras-chave.
        caminhos = []
        try:
            filtros = [{"$regex": _regex_palavra(raiz)} for raiz in palavras_chave]
            encontrados = colecao_docs.get(
                where={"tipo_fonte": "obsidian"},
                where_document=filtros[0] if len(filtros) == 1 else {"$or": filtros},
                include=["metadatas"],
                limit=200,
            )
            caminhos = sorted({m.get("caminho") for m in encontrados.get("metadatas") or [] if m.get("caminho")})
        except Exception as exc:
            _registrar_auditoria("obsidian_keyword_search_failed", consulta=consulta, erro=str(exc))
        if not caminhos:
            # Nota ainda não indexada: recorre à leitura das notas da pasta.
            caminhos = listar_arquivos_notelooklm(base, extensoes=(".md",))

        resultados = []
        for caminho in caminhos:
            real = os.path.realpath(caminho)
            if not os.path.isfile(real) or os.path.commonpath([os.path.realpath(base), real]) != os.path.realpath(base):
                continue
            try:
                with open(caminho, "r", encoding="utf-8", errors="ignore") as arquivo:
                    conteudo = arquivo.read()
            except OSError:
                continue
            if not contem_todas(f"{os.path.basename(caminho)} {conteudo}"):
                continue
            trecho = conteudo[:500].strip().replace("\n", " ")
            relativo = os.path.relpath(caminho, VAULT_OBSIDIAN)
            resultados.append(f"[Nota: {relativo}]\n{trecho}")
            if len(resultados) >= limite:
                break

        return "\n\n".join(resultados) if resultados else "Nenhuma nota correspondente no Obsidian."
    except Exception as exc:
        return f"Erro ao buscar notas no Obsidian: {exc}"


def _extrair_texto_pptx(caminho, limite=50000):
    try:
        with zipfile.ZipFile(caminho) as zf:
            slides = []
            for nome in zf.namelist():
                if nome.startswith("ppt/slides/slide") and nome.endswith(".xml"):
                    xml = zf.read(nome)
                    root = ET.fromstring(xml)
                    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main", "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
                    textos = []
                    for texto in root.findall(".//a:t", ns):
                        if texto.text:
                            textos.append(texto.text)
                    if textos:
                        slides.append(" ".join(textos))
            return "\n\n".join(slides)[:limite]
    except Exception:
        return ""


def _extrair_texto_imagem(caminho, limite=50000):
    nome = os.path.basename(caminho)
    try:
        from PIL import Image
        try:
            import pytesseract  # type: ignore
            imagem = Image.open(caminho)
            texto = pytesseract.image_to_string(imagem, lang="por")
            if texto and texto.strip():
                return texto[:limite]
        except Exception:
            pass
    except Exception:
        pass
    return f"[Imagem detectada: {nome}] Arquivo visual carregado. Sem OCR disponível para extrair texto automaticamente."


def _extrair_texto_html(caminho, limite=50000):
    try:
        with open(caminho, "r", encoding="utf-8", errors="ignore") as f:
            html = f.read()
        soup = BeautifulSoup(html, "html.parser")
        texto = soup.get_text(" ", strip=True)
        return texto[:limite]
    except Exception:
        return ""


def _extrair_unidades(caminho, limite=50000):
    """Extrai unidades com localizadores para permitir citações úteis."""
    ext = os.path.splitext(caminho)[1].lower()
    try:
        if ext == ".pdf":
            reader = pypdf.PdfReader(caminho)
            return [(pagina.extract_text() or "", f"página {indice}") for indice, pagina in enumerate(reader.pages, start=1)]
        if ext == ".docx":
            from docx import Document
            documento = Document(caminho)
            return [(paragrafo.text, f"parágrafo {indice}") for indice, paragrafo in enumerate(documento.paragraphs, start=1) if paragrafo.text.strip()]
        if ext == ".xlsx":
            from openpyxl import load_workbook
            planilha = load_workbook(caminho, read_only=True, data_only=True)
            unidades = []
            for nome in planilha.sheetnames:
                linhas = []
                for linha in planilha[nome].iter_rows(values_only=True):
                    valores = [str(valor) for valor in linha if valor is not None]
                    if valores:
                        linhas.append(" | ".join(valores))
                if linhas:
                    unidades.append(("\n".join(linhas), f"planilha {nome}"))
            return unidades
        if ext == ".jsonl":
            unidades = []
            with open(caminho, "r", encoding="utf-8", errors="ignore") as arquivo:
                for indice, linha in enumerate(arquivo, start=1):
                    if linha.strip():
                        try:
                            texto = json.dumps(json.loads(linha), ensure_ascii=False)
                        except json.JSONDecodeError:
                            texto = linha.strip()
                        unidades.append((texto, f"registro {indice}"))
            return unidades
        if ext in {".txt", ".md", ".json", ".html", ".htm", ".ppt", ".pptx", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}:
            return [(extrair_texto_arquivo_otimizado(caminho, limite), "documento")]
    except Exception as exc:
        _registrar_auditoria("ingest_failed", caminho=caminho, erro=str(exc))
    return []


def extrair_texto_arquivo_otimizado(caminho, limite=50000):
    ext = os.path.splitext(caminho)[1].lower()
    try:
        if ext in {".txt", ".md", ".json"}:
            with open(caminho, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(limite)
        elif ext == ".pdf":
            reader = pypdf.PdfReader(caminho)
            texto = ""
            for pagina in reader.pages:
                texto += (pagina.extract_text() or "") + "\n"
                if len(texto) > limite:
                    return texto[:limite]
            return texto
        elif ext in {".png", ".jpg", ".jpeg", ".bmp", ".webp"}:
            return _extrair_texto_imagem(caminho, limite=limite)
        elif ext == ".pptx":
            return _extrair_texto_pptx(caminho, limite=limite)
        elif ext == ".docx":
            return "\n".join(texto for texto, _ in _extrair_unidades(caminho, limite))[:limite]
        elif ext == ".xlsx":
            return "\n".join(texto for texto, _ in _extrair_unidades(caminho, limite))[:limite]
        elif ext == ".ppt":
            return f"[PowerPoint detectado: {os.path.basename(caminho)}] Arquivo de apresentação carregado. Sem parser nativo disponível para extrair conteúdo automaticamente."
        elif ext in {".html", ".htm"}:
            return _extrair_texto_html(caminho, limite=limite)
    except Exception:
        return ""
    return ""


# --- Trechos ---
# A biblioteca guarda cada arquivo em trechos de ~TAMANHO_TRECHO caracteres. O corte procura
# o fim de um parágrafo ou de uma frase, e cada trecho repete o final do anterior
# (SOBREPOSICAO_TRECHO caracteres), para a ideia que atravessa o corte não se perder.

# Do maior corte natural para o menor: parágrafo, frase, linha e palavra (com o separador
# usado para juntar os pedaços de novo).
_NIVEIS_DE_CORTE = [
    (re.compile(r"\n\s*\n"), "\n\n"),
    (re.compile(r"(?<=[.!?…;])\s+"), " "),
    (re.compile(r"\n"), "\n"),
    (re.compile(r"\s+"), " "),
]
_FIM_DE_FRASE = _NIVEIS_DE_CORTE[1][0]
VERSAO_TRECHOS = f"{TAMANHO_TRECHO}/{SOBREPOSICAO_TRECHO}"


def _segmentar(texto, limite, nivel=0, separador="\n\n"):
    """Pedaços de até `limite` caracteres, cada um com o separador que vinha antes dele.

    Um texto grande demais desce de nível: parágrafos viram frases, frases viram linhas, linhas
    viram palavras, e uma "palavra" enorme (URL, base64) é cortada seca.
    """
    texto = texto.strip()
    if not texto:
        return
    if len(texto) <= limite:
        yield texto, separador
        return
    if nivel == len(_NIVEIS_DE_CORTE):
        for inicio in range(0, len(texto), limite):
            yield texto[inicio:inicio + limite], separador if inicio == 0 else ""
        return
    padrao, separador_do_nivel = _NIVEIS_DE_CORTE[nivel]
    for indice, parte in enumerate(padrao.split(texto)):
        yield from _segmentar(parte, limite, nivel + 1, separador if indice == 0 else separador_do_nivel)


def _final_do_texto(texto, limite):
    """Os últimos `limite` caracteres do texto, começando numa frase (ou ao menos numa palavra)."""
    if len(texto) <= limite:
        return texto
    final = texto[-limite:]
    frase = _FIM_DE_FRASE.search(final)
    if frase and len(final) - frase.end() >= limite // 3:
        return final[frase.end():]
    espaco = re.search(r"\s", final)
    return final[espaco.end():] if espaco else final


def _cauda(pedacos, limite):
    """Os últimos pedaços que cabem em `limite` caracteres: a sobreposição do trecho seguinte."""
    cauda, total = [], 0
    for pedaco, separador, unidade in reversed(pedacos):
        if total + len(pedaco) + len(separador) > limite:
            break
        cauda.insert(0, (pedaco, separador, unidade))
        total += len(pedaco) + len(separador)
    if not cauda and pedacos and limite > 0:
        # O último pedaço sozinho já passa do limite: repete só o final dele.
        pedaco, separador, unidade = pedacos[-1]
        final = _final_do_texto(pedaco, limite)
        return [(final, separador, unidade)] if final else []
    return cauda


def _localizador_intervalo(primeiro, ultimo):
    """'página 3' e 'página 5' viram 'páginas 3 a 5'."""
    if primeiro == ultimo:
        return primeiro
    inicio = re.fullmatch(r"(\w+) (\d+)", primeiro or "")
    fim = re.fullmatch(r"(\w+) (\d+)", ultimo or "")
    if inicio and fim and inicio.group(1) == fim.group(1):
        return f"{inicio.group(1)}s {inicio.group(2)} a {fim.group(2)}"
    return f"{primeiro} a {ultimo}"


def _trechos_das_unidades(unidades, tamanho=None, sobreposicao=None):
    """[(texto, localizador)] de um arquivo -> [(trecho, localizador)].

    Unidades pequenas vizinhas (parágrafos de um Word, registros do histórico, páginas curtas)
    se juntam num trecho só, e o localizador passa a dizer o intervalo ("parágrafos 3 a 7").
    """
    tamanho = max(int(tamanho or TAMANHO_TRECHO), 50)
    sobreposicao = max(0, min(int(SOBREPOSICAO_TRECHO if sobreposicao is None else sobreposicao), tamanho // 2))
    # Os pedaços são menores que o trecho para a sobreposição sempre caber junto com eles.
    limite = tamanho - sobreposicao
    trechos = []
    atual = []  # [(pedaço, separador, índice da unidade)]
    tamanho_atual = 0
    novos = 0  # pedaços do trecho atual que não vieram da sobreposição

    def fechar():
        texto = "".join((separador if posicao else "") + pedaco for posicao, (pedaco, separador, _) in enumerate(atual))
        local = _localizador_intervalo(unidades[atual[0][2]][1], unidades[atual[-1][2]][1])
        trechos.append((texto, local))

    for indice, (texto, _) in enumerate(unidades):
        texto = (texto or "").replace("\r\n", "\n").replace("\r", "\n")
        for pedaco, separador in _segmentar(texto, limite):
            if novos and tamanho_atual + len(separador) + len(pedaco) > tamanho:
                fechar()
                atual = _cauda(atual, sobreposicao)
                tamanho_atual = sum(len(p) + len(s) for p, s, _ in atual)
                novos = 0
            tamanho_atual += (len(separador) if atual else 0) + len(pedaco)
            atual.append((pedaco, separador, indice))
            novos += 1
    if novos:
        fechar()
    return trechos


def dividir_em_trechos(texto, tamanho=None, sobreposicao=None):
    """Trechos de ~`tamanho` caracteres de um texto, cortados em parágrafo ou frase, com sobreposição."""
    return [trecho for trecho, _ in _trechos_das_unidades([(texto, "documento")], tamanho, sobreposicao)]


def _salvar_manifesto(manifesto):
    if BANCO_EM_RAM:
        _MANIFESTO_EM_RAM.clear()
        _MANIFESTO_EM_RAM.update(manifesto)
    # Mesmo com o banco em RAM o arquivo é gravado: o painel lê dele as fontes por assunto.
    os.makedirs(os.path.dirname(CAMINHO_MANIFESTO_MEMORIA), exist_ok=True)
    with open(CAMINHO_MANIFESTO_MEMORIA, "w", encoding="utf-8") as arquivo:
        json.dump(manifesto, arquivo, ensure_ascii=False, indent=2)


def _carregar_manifesto():
    if BANCO_EM_RAM:
        return dict(_MANIFESTO_EM_RAM)
    try:
        with open(CAMINHO_MANIFESTO_MEMORIA, "r", encoding="utf-8") as arquivo:
            return json.load(arquivo)
    except (OSError, json.JSONDecodeError):
        return {}


def indexar_pasta_documentos():
    """Indexa documentos locais, Vault do Obsidian e histórico com IDs estáveis."""
    global _INDEXACAO_EM_ANDAMENTO, _INDEXACAO_PENDENTE

    with _INDEXACAO_LOCK:
        if _INDEXACAO_EM_ANDAMENTO:
            _INDEXACAO_PENDENTE = True
            _atualizar_status_indexacao("carregando", "Indexação já em andamento; nova varredura foi adiada.")
            return "Indexação já em andamento em segundo plano."
        _INDEXACAO_EM_ANDAMENTO = True

    try:
        _avisar_embedding()
        _atualizar_status_indexacao("carregando", "Varredura inicial da biblioteca em andamento...", progresso=0, total=0)
        manifesto = _carregar_manifesto()
        caminhos = []
        extensoes = (".txt", ".md", ".json", ".jsonl", ".pdf", ".png", ".jpg", ".jpeg", ".ppt", ".pptx", ".docx", ".xlsx", ".html", ".htm")
        for raiz, tipo in _fontes_raiz():
            if not os.path.exists(raiz):
                continue
            if os.path.isfile(raiz):
                if os.path.splitext(raiz)[1].lower() in extensoes:
                    caminhos.append((raiz, tipo))
                continue
            for caminho in listar_arquivos_notelooklm(raiz, extensoes=extensoes):
                if os.path.isfile(caminho) and os.path.basename(caminho) != os.path.basename(CAMINHO_MANIFESTO_MEMORIA):
                    caminhos.append((caminho, tipo))

        total_arquivos = len(caminhos)
        _atualizar_status_indexacao("carregando", "Verificando arquivos da biblioteca...", progresso=0, total=max(total_arquivos, 1))

        caminhos_atuais = {os.path.realpath(caminho) for caminho, _ in caminhos}
        for source_id, registro in list(manifesto.items()):
            caminho_registrado = os.path.realpath(registro.get("caminho", ""))
            if caminho_registrado and caminho_registrado not in caminhos_atuais:
                colecao_docs.delete(where={"source_id": source_id})
                manifesto.pop(source_id, None)
                _registrar_auditoria("source_deleted", caminho=caminho_registrado)

        total = 0
        falhas = 0
        sem_ollama = False
        for indice, (caminho, tipo) in enumerate(caminhos, start=1):
            _atualizar_status_indexacao(
                "carregando",
                f"Indexando arquivo {indice}/{total_arquivos}...",
                progresso=min(indice, total_arquivos),
                total=max(total_arquivos, 1),
            )
            estat = os.stat(caminho)
            source_id = _id_fonte(caminho)
            assinatura = f"{estat.st_mtime_ns}:{estat.st_size}"
            registro = manifesto.get(source_id, {})
            # Arquivo igual e cortado do mesmo jeito: nada a fazer. Fontes vazias não têm trechos.
            if registro.get("assinatura") == assinatura and (registro.get("vazio") or registro.get("trechos") == VERSAO_TRECHOS):
                continue

            unidades = _extrair_unidades(caminho)
            if not any(texto.strip() for texto, _ in unidades):
                colecao_docs.delete(where={"source_id": source_id})
                manifesto[source_id] = {"caminho": caminho, "tipo": tipo, "assinatura": assinatura, "vazio": True}
                _registrar_auditoria("ingest_empty", caminho=caminho)
                continue

            colecao_docs.delete(where={"source_id": source_id})
            ids = []
            documentos = []
            metadados = []
            for bloco, localizador in _trechos_das_unidades(unidades):
                chunk_id = f"{source_id}_{len(ids)}"
                ids.append(chunk_id)
                documentos.append(bloco)
                metadados.append({
                    "source_id": source_id,
                    "arquivo": os.path.basename(caminho),
                    "caminho": caminho,
                    "tipo_fonte": tipo,
                    "localizador": localizador,
                    "titulo": os.path.splitext(os.path.basename(caminho))[0],
                    "hash": _hash_conteudo(bloco),
                    "modificado_em": datetime.fromtimestamp(estat.st_mtime).isoformat(),
                })
            if ids:
                try:
                    colecao_docs.upsert(documents=documentos, metadatas=metadados, ids=ids)
                    manifesto[source_id] = {
                        "caminho": caminho,
                        "tipo": tipo,
                        "assinatura": assinatura,
                        "trechos": VERSAO_TRECHOS,
                        "hash": _hash_conteudo("\n".join(documentos)),
                        "atualizado_em": datetime.now().isoformat(),
                    }
                    total += 1
                    _registrar_auditoria("ingest_succeeded", caminho=caminho, chunks=len(ids))
                except Exception as exc:
                    falhas += 1
                    _registrar_auditoria("ingest_upsert_failed", caminho=caminho, erro=str(exc))
                    if isinstance(exc, ConnectionError):
                        # Ollama fora do ar: as outras fontes falhariam igual. A próxima varredura retoma.
                        sem_ollama = True
                        break
                    _atualizar_status_indexacao("falhou", f"Falha na indexação de '{os.path.basename(caminho)}': {exc}")
                    continue

        _salvar_manifesto(manifesto)
        _CACHE_INDEXACAO["timestamp"] = time.time()
        if falhas:
            _atualizar_status_indexacao(
                "falhou",
                "Ollama fora do ar: faltou indexar parte da biblioteca; a próxima varredura retoma."
                if sem_ollama else f"Biblioteca parcialmente carregada: {falhas} falhas de indexação.",
                progresso=max(total_arquivos, 1),
                total=max(total_arquivos, 1),
            )
        else:
            _atualizar_status_indexacao("pronta", f"Biblioteca pronta. {total} fontes indexadas ou atualizadas.", progresso=max(total_arquivos, 1), total=max(total_arquivos, 1))
        return f"{total} fontes indexadas ou atualizadas."
    except Exception as exc:
        _registrar_auditoria("ingest_failed", erro=str(exc))
        _atualizar_status_indexacao("falhou", f"Falha na biblioteca: {exc}")
        return f"Erro: {exc}"
    finally:
        with _INDEXACAO_LOCK:
            _INDEXACAO_EM_ANDAMENTO = False
            if _INDEXACAO_PENDENTE:
                _INDEXACAO_PENDENTE = False
                threading.Thread(target=indexar_pasta_documentos, daemon=True).start()

def deve_consultar_documentos(pergunta, modo=None):
    """Consulta documentos automaticamente para perguntas relevantes, mas não bloqueia comandos simples do sistema nem ferramentas."""
    texto = (pergunta or "").strip()
    if not texto:
        return False

    modo_ativo = (modo or MODO_CONVERSA_ATUAL or MODO_CONVERSA_PADRAO).strip().lower()
    if modo_ativo == "comunicacao":
        return True

    texto_lower = texto.lower()

    termos_explicitos = [
        "não use documentos",
        "sem documentos",
        "ignora documentos",
        "não consultar documentos",
        "não pesquise documentos",
    ]
    if any(termo in texto_lower for termo in termos_explicitos):
        return False

    perguntas_simples = [
        "qual a hora",
        "que horas",
        "qual a data",
        "data de hoje",
        "hoje",
        "agendamento",
        "agenda",
        "me lembre",
        "lembrete",
        "criar evento",
        "marcar compromisso",
        "listar compromissos",
        "próximos eventos",
        "meus compromissos",
        "abra o navegador",
        "abrir navegador",
        "pesquise no youtube",
        "pesquise no spotify",
        "abra o google",
        "abrir o google",
        "pesquise por",
        "abrir app",
        "abra o app",
        "mostrar rota",
        "calendario",
        "abrir calendário",
        "abrir agenda",
        "whatsapp",
        "e-mail",
        "email",
    ]
    if any(termo in texto_lower for termo in perguntas_simples):
        return False

    return True


def buscar_nos_documentos_modo(pergunta, modo=None, quantidade_trechos=None):
    """Os trechos mais relevantes (TRECHOS_POR_RESPOSTA, ou `quantidade_trechos`) com [Fonte: ...] antes de cada um."""
    try:
        if not deve_consultar_documentos(pergunta, modo=modo):
            return "Consulta de documentos desativada por padrão; só será usada quando o usuário pedir ou quando o contexto requerer."

        texto = (pergunta or "").strip()
        if texto.startswith(("http://", "https://")):
            try:
                resposta = requests.get(texto, timeout=10)
                resposta.raise_for_status()
                soup = BeautifulSoup(resposta.text, "html.parser")
                pagina = soup.get_text(" ", strip=True)[:5000]
                if pagina:
                    return f"[Site da web]\n{pagina}"
            except Exception:
                pass

        _avisar_embedding()
        total_documentos = colecao_docs.count()
        if total_documentos == 0:
            indexar_pasta_documentos()
            total_documentos = colecao_docs.count()
        if total_documentos == 0:
            return "Nenhuma informação encontrada."
        modo_busca = (modo or MODO_CONVERSA_ATUAL or MODO_CONVERSA_PADRAO).strip().lower()
        palavras_chave = extrair_palavras_chave(pergunta)
        n_documentos = max(1, min(24, total_documentos))
        filtros = [{"$regex": _regex_palavra(raiz)} for raiz in palavras_chave]
        filtro_alguma = (filtros[0] if len(filtros) == 1 else {"$or": filtros}) if filtros else None

        # 1) Candidatos por significado: os trechos com alguma palavra-chave (filtro pelo índice de
        #    texto do ChromaDB, sem varrer a coleção), os com todas elas e, sem filtro, os parecidos
        #    sem palavras em comum (sinônimos, paráfrases). O vetor da pergunta é calculado uma vez.
        consultas = []
        if filtro_alguma:
            consultas.append((filtro_alguma, n_documentos))
        if len(filtros) > 1:
            consultas.append(({"$and": filtros}, min(12, n_documentos)))
        consultas.append((None, min(12, n_documentos)))

        candidatos = {}
        falhas = []
        for filtro, quantidade in consultas:
            argumentos = {"query_texts": [pergunta], "n_results": quantidade}
            if filtro is not None:
                argumentos["where_document"] = filtro
            try:
                _juntar_candidatos(candidatos, colecao_docs.query(**argumentos), modo_busca, por_significado=True)
            except Exception as exc:
                falhas.append(str(exc))
        if falhas:
            _registrar_auditoria("vector_search_failed", consulta=pergunta, erros=falhas[:3])
        if len(falhas) == len(consultas) and filtro_alguma:
            # Nenhuma busca por significado funcionou (Ollama fora do ar): só as palavras, direto
            # no índice de texto do ChromaDB, que não precisa do embedding.
            try:
                achados = colecao_docs.get(where_document=filtro_alguma, limit=60, include=["documents", "metadatas"])
                _juntar_candidatos(
                    candidatos,
                    {"documents": [achados.get("documents") or []], "metadatas": [achados.get("metadatas") or []]},
                    modo_busca,
                    por_significado=False,
                )
            except Exception as exc:
                _registrar_auditoria("keyword_search_failed", consulta=pergunta, erro=str(exc))

        # 2) Duas ordens dos mesmos candidatos, por significado e por palavras (BM25), fundidas.
        chaves = list(candidatos)
        por_significado = sorted(
            (chave for chave in chaves if candidatos[chave]["distancia"] is not None),
            key=lambda chave: candidatos[chave]["distancia"],
        )
        textos = [f"{candidatos[chave]['metadado'].get('titulo', '')} {candidatos[chave]['documento']}" for chave in chaves]
        por_palavras = [chaves[indice] for indice in ordenar_por_bm25(palavras_chave, textos)]
        ordenados = fundir_rankings(por_palavras, por_significado)[:quantidade_trechos or TRECHOS_POR_RESPOSTA]
        if not ordenados:
            return "Nenhuma informação encontrada."
        resposta = "\n\n".join(
            f"[Fonte: {m.get('arquivo')} | {m.get('tipo_fonte')} | {m.get('localizador')} | {m.get('caminho')}]\n{d}"
            for d, m in ((candidatos[chave]["documento"], candidatos[chave]["metadado"]) for chave in ordenados)
        )
        _registrar_auditoria(
            "search_executed",
            consulta=pergunta,
            palavras_chave=palavras_chave,
            resultados=len(ordenados),
            por_palavras=len(por_palavras),
            por_significado=len(por_significado),
        )
        return resposta
    except Exception as e:
        return f"Erro: {e}"


# --- Busca híbrida ---
# Os candidatos são ordenados duas vezes, por significado (distância dos vetores) e por palavras
# (BM25), e as duas ordens se juntam pela fusão por posição recíproca (RRF): cada lista dá
# 1/(60 + posição) ao trecho. Quem vai bem nas duas sobe; quem só aparece numa ainda entra.
CONSTANTE_RRF = 60
TRECHOS_POR_RESPOSTA = 5
# Por palavras, só contam os trechos com ao menos metade das palavras-chave (pesadas pelo
# tamanho): um trecho que só repete a palavra mais comum da pergunta não é um achado.
COBERTURA_MINIMA_PALAVRAS = 0.5


def _juntar_candidatos(candidatos, resultado, modo, por_significado):
    """Acrescenta os trechos de um resultado do ChromaDB, sem repetir e só os do modo atual."""
    documentos = (resultado.get("documents") or [[]])[0] or []
    metadados = (resultado.get("metadatas") or [[]])[0] or []
    distancias = (resultado.get("distances") or [[]])[0] or []
    for posicao, (documento, metadado) in enumerate(zip(documentos, metadados)):
        metadado = metadado or {}
        if not documento or not _metadado_pertence_ao_modo(metadado, modo):
            continue
        distancia = None
        if por_significado:
            # Sem as distâncias no resultado, a posição na lista faz o papel delas.
            distancia = distancias[posicao] if posicao < len(distancias) and distancias[posicao] is not None else float(posicao)
        chave = (metadado.get("source_id"), metadado.get("hash"))
        atual = candidatos.get(chave)
        if atual is None:
            candidatos[chave] = {"documento": documento, "metadado": metadado, "distancia": distancia}
        elif distancia is not None and (atual["distancia"] is None or distancia < atual["distancia"]):
            atual["distancia"] = distancia


def _tokens(texto):
    """Raízes das palavras do texto, sem as palavras vazias."""
    tokens = []
    for palavra in re.findall(r"[\wÀ-ÿ]+", texto or ""):
        normal = _sem_acentos(palavra)
        if len(normal) >= 3 and normal not in PALAVRAS_VAZIAS:
            tokens.append(_raiz(normal))
    return tokens


def ordenar_por_bm25(palavras_chave, textos):
    """Índices dos textos com ao menos metade das palavras-chave, do maior para o menor BM25.

    As raízes casam por prefixo, como no resto da busca ('document' acha 'documentos'), e o IDF
    vem dos próprios candidatos: a palavra-chave rara entre eles pesa mais que a comum.
    """
    if not palavras_chave or not textos:
        return []
    corpus = [_tokens(texto) for texto in textos]
    prefixos = tuple(palavras_chave)
    termos = sorted({token for tokens in corpus for token in tokens if token.startswith(prefixos)})
    escolhidos = [i for i, texto in enumerate(textos) if _cobertura(palavras_chave, texto) >= COBERTURA_MINIMA_PALAVRAS]
    if not termos or not escolhidos:
        return []
    try:
        from rank_bm25 import BM25Plus

        notas = list(BM25Plus(corpus).get_scores(termos))
    except Exception as exc:  # sem o rank-bm25, a cobertura das palavras-chave ordena sozinha
        _registrar_auditoria("bm25_failed", erro=str(exc))
        notas = [_cobertura(palavras_chave, texto) for texto in textos]
    return sorted(escolhidos, key=lambda indice: notas[indice], reverse=True)


def fundir_rankings(*rankings, k=CONSTANTE_RRF):
    """Fusão por posição recíproca (RRF) de listas ordenadas de chaves; empates ficam na ordem de chegada."""
    notas = {}
    for ranking in rankings:
        for posicao, chave in enumerate(ranking, start=1):
            notas[chave] = notas.get(chave, 0.0) + 1.0 / (k + posicao)
    return sorted(notas, key=lambda chave: notas[chave], reverse=True)


def _metadado_pertence_ao_modo(metadado, modo):
    """Filtra fontes por modo sem esconder documentos da pasta geral."""
    if modo in {"geral", "padrão"}:
        return True
    caminho = os.path.normcase(os.path.realpath(metadado.get("caminho", "")))
    pastas_modo = [pasta for pasta in (MODOS_PASTAS.get(modo), PASTAS_OBSIDIAN.get(modo)) if pasta]
    if not pastas_modo:
        return True
    return any(
        caminho == os.path.normcase(os.path.realpath(pasta))
        or caminho.startswith(os.path.normcase(os.path.realpath(pasta)) + os.sep)
        for pasta in pastas_modo
    )


def auditar_biblioteca_conhecimento():
    """Audita qualidade documental com regras locais e retorna um relatório."""
    agora = time.time()
    arquivos = []
    for raiz, tipo in _fontes_raiz():
        if os.path.exists(raiz):
            if os.path.isfile(raiz):
                arquivos.append((raiz, tipo))
            else:
                arquivos.extend((caminho, tipo) for caminho in listar_arquivos_notelooklm(raiz) if os.path.isfile(caminho))

    hashes = {}
    relatorio = {"total": len(arquivos), "duplicadas": [], "sem_titulo": [], "sem_fonte": [], "antigas": [], "vazias": [], "possiveis_contradicoes": []}
    textos = []
    for caminho, tipo in arquivos:
        try:
            tamanho = os.path.getsize(caminho)
            if tamanho == 0:
                relatorio["vazias"].append(caminho)
                continue
            unidades = _extrair_unidades(caminho)
            texto = "\n".join(unidade for unidade, _ in unidades).strip()
            digest = _hash_conteudo(texto)
            hashes.setdefault(digest, []).append(caminho)
            nome = os.path.splitext(os.path.basename(caminho))[0].strip()
            if not nome or nome.lower() in {"sem título", "untitled", "new note"}:
                relatorio["sem_titulo"].append(caminho)
            extensao = os.path.splitext(caminho)[1].lower()
            if tipo in {"obsidian", "documento_local"} and extensao in {".md", ".txt", ".json"}:
                if "fonte:" not in texto.lower() and "source:" not in texto.lower() and "url:" not in texto.lower():
                    relatorio["sem_fonte"].append(caminho)
            if agora - os.path.getmtime(caminho) > IDADE_MAXIMA_DOCUMENTO_DIAS * 86400:
                relatorio["antigas"].append(caminho)
            if texto:
                textos.append((caminho, texto.lower()))
        except OSError:
            continue

    relatorio["duplicadas"] = [grupo for grupo in hashes.values() if len(grupo) > 1]
    for indice, (caminho_a, texto_a) in enumerate(textos):
        palavras_a = set(re.findall(r"[\wÀ-ÿ]+", texto_a))
        for caminho_b, texto_b in textos[indice + 1:]:
            palavras_b = set(re.findall(r"[\wÀ-ÿ]+", texto_b))
            semelhanca = len(palavras_a & palavras_b) / max(len(palavras_a | palavras_b), 1)
            if semelhanca > 0.75 and ((" não " in f" {texto_a} ") != (" não " in f" {texto_b} ")):
                relatorio["possiveis_contradicoes"].append([caminho_a, caminho_b, round(semelhanca, 2)])

    _registrar_auditoria("quality_audit", resumo={chave: len(valor) if isinstance(valor, list) else valor for chave, valor in relatorio.items()})
    return json.dumps(relatorio, ensure_ascii=False, indent=2)


def importar_fonte_local(caminho: str, categoria: str = "referencias") -> str:
    """Copia um arquivo suportado para a pasta relacionada da Vault e o indexa."""
    try:
        origem = os.path.realpath(caminho)
        if not os.path.isfile(origem):
            return f"Arquivo não encontrado: {caminho}"
        categoria_segura = re.sub(r"[^\w-]+", "_", (categoria or "referencias").strip().lower()).strip("._") or "referencias"
        pasta_destino = _resolver_pasta_obsidian(categoria_segura)
        os.makedirs(pasta_destino, exist_ok=True)
        destino = os.path.join(pasta_destino, os.path.basename(origem))
        shutil.copy2(origem, destino)
        indexar_pasta_documentos()
        return f"Fonte importada e indexada: {destino}"
    except Exception as exc:
        return f"Erro ao importar fonte: {exc}"


def importar_url(url: str, categoria: str = "fontes_web") -> str:
    """Salva uma página web como Markdown com URL e data para citação."""
    try:
        resposta = requests.get(url.strip(), headers={"User-Agent": "Beatriz/1.0"}, timeout=20)
        resposta.raise_for_status()
        soup = BeautifulSoup(resposta.text, "html.parser")
        titulo = soup.title.get_text(" ", strip=True) if soup.title else url
        texto = soup.get_text("\n", strip=True)
        nome = re.sub(r"[^\w-]+", "-", titulo.lower()).strip("-")[:80] or "pagina"
        categoria_segura = re.sub(r"[^\w-]+", "_", (categoria or "fontes_web").strip().lower()).strip("._") or "fontes_web"
        pasta = os.path.join(VAULT_OBSIDIAN, categoria_segura)
        if os.path.commonpath([os.path.realpath(VAULT_OBSIDIAN), os.path.realpath(pasta)]) != os.path.realpath(VAULT_OBSIDIAN):
            return "Erro ao importar URL: categoria inválida."
        os.makedirs(pasta, exist_ok=True)
        caminho = os.path.join(pasta, f"{nome}.md")
        with open(caminho, "w", encoding="utf-8") as arquivo:
            arquivo.write(f"---\ntitulo: {titulo}\nfonte: {url}\ndata_captura: {datetime.now().isoformat()}\n---\n\n{texto}\n")
        indexar_pasta_documentos()
        return f"Página importada e indexada: {caminho}"
    except Exception as exc:
        return f"Erro ao importar URL: {exc}"


ESPERA_WHISPER_VIDEO_SEGUNDOS = 600  # na primeira vez, o modelo da voz pode estar baixando (~480 MB)
_WHISPER_VIDEO = None  # só quando WHISPER_MODEL não é o modelo da voz
_WHISPER_VIDEO_LOCK = threading.Lock()


def _audio_do_video(caminho: str, taxa: int = 16000):
    """A trilha de áudio do arquivo em 16 kHz mono float32, o formato que o faster-whisper aceita.

    Decodifica com o PyAV aqui mesmo porque o `faster_whisper.decode_audio` (1.2.1) passa
    `metadata_errors` para o `av.open`, e o PyAV 19 tirou esse parâmetro (TypeError).
    """
    import av

    partes = []
    with av.open(caminho) as entrada:
        if not entrada.streams.audio:
            raise ValueError("o arquivo não tem trilha de áudio.")
        reamostrador = av.AudioResampler(format="flt", layout="mono", rate=taxa)
        try:
            # Pacote a pacote: um pacote corrompido perde só ele (o gerador do decode() pararia ali).
            # O último pacote, vazio, esvazia o decodificador.
            for pacote in entrada.demux(entrada.streams.audio[0]):
                try:
                    quadros = pacote.decode()
                except av.error.InvalidDataError:
                    continue
                for quadro in quadros:
                    partes.extend(q.to_ndarray().reshape(-1) for q in reamostrador.resample(quadro))
        except av.error.InvalidDataError:
            logger.warning("Vídeo truncado (%s); transcrevendo o que deu para ler.", os.path.basename(caminho))
        partes.extend(q.to_ndarray().reshape(-1) for q in reamostrador.resample(None))
    if not partes:
        raise ValueError("não deu para ler o áudio do arquivo.")
    return np.concatenate(partes)


def _modelo_whisper_video():
    """O Whisper que a voz já carrega (`voice._WHISPER`), para o vídeo não ocupar outros ~500 MB.

    Os dois podem transcrever ao mesmo tempo: o CTranslate2 aceita chamadas de várias threads.
    Outro modelo só se WHISPER_MODEL for diferente do da voz ou o dela não carregar, e uma vez só.
    """
    global _WHISPER_VIDEO
    import voice

    nome = os.getenv("WHISPER_MODEL", "small")
    if nome == voice.WHISPER_MODELO_VOZ:
        modelo = voice._WHISPER.obter(espera=ESPERA_WHISPER_VIDEO_SEGUNDOS)
        if modelo is not None:
            return modelo
    with _WHISPER_VIDEO_LOCK:
        if _WHISPER_VIDEO is None:
            from faster_whisper import WhisperModel

            # Na CPU, como a voz: com device="auto" ele escolhe a GPU e, sem as DLLs do
            # cuBLAS/cuDNN, a transcrição falha.
            _WHISPER_VIDEO = WhisperModel(nome, device="cpu", compute_type="int8")
        return _WHISPER_VIDEO


def importar_video_transcricao(caminho: str, categoria: str = "referencias") -> str:
    """Transcreve o vídeo (ou áudio) com o Whisper local e salva a transcrição na Vault."""
    try:
        import faster_whisper  # noqa: F401  (traz o PyAV, que já inclui o FFmpeg)
    except ImportError:
        return "Para importar vídeos, instale o 'faster-whisper'."
    try:
        origem = os.path.realpath(caminho)
        audio = _audio_do_video(origem)  # antes do modelo: um arquivo sem áudio falha logo
        segmentos, info = _modelo_whisper_video().transcribe(audio, language="pt")
        texto = "\n".join(segmento.text.strip() for segmento in segmentos if segmento.text.strip())
        categoria_segura = re.sub(r"[^\w-]+", "_", (categoria or "referencias").strip().lower()).strip("._") or "referencias"
        pasta = _resolver_pasta_obsidian(categoria_segura)
        os.makedirs(pasta, exist_ok=True)
        destino = os.path.join(pasta, os.path.splitext(os.path.basename(origem))[0] + ".md")
        with open(destino, "w", encoding="utf-8") as arquivo:
            arquivo.write(f"---\ntitulo: {os.path.basename(origem)}\nfonte: {origem}\nidioma: {info.language}\n---\n\n{texto}\n")
        indexar_pasta_documentos()
        return f"Vídeo transcrito e indexado: {destino}"
    except Exception as exc:
        return f"Erro ao transcrever vídeo: {exc}"


def importar_emails_importantes(limite: int = 50) -> str:
    """Importa mensagens marcadas como importantes do Gmail, se OAuth estiver configurado."""
    try:
        from config import GMAIL_CREDENTIALS_FILE, GMAIL_TOKEN_FILE
        from google.oauth2.credentials import Credentials
        from googleapiclient.discovery import build
        if not os.path.exists(GMAIL_TOKEN_FILE) or not os.path.exists(GMAIL_CREDENTIALS_FILE):
            return "Gmail não configurado: informe as credenciais OAuth e o token."
        credenciais = Credentials.from_authorized_user_file(GMAIL_TOKEN_FILE, ["https://www.googleapis.com/auth/gmail.readonly"])
        servico = build("gmail", "v1", credentials=credenciais, cache_discovery=False)
        mensagens = servico.users().messages().list(userId="me", q="is:important", maxResults=limite).execute().get("messages", [])
        pasta = os.path.join(VAULT_OBSIDIAN, "referencias", "emails")
        os.makedirs(pasta, exist_ok=True)
        for item in mensagens:
            dados = servico.users().messages().get(userId="me", id=item["id"], format="full").execute()
            cabecalhos = {h["name"].lower(): h["value"] for h in dados.get("payload", {}).get("headers", [])}
            corpo = dados.get("snippet", "")
            titulo = re.sub(r"[^\w-]+", "-", cabecalhos.get("subject", item["id"]).lower()).strip("-")[:80]
            with open(os.path.join(pasta, f"{item['id']}-{titulo or 'email'}.md"), "w", encoding="utf-8") as arquivo:
                arquivo.write(f"---\ntitulo: {cabecalhos.get('subject', '(sem assunto)')}\nfonte: gmail:{item['id']}\nde: {cabecalhos.get('from', '')}\ndata: {cabecalhos.get('date', '')}\n---\n\n{corpo}\n")
        indexar_pasta_documentos()
        return f"{len(mensagens)} e-mails importantes importados e indexados."
    except Exception as exc:
        return f"Erro ao importar e-mails: {exc}"


def iniciar_observador_memoria():
    """Observa documentos e Vault, reindexando alterações fora do callback."""
    global _OBSERVADOR_MEMORIA
    if _OBSERVADOR_MEMORIA is not None:
        return "Observador da memória já está ativo."
    try:
        from watchdog.events import FileSystemEventHandler
        from watchdog.observers import Observer

        class Handler(FileSystemEventHandler):
            def __init__(self):
                self.ultima_execucao = 0

            def on_any_event(self, event):
                if event.is_directory or time.time() - self.ultima_execucao < INTERVALO_OBSERVADOR_MEMORIA:
                    return
                self.ultima_execucao = time.time()
                threading.Thread(target=indexar_pasta_documentos, daemon=True).start()

        observador = Observer()
        handler = Handler()
        for raiz, _ in _fontes_raiz():
            if os.path.isdir(raiz):
                observador.schedule(handler, raiz, recursive=True)
        observador.start()
        _OBSERVADOR_MEMORIA = observador
        return "Observador da memória iniciado."
    except ImportError:
        return "Instale watchdog para ativar a atualização automática."
    except Exception as exc:
        return f"Erro ao iniciar observador: {exc}"


def parar_observador_memoria():
    global _OBSERVADOR_MEMORIA
    if _OBSERVADOR_MEMORIA is not None:
        _OBSERVADOR_MEMORIA.stop()
        _OBSERVADOR_MEMORIA.join(timeout=5)
        _OBSERVADOR_MEMORIA = None

def atualizar_memoria_usuario(novo_fato):
    fato = (novo_fato or "").strip()
    if not fato:
        return "Não encontrei um fato para guardar."

    try:
        os.makedirs(PASTA_MEMORIA_USUARIO, exist_ok=True)
        if not os.path.exists(ARQUIVO_MEMORIA_USUARIO):
            with open(ARQUIVO_MEMORIA_USUARIO, "w", encoding="utf-8") as arquivo:
                arquivo.write(
                    "---\n"
                    "tipo: memoria_usuario\n"
                    "criado_por: Beatriz\n"
                    "---\n\n"
                    "# Memória da Beatriz\n\n"
                )

        with open(ARQUIVO_MEMORIA_USUARIO, "a", encoding="utf-8") as arquivo:
            arquivo.write(f"- {datetime.now().isoformat(timespec='seconds')}: {fato}\n")

        indexar_pasta_documentos()
        return "Fato gravado na memória da Beatriz, no Obsidian/geral."
    except OSError as exc:
        return f"Não consegui gravar na memória do Obsidian: {exc}"

class IndicePalavras:
    """Índice invertido (raiz da palavra -> itens) para arquivos que só crescem.

    Lê apenas as linhas novas desde a última consulta, então cada busca custa
    o tamanho da resposta, não o tamanho do arquivo.
    """

    def __init__(self, obter_caminho, interpretar_linha):
        self.obter_caminho = obter_caminho  # função, para acompanhar mudanças do caminho configurado
        self.interpretar_linha = interpretar_linha  # linha -> (item, texto) ou None
        self.lock = threading.Lock()
        self.caminho = None
        self._zerar()

    def _zerar(self):
        self.posicao = 0
        self.itens = []
        self.indice = {}

    def _atualizar(self):
        caminho = self.obter_caminho()
        if caminho != self.caminho:
            self.caminho = caminho
            self._zerar()
        if not os.path.exists(caminho):
            self._zerar()
            return
        if os.path.getsize(caminho) < self.posicao:  # arquivo foi reescrito ou truncado
            self._zerar()
        with open(caminho, "r", encoding="utf-8", errors="ignore") as arquivo:
            arquivo.seek(self.posicao)
            for linha in iter(arquivo.readline, ""):
                if not linha.endswith("\n"):  # linha ainda sendo escrita
                    break
                self.posicao = arquivo.tell()
                interpretado = self.interpretar_linha(linha)
                if not interpretado:
                    continue
                item, texto = interpretado
                posicao_item = len(self.itens)
                self.itens.append(item)
                for raiz in raizes_do_texto(texto):
                    self.indice.setdefault(raiz, set()).add(posicao_item)

    def buscar(self, pergunta, limite):
        """Itens com mais palavras-chave em comum; palavras raras pesam mais e os recentes desempatam."""
        palavras_chave = extrair_palavras_chave(pergunta)
        if not palavras_chave:
            return []
        with self.lock:
            self._atualizar()
            total = max(len(self.itens), 1)
            pontuacao = {}
            peso_total = 0.0
            for raiz in palavras_chave:
                # Casa por prefixo: 'document' encontra 'documento' e 'documentos'.
                posicoes = set()
                for chave in [raiz] if raiz in self.indice else [c for c in self.indice if c.startswith(raiz)]:
                    posicoes |= self.indice[chave]
                # Palavra rara (ou ausente) pesa mais que palavra comum.
                peso = 1 + math.log(total / max(len(posicoes), 1))
                peso_total += peso
                for posicao in posicoes:
                    pontuacao[posicao] = pontuacao.get(posicao, 0) + peso
            # Exige ao menos metade do peso da pergunta, para uma palavra comum sozinha não bastar.
            minimo = peso_total * 0.5
            relevantes = [p for p in pontuacao if pontuacao[p] >= minimo]
            melhores = sorted(relevantes, key=lambda p: (pontuacao[p], p), reverse=True)[:limite]
            return [self.itens[p] for p in melhores]


def _interpretar_turno(linha):
    try:
        turno = json.loads(linha)
    except json.JSONDecodeError:
        return None
    if not turno.get("pergunta") or not turno.get("resposta"):
        return None
    return turno, f"{turno['pergunta']} {turno['resposta']}"


def _interpretar_fato(linha):
    if not linha.lstrip().startswith("-"):
        return None
    fato = linha.strip().lstrip("- ").strip()
    return (fato, fato) if fato else None


_INDICE_HISTORICO = IndicePalavras(lambda: ARQUIVO_HISTORICO_CONVERSAS, _interpretar_turno)
_INDICE_MEMORIA_USUARIO = IndicePalavras(lambda: ARQUIVO_MEMORIA_USUARIO, _interpretar_fato)


def buscar_memoria(pergunta, limite=8):
    """Fatos guardados sobre o usuário que compartilham palavras-chave com a pergunta."""
    if not (pergunta or "").strip() or not os.path.exists(ARQUIVO_MEMORIA_USUARIO):
        return "Nenhuma memória."
    try:
        resultados = _INDICE_MEMORIA_USUARIO.buscar(pergunta, limite)
    except OSError:
        return "Nenhuma memória."
    if resultados:
        return "Memórias:\n" + "\n".join(f"- {fato}" for fato in resultados)
    return "Nenhuma memória correspondente."


def registrar_conversa(pergunta, resposta):
    """Persiste cada turno para que conversas futuras possam recuperar contexto."""
    if not pergunta or not resposta:
        return
    try:
        with open(ARQUIVO_HISTORICO_CONVERSAS, "a", encoding="utf-8") as arquivo:
            arquivo.write(json.dumps({"pergunta": pergunta, "resposta": resposta, "timestamp": time.time()}, ensure_ascii=False) + "\n")
    except Exception:
        pass


def buscar_historico_conversa(pergunta, limite=6):
    """Retorna turnos anteriores com palavras relevantes para a pergunta atual."""
    if not os.path.exists(ARQUIVO_HISTORICO_CONVERSAS):
        return "Nenhum histórico de conversa."
    try:
        turnos = _INDICE_HISTORICO.buscar(pergunta, limite)
    except Exception:
        return "Nenhum histórico relacionado."
    return "\n\n".join(
        f"Usuário: {turno['pergunta']}\nBeatriz: {turno['resposta']}" for turno in turnos
    ) or "Nenhum histórico relacionado."


def notelooklm_assistente(pergunta, modo=None):
    """Responde em modo assistente/conversa usando documentos locais e memória do usuário."""
    if not pergunta or not str(pergunta).strip():
        return "Não recebi uma pergunta para responder."

    modo_ativo = (modo or MODO_CONVERSA_ATUAL or MODO_CONVERSA_PADRAO).strip().lower()
    deve_consultar = deve_consultar_documentos(pergunta, modo=modo_ativo)
    contexto_docs = buscar_nos_documentos_modo(pergunta, modo=modo_ativo) if deve_consultar else "A consulta automática de documentos está desativada neste momento."
    memoria = buscar_memoria(pergunta)
    historico = buscar_historico_conversa(pergunta)

    prompt = (
        "Você é a Beatriz, uma assistente pessoal e conversacional. "
        "Responda de forma direta, útil e natural, no seu jeito leve e informal de conversar. "
        "Para pedidos que envolvam escolha, planejamento ou opinião, diga o que você recomenda e por quê, "
        "como numa conversa, citando outra opção só se valer a pena; nada de formato de relatório. "
        "Para perguntas simples, responda em poucas frases. "
        "Use somente o contexto fornecido e não invente fatos. "
        "Quando usar documentos, preserve as citações no formato [Fonte: ...] e não remova o caminho ou o localizador. "
        "Se não houver informação suficiente, diga que não sabe com base no contexto.\n\n"
        f"Modo: {modo_ativo}\n\n"
        f"Contexto dos documentos:\n{contexto_docs}\n\n"
        f"Memória do usuário:\n{memoria}\n\n"
        f"Histórico relacionado de conversas anteriores:\n{historico}\n\n"
        f"Pergunta do usuário:\n{pergunta}\n\n"
        "Resposta:"
    )

    try:
        from llm_clients import llm_clients
        if llm_clients.is_ollama_available():
            return llm_clients.chat_beatriz(
                prompt,
                contexto=(
                    "Você está respondendo com memória local. Diferencie fatos recuperados "
                    "de inferências e, quando fizer uma recomendação, mostre resumidamente "
                    "por que ela é a mais adequada."
                ),
            )
    except Exception:
        pass

    linhas = []
    if contexto_docs and contexto_docs.strip() not in {"Nenhuma informação encontrada."}:
        linhas.append(contexto_docs)
    if memoria and memoria.strip() not in {"Nenhuma memória.", "Nenhuma memória correspondente."}:
        linhas.append(memoria)

    if linhas:
        return "\n\n".join(linhas[:3])

    return "Posso ajudar, mas ainda não encontrei contexto suficiente para responder com segurança neste modo de conversa."


def notelooklm_responder(pergunta, modo=None):
    """Alias focado em conversa/assistente para o NotelookLM."""
    return notelooklm_assistente(pergunta, modo=modo)